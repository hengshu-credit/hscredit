"""建模流程中的特征筛选汇总与过程数据。

``selection_report`` 只返回两个 DataFrame，供模型报告组合使用，不写入文件。
独立的快照存取函数用于显式归档，不参与筛选报告的数据生成。
"""

import copy
from datetime import date, datetime
import json
import logging
import math
from numbers import Integral
from os import PathLike
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple, Union
from uuid import uuid4

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

from ..core.selectors.reporting import (
    SelectionReport,
    collect_selection_report,
    SCHEMA_VERSION,
    DETAIL_COLUMNS,
    SUMMARY_COLUMNS,
    METRIC_COLUMNS,
    HISTORY_COLUMNS,
    _INT_COLUMNS,
    _FLOAT_COLUMNS,
    _BOOL_COLUMNS,
)
from ..exceptions import ValidationError
from .result import ReportResult

_TABLES = {"summary": "筛选汇总", "details": "特征决策", "metrics": "指标明细", "history": "迭代记录"}
_Path = Union[str, PathLike]
_ExcelPlan = Tuple[str, str, pd.DataFrame, int, int]
logger = logging.getLogger(__name__)


def to_report_result(report: SelectionReport) -> ReportResult:
    """将筛选快照转换为包含四张中文结果表的独立报告。

    :param report: 筛选报告快照。
    :return: 以筛选汇总、特征决策、指标明细、迭代记录为键的 ReportResult。
        保留列顺序、数据类型和原始精度；修改结果不会影响快照或筛选器。
    :raises ValidationError: 输入不是 SelectionReport。
    """
    if not isinstance(report, SelectionReport):
        raise ValidationError("报告对象必须为 SelectionReport")
    metadata = report.metadata
    result = ReportResult(mode="strict" if metadata.get("完整", True) else "best_effort", metadata=metadata)
    for attribute, name in _TABLES.items():
        frame = getattr(report, attribute)
        result.plan(name)
        result.success(name, frame)
    if not metadata.get("完整", True):
        result.failure("来源完整性", "；".join(map(str, metadata.get("诊断", []))) or "报告来源不完整")
    return result


def selection_report(
    selector: Any,
    *,
    relation: str = "auto",
    strict: bool = True,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """聚合筛选器核心层的 report 输出，返回汇总和具体过程两张表。

    :param selector: 已拟合的单筛选器、组合筛选器、Pipeline、筛选器列表或
        SelectionReport 快照。列表可混合上述筛选流程，也可使用 ``(名称, 筛选器)``
        为每个流程命名；不会创建筛选器，也不会调用 fit 或 transform。
    :param relation: 默认 'auto'，列表按独立筛选结果汇总，Pipeline 按实际步骤串行收集。
        串行列表须显式指定 'sequential' 并满足输入/输出字段连续；
        'intersection' / 'union' 按共同输入字段生成交集/并集根决策。
        'independent' 表示独立比较，不合成最终特征数。
    :param strict: 默认 True，拒绝未拟合或不完整的来源；False 保留已有记录，
        完整性及诊断写入两张表各自的 attrs，不添加诊断列。
    :return: ``(summary, process)`` 两个独立 DataFrame。
        summary 按阶段给出路径、父路径、筛选器、输入/选中/剔除特征数及保留率；
        根流程汇总与子阶段分别呈现，不能再次累加为唯一特征数量。
        process 用“记录类型”区分“特征决策”“指标明细”“迭代事件”，包含特征、
        指标值、阈值、筛选结果、原因及迭代轮次等原始字段。按汇总阶段顺序排列，
        同阶段按上述记录类型排列，各类型内部保持原始顺序；不做多对多拼接，
        不重复计算指标，不舍入，不修改筛选器。未记录的历史不会被补造。
        两表列名为中文，空表也保留固定列结构及 dtype。
    :raises ValidationError: 参数无效、字段链不连续或严格模式下来源不完整。
    :raises NotFittedError: 严格模式下筛选器尚未拟合。

    **参考样例**

    >>> summary, process = selection_report(selector=fitted_selector)
    >>> summary, process = selection_report([('缺失', null_selector), ('相关性', corr_selector)], relation='sequential')
    >>> summary, process = selection_report(fitted_pipeline)
    """
    _validate_bool(strict, "strict")
    if not isinstance(relation, str):
        raise ValidationError("relation 必须为筛选关系名称字符串")
    report = collect_selection_report(selector, relation=relation, strict=strict)
    return report.report(strict=strict)


def _validate_bool(value: bool, name: str) -> None:
    if not isinstance(value, bool):
        raise ValidationError(f"{name} 必须为布尔值")


def _validate_path(path: _Path) -> Path:
    if not isinstance(path, (str, PathLike)) or not str(path).strip():
        raise ValidationError("文件路径必须为非空字符串或路径对象")
    return Path(path).expanduser().resolve()


def _pack(value: Any) -> Any:
    if value is pd.NA or value is pd.NaT:
        return {"类型": "空值"}
    if isinstance(value, np.generic):
        value = value.item()
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        if math.isnan(value):
            return {"类型": "浮点特殊值", "值": "NaN"}
        if math.isinf(value):
            return {"类型": "浮点特殊值", "值": "+Inf" if value > 0 else "-Inf"}
        return value
    if isinstance(value, pd.Timestamp):
        return {"类型": "时间戳", "值": value.isoformat()}
    if isinstance(value, datetime):
        return {"类型": "日期时间", "值": value.isoformat()}
    if isinstance(value, date):
        return {"类型": "日期", "值": value.isoformat()}
    if isinstance(value, tuple):
        return {"类型": "元组", "值": [_pack(item) for item in value]}
    if isinstance(value, list):
        return [_pack(item) for item in value]
    if isinstance(value, dict):
        # 保留整数/字符串键，并避免用户字典与类型标记冲突。
        return {"类型": "映射", "值": [[_pack(key), _pack(item)] for key, item in value.items()]}
    raise ValidationError(f"筛选报告含不支持的 JSON 类型：{type(value).__name__}")


def _unpack(value: Any) -> Any:
    if isinstance(value, list):
        return [_unpack(item) for item in value]
    if not isinstance(value, dict):
        return value
    kind = value.get("类型")
    if kind == "空值":
        return pd.NA
    if kind == "浮点特殊值":
        return {"NaN": float("nan"), "+Inf": float("inf"), "-Inf": -float("inf")}[value["值"]]
    if kind == "日期时间":
        return datetime.fromisoformat(value["值"])
    if kind == "日期":
        return date.fromisoformat(value["值"])
    if kind == "时间戳":
        return pd.Timestamp(value["值"])
    if kind == "元组":
        return tuple(_unpack(item) for item in value["值"])
    if kind == "映射":
        return {_unpack(key): _unpack(item) for key, item in value["值"]}
    raise ValidationError("无法识别筛选报告 JSON 类型标记")


def _write_json(report: SelectionReport, path: _Path, run_id: str, sheet_mapping: List[Dict[str, Any]]) -> None:
    # 逐行写入，避免再构造全部表的字典列表副本。
    with Path(path).open("w", encoding="utf-8") as handle:
        handle.write('{"格式":"hscredit.selection-report","协议版本":')
        handle.write(str(SCHEMA_VERSION))
        handle.write(',"run_id":')
        json.dump(run_id, handle)
        handle.write(',"元数据":')
        json.dump(_pack(report.metadata), handle, ensure_ascii=False, allow_nan=False)
        handle.write(',"工作表映射":')
        json.dump(sheet_mapping, handle, ensure_ascii=False, allow_nan=False)
        handle.write(',"表":{')
        for index, (attribute, name) in enumerate(_TABLES.items()):
            if index:
                handle.write(",")
            json.dump(attribute, handle)
            handle.write(':{"名称":')
            json.dump(name, handle, ensure_ascii=False)
            frame = getattr(report, "_" + attribute)
            handle.write(',"列":')
            json.dump(list(frame.columns), handle, ensure_ascii=False)
            handle.write(',"行":[')
            for position, row in enumerate(frame.itertuples(index=False, name=None)):
                if position:
                    handle.write(",")
                json.dump([_pack(value) for value in row], handle, ensure_ascii=False, allow_nan=False)
            handle.write("]}")
        handle.write("}}")


def load_selection_report(path: _Path) -> SelectionReport:
    """读取筛选报告 JSON，恢复列顺序、字段类型、数值精度及元数据。

    :param path: save_selection_report 返回的 ``结果['路径']['json']``，支持路径对象。
    :return: 独立的 SelectionReport 快照，不执行代码或重新拟合。
    :raises ValidationError: JSON 内容损坏、类型标记错误或协议结构不兼容。
    :raises OSError: 文件不存在或不可读取。
    """
    path = _validate_path(path)
    try:
        with path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
        return _report_from_payload(payload)
    except ValidationError:
        raise
    except (ValueError, TypeError, KeyError, OverflowError) as error:
        raise ValidationError("筛选报告 JSON 内容或字段类型无效") from error


def _report_from_payload(payload: Any) -> SelectionReport:
    if (
        not isinstance(payload, dict)
        or payload.get("格式") != "hscredit.selection-report"
        or type(payload.get("协议版本")) is not int
        or payload.get("协议版本") != SCHEMA_VERSION
    ):
        raise ValidationError("不是支持的筛选报告 JSON 协议")
    metadata = _unpack(payload.get("元数据"))
    if (
        not isinstance(metadata, dict)
        or type(metadata.get("schema_version")) is not int
        or metadata.get("schema_version") != SCHEMA_VERSION
        or type(metadata.get("完整")) is not bool
    ):
        raise ValidationError("筛选报告元数据的协议版本或完整性状态无效")
    if not isinstance(metadata.get("诊断"), list):
        raise ValidationError("筛选报告诊断必须为列表")
    schemas = {
        "summary": SUMMARY_COLUMNS,
        "details": DETAIL_COLUMNS,
        "metrics": METRIC_COLUMNS,
        "history": HISTORY_COLUMNS,
    }
    if not isinstance(payload.get("表"), dict) or set(payload["表"]) != set(schemas):
        raise ValidationError("筛选报告必须包含全部四个固定结构表")
    values = {}
    for attribute in _TABLES:
        table = payload["表"][attribute]
        if (
            not isinstance(table, dict)
            or table.get("列") != schemas[attribute]
            or not isinstance(table.get("行"), list)
        ):
            raise ValidationError(f"筛选报告 {attribute} 的列结构与协议不一致")
        if any(not isinstance(row, list) or len(row) != len(table["列"]) for row in table["行"]):
            raise ValidationError(f"筛选报告 {attribute} 的行长度与列结构不一致")
        rows = [[_unpack(value) for value in row] for row in table["行"]]
        for row in rows:
            for column, value in zip(table["列"], row):
                if value is None or value is pd.NA:
                    continue
                valid = True
                if column in _INT_COLUMNS:
                    valid = type(value) is int
                elif column in _FLOAT_COLUMNS:
                    valid = type(value) in {int, float}
                elif column in _BOOL_COLUMNS:
                    valid = type(value) is bool
                if not valid:
                    raise ValidationError(f"筛选报告 {attribute} 的 {column} 列值类型无效，不能静默转换为缺失")
        values[attribute] = pd.DataFrame(rows, columns=table["列"], dtype=object)
    return SelectionReport(metadata=metadata, **values)


def _sheet_name(name: str, used: Set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", str(name)).strip("'") or "报告"
    base = cleaned[:31]
    candidate, suffix = base, 1
    while candidate.lower() in used:
        suffix += 1
        candidate = f"{base[:31-len(str(suffix))-1]}_{suffix}"
    used.add(candidate.lower())
    return candidate


def _excel_value(value: Any) -> Any:
    if value is pd.NA or value is pd.NaT or value is None:
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return "正无穷" if value == float("inf") else "负无穷" if value == -float("inf") else None
    if isinstance(value, (dict, tuple, list)):
        return json.dumps(_pack(value), ensure_ascii=False, allow_nan=False)
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if isinstance(value, int) and not isinstance(value, bool) and abs(value) >= 10**15:
        return str(value)
    return value


def _plan_excel(
    report: SelectionReport, max_rows_per_sheet: int, max_excel_cells: int, max_sheets: int, by_selector: bool
) -> List[_ExcelPlan]:
    if (
        not isinstance(max_rows_per_sheet, Integral)
        or isinstance(max_rows_per_sheet, bool)
        or not 1 <= max_rows_per_sheet <= 1048575
    ):
        raise ValidationError("max_rows_per_sheet 必须为 1 至 1048575 的整数（不含表头）")
    for name, value in (("max_excel_cells", max_excel_cells), ("max_sheets", max_sheets)):
        if isinstance(value, bool) or not isinstance(value, Integral) or value < 1:
            raise ValidationError(f"{name} 必须为正整数")
    _validate_bool(by_selector, "by_selector")
    max_rows_per_sheet = int(max_rows_per_sheet)
    frames = [(name, getattr(report, "_" + attribute)) for attribute, name in _TABLES.items()]
    status = pd.DataFrame(
        [
            {"项目": "报告完整", "内容": report._metadata.get("完整", True)},
            {"项目": "诊断", "内容": report._metadata.get("诊断", [])},
        ]
    )
    snapshots = report._metadata.get("阶段快照", {})
    parameters = pd.DataFrame(
        [
            {"阶段路径": stage, "参数": snapshot.get("拟合参数", {}), "拟合标识": snapshot.get("fit_id")}
            for stage, snapshot in snapshots.items()
        ],
        columns=["阶段路径", "参数", "拟合标识"],
    )
    frames = [("执行状态", status)] + frames + [("参数说明", parameters)]
    if by_selector:
        for path, part in report._details.groupby("阶段路径", sort=False):
            frames.append((str(path), part))
    total_cells = sum(
        (len(frame) + max(1, math.ceil(len(frame) / max_rows_per_sheet))) * len(frame.columns) for _, frame in frames
    )
    if total_cells > max_excel_cells:
        raise ValidationError(
            f"Excel 预计 {total_cells} 个单元格，超出 max_excel_cells={max_excel_cells}；请仅保存完整 JSON 或提高预算"
        )
    plans, used = [], set()
    for label, frame in frames:
        chunks = max(1, math.ceil(len(frame) / max_rows_per_sheet))
        for chunk in range(chunks):
            start, end = chunk * max_rows_per_sheet, min(len(frame), (chunk + 1) * max_rows_per_sheet)
            sheet = _sheet_name(label if chunks == 1 else f"{label}_{chunk+1}", used)
            plans.append((sheet, label, frame, start, end))
    if len(plans) > max_sheets:
        raise ValidationError(f"Excel 需要 {len(plans)} 个工作表，超出 max_sheets={max_sheets}")
    return plans


def _write_excel(plans: Sequence[_ExcelPlan], path: _Path, run_id: Optional[str] = None) -> None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment

    workbook = Workbook(write_only=True)
    workbook.properties.identifier = run_id
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="244A73")
    text_alignment = Alignment(vertical="top", wrap_text=True)
    number_alignment = Alignment(vertical="top")
    widths = {
        "阶段路径": 42,
        "父路径": 36,
        "特征": 30,
        "阶段名称": 26,
        "筛选器": 28,
        "筛选方法": 24,
        "指标名称": 32,
        "关联特征": 30,
        "筛选原因": 54,
        "停止原因": 54,
        "原因": 54,
        "统计口径": 46,
        "文本值": 38,
        "参数": 64,
        "补充信息": 64,
        "候选子集": 48,
        "拟合标识": 36,
        "内容": 64,
    }
    try:
        for sheet_name, _, frame, start, end in plans:
            sheet = workbook.create_sheet(sheet_name)
            sheet.freeze_panes = "A2"
            sheet.row_dimensions[1].height = 30
            for position, column in enumerate(frame.columns, start=1):
                sheet.column_dimensions[get_column_letter(position)].width = widths.get(column, 18)
            if len(frame.columns):
                sheet.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}{end-start+1}"

            def cells(values, header=False):
                row = []
                for value in values:
                    value = _excel_value(value)
                    if isinstance(value, str) and len(value) > 32767:
                        raise ValidationError("Excel 单元格文本超出 32767 字符；请仅保存 JSON，不会静默截断")
                    cell = WriteOnlyCell(sheet, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"  # 字符串始终按文本写入，不能执行用户提供的公式。
                    cell.alignment = text_alignment if isinstance(value, str) else number_alignment
                    if header:
                        cell.font = header_font
                        cell.fill = header_fill
                    row.append(cell)
                return row

            sheet.append(cells(frame.columns, header=True))
            for row in frame.iloc[start:end].itertuples(index=False, name=None):
                sheet.append(cells(row))
        workbook.save(path)
    except BaseException:
        # 关闭流式 XML 生成器后清理临时文件；不能用清理异常覆盖原始错误。
        for sheet in workbook.worksheets:
            try:
                if not sheet.closed:
                    sheet.close()
                writer = getattr(sheet, "_writer", None)
                if writer is not None and Path(writer.out).exists():
                    writer.cleanup()
            except Exception as error:
                logger.warning("筛选报告 Excel 临时资源清理失败：%s", error)
        raise
    finally:
        workbook.close()


def save_selection_report(
    report: SelectionReport,
    path: _Path,
    *,
    formats: Union[str, Sequence[str]] = ("xlsx", "json"),
    overwrite: bool = False,
    max_rows_per_sheet: int = 1048575,
    max_excel_cells: int = 5000000,
    max_sheets: int = 200,
    by_selector: bool = False,
    allow_incomplete: bool = False,
) -> Dict[str, Any]:
    """将完整报告保存为同一版本制品，返回实际路径、校验摘要和映射。

    **参数**

    :param report: SelectionReport 快照。
    :param path: 输出名称或带 .xlsx/.json 后缀的期望名称；实际文件位于版本目录。
    :param formats: 支持 xlsx/json，默认同时生成。
    :param overwrite: 默认值为 ``False`` 。是否允许更新同名发布清单，不删除旧版本
    :param max_rows_per_sheet: 每个工作表最多数据行数，默认 1048575，不含表头。
    :param max_excel_cells: Excel 总单元格预算，默认 5000000，包含表头。
    :param max_sheets: 本次最多生成工作表数，默认 200。
    :param by_selector: 默认 False；是否额外按筛选阶段导出特征决策表。
    :param allow_incomplete: 默认 False；是否允许归档带诊断的不完整快照。
    :return: 含路径、清单、协议版本、制品、工作表映射、表行数、完整、诊断的字典；
        run_id 是版本标识。始终使用返回路径定位文件；建模流程中的两张结果表请用 selection_report。
    :raises ValidationError: 参数无效、报告不完整或超过 Excel 容量预算。

    **参考样例**

    >>> paths = report.save('输出/筛选报告', overwrite=True)
    >>> paths['路径']['xlsx']
    """
    if not isinstance(report, SelectionReport):
        raise ValidationError("保存对象必须为 SelectionReport")
    for name, value in (("overwrite", overwrite), ("by_selector", by_selector), ("allow_incomplete", allow_incomplete)):
        _validate_bool(value, name)
    if not report._metadata.get("完整", True) and not allow_incomplete:
        raise ValidationError("报告不完整，默认拒绝保存；如需诊断文档请显式设置 allow_incomplete=True")
    if isinstance(formats, str):
        formats = (formats,)
    if (
        not isinstance(formats, (list, tuple))
        or not formats
        or any(not isinstance(item, str) or item not in {"xlsx", "json"} for item in formats)
        or len(set(formats)) != len(formats)
    ):
        raise ValidationError("formats 只能包含不重复的 xlsx 和 json")
    output = _validate_path(path)
    if output.suffix and output.suffix.lower() not in {".xlsx", ".json"}:
        raise ValidationError("筛选报告文件后缀只支持 .xlsx 或 .json，也可不指定后缀")
    name = output.stem if output.suffix else output.name
    plans = (
        _plan_excel(report, max_rows_per_sheet, max_excel_cells, max_sheets, by_selector) if "xlsx" in formats else []
    )
    mapping = [
        {"工作表": sheet, "来源": source, "起始行": start, "结束行": end, "行数": end - start}
        for sheet, source, _, start, end in plans
    ]
    # 仅在归档时加载发布事务，常规报告不依赖制品运行时。
    from ..skills_runtime.artifacts import ArtifactTransaction

    transaction = ArtifactTransaction({"directory": str(output.parent), "name": name, "overwrite": overwrite})
    run_id = uuid4().hex
    paths = {}
    with transaction:
        if "json" in formats:
            staged = transaction.stage_path(f"{name}.json")
            _write_json(report, staged, run_id, mapping)
            paths["json"] = transaction.publish(staged, artifact_type="selection_report_json")["path"]
        if "xlsx" in formats:
            staged = transaction.stage_path(f"{name}.xlsx")
            _write_excel(plans, staged, run_id=run_id)
            paths["xlsx"] = transaction.publish(staged, artifact_type="selection_report_excel")["path"]
        metadata = transaction.stage_path(f"{name}.index.json")
        metadata.write_text(
            json.dumps(
                {
                    "run_id": run_id,
                    "协议版本": SCHEMA_VERSION,
                    "工作表映射": mapping,
                    "表行数": {key: len(getattr(report, "_" + key)) for key in _TABLES},
                    "完整": report._metadata.get("完整", True),
                },
                ensure_ascii=False,
                allow_nan=False,
            ),
            encoding="utf-8",
        )
        transaction.publish(metadata, artifact_type="selection_report_index")
    manifest = json.loads(transaction.manifest_path.read_text(encoding="utf-8"))
    return {
        "路径": paths,
        "清单": str(transaction.manifest_path),
        "run_id": run_id,
        "协议版本": SCHEMA_VERSION,
        "制品": manifest["制品"],
        "工作表映射": mapping,
        "表行数": {key: len(getattr(report, "_" + key)) for key in _TABLES},
        "完整": report._metadata.get("完整", True),
        "诊断": copy.deepcopy(report._metadata.get("诊断", []) + transaction.cleanup_errors),
    }


__all__ = ["selection_report", "save_selection_report", "load_selection_report", "to_report_result"]
