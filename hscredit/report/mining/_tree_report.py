"""复用模型报告样本统计与 scorecardpipeline 策略报告布局的树挖掘导出。"""

from copy import copy, deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import numpy as np
import pandas as pd

from ...utils.targets import bind_targets
from ...utils.dates import date_groups
from .._sample_stats import build_sample_stats_table


def _report_metric_name(table, column):
    """兼容指标在首层的统计宽表、指标在末层的规则宽表和单层长表。"""
    if not isinstance(column, tuple):
        return column
    if ("分箱详情", "节点编号") in table.columns or column[0] in {"统计详情", "分箱详情"}:
        return column[-1]
    return column[0]


def _write_report_table(data, excel_writer, *, amount_basis=False, merge_columns=None, **kwargs):
    """通过 dataframe2excel 写表，按指标含义设置格式，不舍入数值或调整列宽。"""
    from ...excel import dataframe2excel
    from openpyxl.utils import get_column_letter

    # 传入完整列键，避免宽表按最后一层标签误匹配指标。
    for option in ("percent_cols", "custom_cols", "condition_cols", "color_cols", "left_cols", "right_cols"):
        requested = kwargs.get(option)
        if requested:
            kwargs[option] = [c for c in data if c in requested or _report_metric_name(data, c) in requested]
    end_row, end_col = dataframe2excel(data, excel_writer, **kwargs)
    if data.empty:
        return end_row, end_col
    worksheet = excel_writer.get_sheet_by_name(kwargs.get("sheet_name") or "Sheet1")
    first_row = end_row - len(data)
    first_col = kwargs.get("start_col", 2) + (data.index.nlevels if kwargs.get("index", False) else 0)
    basis_col = next((c for c in data if _report_metric_name(data, c) == "统计口径"), None)
    currency_rows = data[basis_col].eq("金额").to_numpy() if basis_col is not None else np.zeros(len(data), dtype=bool)
    counts = {"样本数", "样本总数", "好样本数", "坏样本数", "命中数", "累积好样本数", "累积坏样本数"}
    decimals = {
        "AUC",
        "PSI",
        "LIFT值",
        "累积LIFT值",
        "指标IV值",
        "分档IV值",
        "分档WOE值",
        "风险拒绝比",
        "累计风险拒绝比",
        "累积风险拒绝比",
        "分裂增益",
        "原始IV值",
        "有效IV值",
    }
    identifiers = {"数据集", "数据分组", "目标标签", "逾期标签", "分箱", "分箱标签", "配置"}
    for offset, column in enumerate(data.columns):
        metric = str(_report_metric_name(data, column))
        if metric in identifiers:
            continue
        numeric = pd.api.types.is_numeric_dtype(data[column])
        percent = (
            column in (kwargs.get("percent_cols") or ())
            or metric.endswith(("率", "占比", "KS值"))
            or metric in {"KS", "坏账改善", "累积坏账改善"}
        )
        if percent:
            number_format = "0.00%"
        elif "金额" in metric or (metric in counts and amount_basis):
            number_format = "#,##0.00"
        elif metric in decimals:
            number_format = "0.0000"
        elif metric in counts or metric == "节点编号" or pd.api.types.is_integer_dtype(data[column]):
            number_format = "#,##0"
        elif numeric:
            number_format = "0.0000"
        else:
            continue
        letter = get_column_letter(first_col + offset)
        excel_writer.set_number_format(worksheet, f"{letter}{first_row}:{letter}{end_row-1}", number_format)
        if metric in counts and not amount_basis:
            for position in np.flatnonzero(currency_rows):
                excel_writer.set_number_format(worksheet, f"{letter}{first_row+position}", "#,##0.00")
    if merge_columns:
        _merge_group_cells(data, excel_writer, worksheet, first_row, first_col, merge_columns)
    return end_row, end_col


def _merge_group_cells(data, writer, worksheet, first_row, first_col, columns):
    """按父分组限定连续合并范围，避免相同子组名称跨数据集合并。"""
    keys = [c for name in columns for c in data if _report_metric_name(data, c) == name]
    prefix = []
    for column in keys:
        prefix.append(column)
        boundaries = data[prefix].ne(data[prefix].shift()).any(axis=1).to_numpy()
        starts = np.flatnonzero(boundaries)
        for start, end in zip(starts, list(starts[1:]) + [len(data)]):
            if end - start > 1:
                col = first_col + data.columns.get_loc(column)
                writer.merge_cells(worksheet, (col, first_row + start), (col, first_row + end - 1))


def _strategy_table(table, long_format=True):
    """仅转换 Excel 展示字段；指标直接保留 Rule.report 的数值。"""
    multi = isinstance(table.columns, pd.MultiIndex)

    def key(name):
        return ("分箱详情", name) if multi else name

    root = table.loc[table[key("节点编号")].eq(0)]
    result = table.drop(columns=[key(c) for c in ["规则分类", "分箱"] if key(c) in table]).copy()
    rename = {"指标名称": "组合策略", "样本总数": "命中数", "样本占比": "命中率"}
    result = result.rename(columns=rename, level=1) if multi else result.rename(columns=rename)
    for column in table.columns:
        if (column[-1] if multi else column) == "坏样本率":
            output = (column[0], "样本整体坏率") if multi else "样本整体坏率"
            result[output] = root[column].iloc[0] if len(root) else np.nan
    first = [key(c) for c in ["节点编号", "指标含义", "组合策略", "是否叶子", "是否调整"] if key(c) in result]
    last = [key(c) for c in ["入参字段", "字段含义"] if key(c) in result]
    middle = [c for c in result if c not in first + last]
    if multi:
        groups = dict.fromkeys(column[0] for column in middle)
        middle = [column for group in groups for column in middle if column[0] == group]
    return format_rule_tables(result[first + middle + last], long_format)


def format_rule_tables(result, long_format=True):
    """宽表计算结果转为规则长表；False 保留既有宽表接口。"""
    if not isinstance(long_format, (bool, np.bool_)):
        raise ValueError("long_format 必须为布尔值")
    if not long_format:
        return result
    if isinstance(result, dict):
        return {name: format_rule_tables(table, True) for name, table in result.items()}
    if isinstance(result, (list, tuple)):
        return [format_rule_tables(table, True) for table in result]
    if isinstance(result.columns, pd.MultiIndex):
        common = result["分箱详情"].copy()
        labels = list(dict.fromkeys(c[0] for c in result if c[0] != "分箱详情"))
        tables = []
        for ordinal, label in enumerate(labels):
            table = common.copy()
            for metric in result[label]:
                table[metric] = result[(label, metric)].to_numpy()
            table["目标标签"] = label
            table["_row"] = np.arange(len(table))
            table["_label"] = ordinal
            tables.append(table)
        result = pd.concat(tables, ignore_index=True).sort_values(["_row", "_label"], kind="stable")
        result = result.drop(columns=["_row", "_label"]).reset_index(drop=True)
    result = result.rename(columns={"指标名称": "组合策略"})
    first = [c for c in ["节点编号", "指标含义", "组合策略", "是否叶子", "是否调整", "目标标签"] if c in result]
    return result[first + [c for c in result if c not in first]]


def tree_report_excel(
    tree,
    datasets,
    *,
    target=None,
    overdue=None,
    dpds=None,
    del_grey=False,
    leaf_only=False,
    feature_map=None,
    overdue_operator=">",
    save=None,
    sheet="组合策略汇总",
    excel_writer=None,
    date=None,
    freq="M",
    project_desc=None,
    data_source=None,
    long_format=True,
    methods=("quantile", "uniform", "tree"),
    bin_params=None,
    group=None,
    apply_rule_filter=True,
    **kwargs,
):
    """标签只求值一次，在单页生成基本信息、挖掘摘要与订单/金额口径的树图及规则表。

    :param group: 默认 None；分组字段或字段列表，可与 date 同时使用。
    :param apply_rule_filter: 默认 True；应用初始化的规则筛选条件，False 输出完整规则。
    :param tree: 已拟合的树规则挖掘器。
    :param datasets: DataFrame、数据集列表/元组或名称到 DataFrame 的字典。
    :param target: 默认 None；支持标签列名、Series、ndarray、list、callable 或标签映射；
        Series 按索引对齐，ndarray/list 按行位置对应，callable 接收完整数据集；
        与 overdue + dpds 互斥，未提供标签时使用树的训练目标。
    :param overdue: 默认 None；逾期字段或字段列表，与 dpds 同时提供。
    :param dpds: 默认 None；逾期阈值或阈值列表。
    :param del_grey: 默认 False；是否逐标签剔除灰样本。
    :param leaf_only: 默认 False 返回全部节点效果；True 仅返回叶节点效果。
    :param feature_map: 默认 None；特征名称到中文含义的映射。
    :param overdue_operator: 默认 '>'；逾期比较符，支持 '>'、'>='、'<'、'<='。
    :param save: 默认 None；导出 Excel 路径。
    :param sheet: 默认 '组合策略汇总'；完整挖掘报告的单页名称。
    :param excel_writer: 默认 None；可复用的 hscredit ExcelWriter。
    :param date: 默认 None；用于样本时间分布的日期字段。
    :param freq: 默认 'M'；日期统计频率。
    :param project_desc: 默认 None；项目目标和取样逻辑说明。
    :param data_source: 默认 None；数据来源说明。
    :param methods: 默认等频、等距、决策树三种方法，用于规则特征分箱效果页。
    :param bin_params: 默认 None；公共或按方法配置的分箱参数，传给 feature_binning_summary。
    :param long_format: 默认 True；规则表按标签纵向输出，False 保留规则宽表。样本统计与分布固定使用宽表。
    :param kwargs: 传给树 report 的评估参数，如 amount、round_digits。
    :return: 与输入数据集结构对应的中文效果表、列表或字典。
    """
    from ...excel import ExcelWriter
    from ._rule_filter import filter_rule_tables

    dataframe2excel = _write_report_table

    if excel_writer is not None and not isinstance(excel_writer, ExcelWriter):
        raise TypeError("excel_writer 必须是 hscredit.excel.ExcelWriter 实例")
    if not isinstance(sheet, str) or not sheet or len(sheet) > 24 or any(c in sheet for c in "[]:*?/\\"):
        raise ValueError("sheet 必须是 1 至 24 字符的有效 Excel 工作表名称")
    frames = (
        datasets
        if isinstance(datasets, dict)
        else {
            f"数据集{i + 1}": frame
            for i, frame in enumerate(datasets if isinstance(datasets, (list, tuple)) else [datasets])
        }
    )
    if not frames:
        raise ValueError("导出挖掘报告至少需要一个数据集")
    full_reports, order_reports, y_maps, date_values, plot_labels, contexts = {}, {}, [], [], {}, {}
    amount = kwargs.get("amount")
    for name, frame in frames.items():
        actual_target = target
        if target is None and overdue is None:
            classes = getattr(tree, "classes_", [0, 1])
            actual_target = (
                tree.target
                if set(classes).issubset({0, 1})
                else {f"类别_{value}": frame[tree.target].eq(value).astype(int) for value in classes}
            )
        prepared, bound = bind_targets(frame, actual_target, overdue, dpds, del_grey, overdue_operator)
        contexts[name] = (prepared, bound)
        y_maps.append(bound.y_map(prepared, style="at"))
        plot_labels[name] = bound.names(style="space")[0]
        if date is not None:
            if date not in frame:
                raise ValueError(f"数据集 {name} 缺少日期字段 {date}")
            date_values.append(date_groups(frame[date], freq))
        full_reports[name] = tree.report(
            prepared,
            target=bound,
            leaf_only=False,
            feature_map=feature_map,
            overdue_operator=overdue_operator,
            long_format=False,
            apply_rule_filter=False,
            **kwargs,
        )
        order_reports[name] = (
            full_reports[name]
            if amount is None
            else tree.report(
                prepared,
                target=bound,
                leaf_only=False,
                feature_map=feature_map,
                overdue_operator=overdue_operator,
                long_format=False,
                apply_rule_filter=False,
                **{k: v for k, v in kwargs.items() if k != "amount"},
            )
        )
    results = {}
    for name, table in full_reports.items():
        leaf_column = ("分箱详情", "是否叶子") if isinstance(table.columns, pd.MultiIndex) else "是否叶子"
        results[name] = table.loc[table[leaf_column].eq("是")].reset_index(drop=True) if leaf_only else table
    writer = excel_writer if excel_writer is not None else ExcelWriter()
    basic = sheet
    names = [sheet]
    for name in names:
        if name in writer.workbook.sheetnames:
            del writer.workbook[name]
        writer.get_sheet_by_name(name)
    percent = [
        "好样本占比",
        "坏样本占比",
        "命中率",
        "坏样本率",
        "样本整体坏率",
        "坏账改善",
        "准确率",
        "精确率",
        "召回率",
        "F1分数",
    ]
    try:
        ws = writer.get_sheet_by_name(basic)
        row, _ = writer.insert_value2sheet(
            ws, (2, 2), value="一、基本信息", style="header_middle", align={"horizontal": "left"}
        )
        description = project_desc or "使用树规则划分客群，并比较各数据集的规则命中效果"
        row, _ = writer.insert_value2sheet(
            ws, (row + 2, 2), value="1、项目目标", style="header_middle", align={"horizontal": "left"}
        )
        row, _ = writer.insert_value2sheet(
            ws, (row, 2), value=description, style="middle", align={"horizontal": "left"}
        )
        info = [
            {"统计项": "模型名称", "统计内容": type(tree).__name__},
            {"统计项": "取样逻辑", "统计内容": project_desc or ""},
            {"统计项": "数据源", "统计内容": data_source or ""},
            {"统计项": "指标口径", "统计内容": "、".join(y_maps[0])},
            {
                "统计项": "规则范围",
                "统计内容": (
                    "按训练集订单指标筛选，条件见树的参数配置"
                    if apply_rule_filter and getattr(tree, "rule_filter", None) is not None
                    else ("仅叶节点" if leaf_only else "全部节点（父子规则存在重叠）")
                ),
            },
        ]
        if date is not None:
            dates = pd.concat([pd.to_datetime(frame[date]) for frame in frames.values()]).dropna()
            info.insert(
                0,
                {
                    "统计项": "样本区间",
                    "统计内容": "" if dates.empty else f"{dates.min():%Y-%m-%d} ~ {dates.max():%Y-%m-%d}",
                },
            )
        if kwargs.get("amount") is not None:
            info.append(
                {
                    "统计项": "金额口径",
                    "统计内容": f"金额维度按 {kwargs['amount']} 加权；样本总数、好坏样本数表示对应金额，AUC/KS 同样加权",
                }
            )
        row, _ = dataframe2excel(
            pd.DataFrame(info),
            writer,
            sheet_name=basic,
            title="2、数据样本描述",
            start_row=row + 2,
            left_cols=["统计项", "统计内容"],
        )
        orders, amounts = _sample_tables(tree, frames, y_maps, full_reports, amount, False)
        row = _write_pair(
            writer,
            sheet,
            row + 2,
            "3、数据样本统计",
            orders,
            amounts,
            "按各标签有效样本统计；AUC/KS 使用最终树预测。",
        )
        if date is not None or group is not None:
            grouped_orders, grouped_amounts = _distribution_tables(
                tree, frames, contexts, date_values, amount, False, feature_map, overdue_operator, kwargs, group
            )
            row = _write_pair(
                writer,
                sheet,
                row + 2,
                "4、样本分布情况",
                grouped_orders,
                grouped_amounts,
                "与数据样本统计使用相同指标，按日期及指定字段分组计算。",
                merge_columns=["数据集", "数据分组"] + ([group] if isinstance(group, str) else list(group or [])),
            )
        training_data = tree._data if tree._data is not None else next(iter(frames.values()))
        if tree._data is None:
            training = next(iter(order_reports.values()))
            training_amount = next(iter(full_reports.values())) if amount else None
        else:
            training = tree.get_rule_table(
                training_data, round_digits=False, long_format=False, apply_rule_filter=False
            )
            if amount and amount not in training_data:
                raise ValueError(f"模型未保留训练集金额字段 {amount}，请使用包含该字段的训练数据重新 fit")
            training_amount = (
                tree.get_rule_table(
                    training_data, amount=amount, round_digits=False, long_format=False, apply_rule_filter=False
                )
                if amount
                else None
            )
        results = filter_rule_tables(tree, results, apply_rule_filter, training=training)
        selected = _select_rules(filter_rule_tables(tree, training, apply_rule_filter, training=training), leaf_only)
        node_col = ("分箱详情", "节点编号") if isinstance(selected.columns, pd.MultiIndex) else "节点编号"
        node_order = selected[node_col].tolist()
        row = _write_mining_summary(tree, writer, sheet, row + 2, training, training_amount, len(node_order))
        with TemporaryDirectory(prefix="hscredit-tree-report-") as temporary:
            train_label = "训练集" if tree._data is not None else "首数据集"
            row = _write_tree_pair(
                tree,
                training,
                training_amount,
                writer,
                sheet,
                row + 2,
                f"三、策略详情：{train_label}",
                temporary,
                node_order,
                long_format,
            )
            for name, orders in order_reports.items():
                row = _write_tree_pair(
                    tree,
                    orders,
                    full_reports[name] if amount else None,
                    writer,
                    sheet,
                    row + 2,
                    f"四、最终树应用：{name}",
                    temporary,
                    node_order,
                    long_format,
                    plot_labels[name],
                )
        feature_sheet = _write_feature_binning(
            tree, writer, sheet, contexts, methods, bin_params, amount, del_grey, overdue_operator, node_order
        )
        _fit_report_layout(writer, names + [feature_sheet])
        writer.adjust_columns_width(writer.get_sheet_by_name(feature_sheet), max_width=25)
        if save is not None:
            images = [(image, image._data()) for ws in writer.workbook for image in ws._images]
            for image, payload in images:
                image.ref = BytesIO(payload)
            try:
                writer.save(str(save), close=excel_writer is None)
            finally:
                for image, payload in images:
                    image.ref = BytesIO(payload)
    finally:
        if excel_writer is None:
            writer.workbook.close()
    if isinstance(datasets, dict):
        return format_rule_tables(results, long_format)
    result = list(results.values()) if isinstance(datasets, (list, tuple)) else next(iter(results.values()))
    return format_rule_tables(result, long_format)


def _fit_report_layout(writer, names):
    """复用模型报告标题布局，沿用模板列宽，关闭自动换行。"""
    from .._excel_layout import adjust_report_title_merges

    for name in names:
        ws = writer.get_sheet_by_name(name)
        ws.sheet_view.showGridLines = False
        adjust_report_title_merges(ws)
        for row in ws:
            for cell in row:
                if cell.value is not None:
                    alignment = copy(cell.alignment)
                    alignment.wrap_text = False
                    cell.alignment = alignment


def _select_rules(table, leaf_only):
    column = ("分箱详情", "是否叶子") if isinstance(table.columns, pd.MultiIndex) else "是否叶子"
    return table.loc[table[column].eq("是")] if leaf_only else table


def _tree_score(tree, data, class_index=1):
    """使用最终树预测；手工二叉树复用叶规则路由与已存储的训练概率。"""
    if hasattr(tree, "predict_proba"):
        probability = tree.predict_proba(data)
        if probability.shape[1] == 1:
            classes = getattr(tree, "classes_", getattr(getattr(tree, "clf", None), "classes_", [0]))
            return np.full(len(data), float(classes[0] == 1))
        return probability[:, min(class_index, probability.shape[1] - 1)]
    from ...core.viz.tree_plots import _extract_tree_data

    nodes = {node["node_id"]: node for node in _extract_tree_data(tree)["nodes"]}
    values = np.full(len(data), nodes[0]["bad_rate"], dtype=float)
    for rule in tree.get_rules():
        node = int(rule.name.rsplit("_", 1)[1])
        values[np.asarray(rule.predict(data), dtype=bool)] = nodes[node]["bad_rate"]
    return values


def _prediction_metrics(tree, frames, y_maps, amount):
    """样本统计的 AUC/KS 直接复用 ModelReport，金额口径使用其 sample_weight。"""
    from ..model_report import ModelReport

    features = getattr(tree, "features_", None) or getattr(tree, "_feature_list", None)
    result = {}
    for label_index, label in enumerate(y_maps[0]):
        class_index = label_index if str(label).startswith("类别_") else 1
        for basis, weight in [("订单", None)] + ([("金额", amount)] if amount else []):
            eligible = {
                name: (frame, y_map[label])
                for (name, frame), y_map in zip(frames.items(), y_maps)
                if len(frame)
                and pd.notna(y_map[label]).any()
                and (weight is None or frame[weight].to_numpy()[np.asarray(pd.notna(y_map[label]))].sum() > 0)
            }
            metrics = pd.DataFrame()
            if eligible:
                report = ModelReport(
                    tree,
                    datasets=eligible,
                    features=features,
                    prediction_method=lambda X: _tree_score(tree, X, class_index),
                    score_direction="higher_risk",
                    sample_weight=weight,
                    del_grey=True,
                    n_jobs=1,
                )
                metrics = report.get_metrics().set_index("统计项")
            for metric in ["AUC", "KS"]:
                if metric in metrics.index or metric in {"AUC", "KS"}:
                    values = [
                        metrics.loc[metric, name] if metric in metrics.index and name in metrics else np.nan
                        for name in frames
                    ]
                    result[(basis, label, metric)] = pd.to_numeric(pd.Series(values), errors="coerce").tolist()
    return result


def _tree_view(tree, table, amount=False):
    """固定节点、边和阈值，仅叠加当前数据集命中量与第一标签的效果。"""
    from ...core.viz.tree_plots import _extract_tree_data

    view = deepcopy(_extract_tree_data(tree))
    multi = isinstance(table.columns, pd.MultiIndex)
    meta = lambda name: ("分箱详情", name) if multi else name
    bad_col = next(c for c in table if (c[-1] if multi else c) == "坏样本率")
    label = bad_col[0] if multi else str(tree.target)
    metric = lambda name: (label, name) if multi else name
    # 剔灰模式的总数在各标签下独立保存。
    count_col = meta("样本总数") if meta("样本总数") in table else metric("样本总数")
    share_col = meta("样本占比") if meta("样本占比") in table else metric("样本占比")
    rows = table.set_index(meta("节点编号"))
    view["total_samples"] = rows.loc[0, count_col]
    view["overall_bad_rate"] = rows.loc[0, bad_col]
    for node in view["nodes"]:
        row = rows.loc[node["node_id"]]
        n, bad = row[count_col], row[metric("坏样本数")]
        node.update(
            n_samples=n,
            sample_pct=row[share_col],
            bad_count=bad,
            good_count=n - bad,
            bad_rate=row[bad_col],
            lift=row[metric("LIFT值")],
        )
        unit = "金额" if amount else "订单数"
        node["metric_rows"] = [
            (f"命中{unit}", f"{n:,.2f}" if amount else f"{n:,.0f}"),
            (f"{unit}占比", f"{row[share_col]:.2%}"),
            ("坏样本金额" if amount else "坏订单数", f"{bad:,.2f}" if amount else f"{bad:,.0f}"),
            ("金额坏率" if amount else "坏订单率", f"{row[bad_col]:.2%}"),
            ("LIFT值", f"{row[metric('LIFT值')]:.4f}"),
        ]
    return view, label


def _write_mining_summary(tree, writer, sheet, row, training, training_amount, selected_count):
    dataframe2excel = _write_report_table
    from ...core.viz.tree_plots import _extract_tree_data

    nodes = _extract_tree_data(tree)["nodes"]
    leaves = sum(node["is_leaf"] for node in nodes)
    summary = pd.DataFrame(
        [
            {
                "规则数（含根节点）": len(nodes),
                "报告规则数": selected_count,
                "叶子节点数": leaves,
                "中间节点数": len(nodes) - leaves,
                "树深度（根为0）": max(node["depth"] for node in nodes),
                "人工调整节点数": int(
                    training[("分箱详情", "是否调整") if isinstance(training.columns, pd.MultiIndex) else "是否调整"]
                    .eq("是")
                    .sum()
                ),
            }
        ]
    )
    row, _ = dataframe2excel(summary, writer, sheet_name=sheet, title="二、规则挖掘摘要", start_row=row)
    model = tree if hasattr(tree, "get_params") else getattr(tree, "clf", getattr(tree, "_sklearn_clf", None))
    parameters = model.get_params(deep=False) if model is not None else {}
    if hasattr(tree, "_branch_limit"):
        if tree.max_n_bins is not None or tree.max_branches is None:
            parameters.pop("max_branches", None)
            parameters["max_n_bins"] = tree._branch_limit()
        else:
            parameters.pop("max_n_bins", None)
    parameters["features"] = getattr(tree, "features_", None) or getattr(tree, "_feature_list", None)
    parameters["rule_filter"] = getattr(tree, "rule_filter", None)
    params = pd.DataFrame([{"参数": name, "配置": str(value)} for name, value in parameters.items()])
    row, _ = dataframe2excel(
        params, writer, sheet_name=sheet, title="树的参数配置", start_row=row + 2, left_cols=["配置"]
    )
    history = getattr(tree, "edit_history_", [])
    if history:
        row, _ = dataframe2excel(
            pd.DataFrame(history),
            writer,
            sheet_name=sheet,
            title="手工调整记录（节点编号为操作当时编号）",
            start_row=row + 2,
        )
    else:
        adjusted = [node for node in nodes if node["is_manual"]]
        records = pd.DataFrame(
            [{"节点编号": n["node_id"], "当前分裂": n["condition"], "阈值": n["threshold_text"]} for n in adjusted],
            columns=["节点编号", "当前分裂", "阈值"],
        )
        row, _ = dataframe2excel(records, writer, sheet_name=sheet, title="当前人工调整节点", start_row=row + 2)
    rows = []
    cuts = [0, 0.25, 0.5, 1, 1.5, 2, 3, 5, np.inf]
    for basis, table in [("订单", training)] + ([("金额", training_amount)] if training_amount is not None else []):
        multi = isinstance(table.columns, pd.MultiIndex)
        leaf_col = ("分箱详情", "是否叶子") if multi else "是否叶子"
        for lift in [c for c in table if (c[-1] if multi else c) == "LIFT值"]:
            for low, high in zip(cuts, cuts[1:]):
                selected = table.loc[table[lift].ge(low) & table[lift].lt(high)]
                rows.append(
                    {
                        "口径": basis,
                        "标签": lift[0] if multi else tree.target,
                        "LIFT区间": f"[{low:g}, {high:g})",
                        "规则数": len(selected),
                        "叶子节点数": int(selected[leaf_col].eq("是").sum()),
                        "中间节点数": int(selected[leaf_col].eq("否").sum()),
                    }
                )
    row, _ = dataframe2excel(
        pd.DataFrame(rows), writer, sheet_name=sheet, title="训练集 LIFT 区间规则分布", start_row=row + 2
    )
    return row


def _sample_tables(tree, frames, y_maps, amount_reports, amount, long_format, target_specific_totals=None):
    """整体与分组统计共用模型报告样本统计和性能入口，不混合订单/金额分母。"""
    labels = list(y_maps[0])
    filtered = any(pd.isna(values).any() for mapping in y_maps for values in mapping.values())
    if target_specific_totals is not None:
        filtered = target_specific_totals
    orders, _ = build_sample_stats_table(
        list(frames), y_maps, labels, flat_total_col="样本总数", target_specific_totals=filtered, n_jobs=1
    )
    amounts = orders.copy() if amount else None
    if amounts is not None:
        multi = isinstance(amounts.columns, pd.MultiIndex)
        for source in ["样本总数", "好样本数", "坏样本数", "坏样本率"]:
            for index, label in enumerate(labels):
                values = []
                for table in amount_reports.values():
                    if isinstance(table.columns, pd.MultiIndex):
                        groups = [col[0] for col in table if col[-1] == "坏样本率"]
                        key = ("分箱详情", source) if ("分箱详情", source) in table else (groups[index], source)
                    else:
                        key = source
                    values.append(table[key].iloc[0])
                key = (
                    (("统计详情", source) if ("统计详情", source) in amounts else (source, label)) if multi else source
                )
                amounts[key] = values
    performance = _prediction_metrics(tree, frames, y_maps, amount)
    for table, basis in [(orders, "订单"), (amounts, "金额")]:
        if table is None:
            continue
        for metric in ["AUC", "KS"]:
            for label in labels:
                if (basis, label, metric) in performance:
                    key = (metric, label) if isinstance(table.columns, pd.MultiIndex) else metric
                    table[key] = performance.get((basis, label, metric), [np.nan] * len(frames))
    return _stats_format(orders, labels, long_format), _stats_format(amounts, labels, long_format)


def _stats_format(table, labels, long_format, detail_group="统计详情", index_name="数据集"):
    if table is None or not isinstance(table.columns, pd.MultiIndex):
        return table
    table = table.copy()
    if index_name is not None:
        table.insert(0, (detail_group, index_name), table.index.to_list())
    table = table.reset_index(drop=True)
    if not long_format:
        return table
    common = table[detail_group]
    outputs = []
    for ordinal, label in enumerate(labels):
        output = common.copy()
        output.insert(1, "目标标签", label)
        for metric, current in table.columns:
            if metric != detail_group and current == label:
                output[metric] = table[metric, current].to_numpy()
        output["_row"] = np.arange(len(output))
        output["_label"] = ordinal
        outputs.append(output)
    return (
        pd.concat(outputs, ignore_index=True)
        .sort_values(["_row", "_label"], kind="stable")
        .drop(columns=["_row", "_label"])
        .reset_index(drop=True)
    )


def _distribution_tables(
    tree, frames, contexts, date_values, amount, long_format, feature_map, operator, kwargs, group=None
):
    group_columns = [group] if isinstance(group, str) else list(group or [])
    if len(set(group_columns)) != len(group_columns) or any(c in {"数据集", "数据分组"} for c in group_columns):
        raise ValueError("group 字段不能重复，也不能使用保留名称 数据集、数据分组")
    names = (["数据分组"] if date_values else []) + group_columns
    if not names:
        raise ValueError("group 至少需要一个分组字段")
    order_tables, amount_tables = [], []
    filtered = any(bound.needs_filtering(prepared) for prepared, bound in contexts.values())
    for ordinal, (name, frame) in enumerate(frames.items()):
        missing = [c for c in group_columns if c not in frame]
        if missing:
            raise ValueError(f"数据集 {name} 缺少分组字段：{missing}")
        groups = frame[group_columns].reset_index(drop=True).copy()
        if date_values:
            groups.insert(0, "数据分组", np.asarray(date_values[ordinal]))
        groups = groups.astype(object).where(groups.notna(), "缺失")
        for values, positions in groups.groupby(names, sort=True, observed=True).indices.items():
            values = (values,) if len(names) == 1 else values
            prepared, bound = contexts[name]
            subset = prepared.iloc[positions]
            grouped_frames = {name: frame.iloc[positions]}
            y_maps = [bound.y_map(subset, style="at")]
            amount_reports = {}
            if amount:
                amount_reports[name] = tree.report(
                    subset,
                    target=bound,
                    feature_map=feature_map,
                    overdue_operator=operator,
                    long_format=False,
                    apply_rule_filter=False,
                    **kwargs,
                )
            orders, amounts = _sample_tables(
                tree, grouped_frames, y_maps, amount_reports, amount, long_format, target_specific_totals=filtered
            )
            for table, outputs in [(orders, order_tables), (amounts, amount_tables)]:
                if table is not None:
                    for index, (column, value) in enumerate(zip(names, values), 1):
                        key = ("统计详情", column) if isinstance(table.columns, pd.MultiIndex) else column
                        table.insert(index, key, value)
                    outputs.append(table)

    def combine(tables):
        if not tables:
            return None
        table = pd.concat(tables, ignore_index=True)
        key = ("统计详情", "数据集") if isinstance(table.columns, pd.MultiIndex) else "数据集"
        return pd.concat([table.loc[table[key].eq(name)] for name in frames], ignore_index=True)

    return combine(order_tables), combine(amount_tables)


def _write_pair(writer, sheet, row, title, left, right, description, figures=None, merge_columns=None):
    """左右并列，中间留一个空列；上方统一标题和说明，下方保留一个空行。"""
    dataframe2excel = _write_report_table
    from PIL import Image

    ws = writer.get_sheet_by_name(sheet)
    left_col = 2
    right_col = left_col + len(left.columns) + 1
    end_col = right_col + len(right.columns) - 1 if right is not None else left_col + len(left.columns) - 1
    writer.insert_value2sheet(
        ws, (row, left_col), value=title, style="header_middle", align={"horizontal": "left"}, end_space=(row, end_col)
    )
    writer.insert_value2sheet(
        ws,
        (row + 1, left_col),
        value=description,
        style="middle",
        end_space=(row + 1, end_col),
        align={"horizontal": "left"},
    )
    subtitle = row + 3
    tables = [(left, left_col, "订单维度"), (right, right_col, "金额维度")]
    table_row = subtitle + 2
    for index, (table, col, label) in enumerate(tables):
        if table is None:
            continue
        writer.insert_value2sheet(
            ws,
            (subtitle, col),
            value=label,
            style="header_middle",
            align={"horizontal": "left"},
            end_space=(subtitle, col + len(table.columns) - 1),
        )
        if figures:
            path = figures[index]
            with Image.open(path) as image:
                width = min(1400, len(table.columns) * 61)
                size = (width, round(width * image.height / image.width))
            end, _ = writer.insert_pic2sheet(ws, path, (subtitle + 2, col), figsize=size)
            ws._images[-1].ref = BytesIO(Path(path).read_bytes())
            table_row = max(table_row, end + 1)
    end_row = table_row
    percent = [
        "命中率",
        "样本占比",
        "好样本占比",
        "坏样本占比",
        "坏样本率",
        "金额坏率",
        "坏账改善",
        "样本整体坏率",
        "准确率",
        "精确率",
        "召回率",
        "F1分数",
        "KS",
    ]
    for table, col, basis in tables:
        if table is not None:
            end, _ = dataframe2excel(
                table,
                writer,
                sheet_name=sheet,
                start_row=table_row,
                start_col=col,
                amount_basis=basis == "金额维度",
                percent_cols=percent,
                condition_cols=["坏样本率", "金额坏率", "LIFT值"],
                decimal=None,
                left_cols=["指标含义", "组合策略"],
                merge_columns=merge_columns,
            )
            end_row = max(end_row, end)
    return end_row + 1


def _write_tree_pair(
    tree, orders, amounts, writer, sheet, row, title, temporary, node_order, long_format, plot_label=None
):
    from ...core.viz import plot_tree
    import matplotlib.pyplot as plt

    paths, tables = [], []
    for basis, full in enumerate([orders, amounts]):
        if full is None:
            tables.append(None)
            continue
        view, label = _tree_view(tree, full, bool(basis))
        path = str(Path(temporary) / f"tree-{row}-{basis}.png")
        figure = plot_tree(view, save=path)
        plt.close(figure)
        paths.append(path)
        node_col = ("分箱详情", "节点编号") if isinstance(full.columns, pd.MultiIndex) else "节点编号"
        strategy = _strategy_table(full, False).set_index(node_col, drop=False).loc[node_order].reset_index(drop=True)
        tables.append(format_rule_tables(strategy, long_format))
    description = (
        f"{len(node_order)} 条节点规则，保持训练集规则顺序；树图标签：{plot_label or label}。"
        "命中率按各标签有效样本计算，LIFT=命中坏率/整体坏率；金额表按金额加权。"
    )
    return _write_pair(writer, sheet, row, title, tables[0], tables[1], description, paths)


def _write_feature_binning(tree, writer, sheet, contexts, methods, bin_params, amount, del_grey, operator, node_order):
    """复用 swap_out_report 使用的 feature_binning_summary，计算规则涉及的字段。"""
    from ..feature_analyzer import feature_binning_summary

    dataframe2excel = _write_report_table
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    features = list(
        dict.fromkeys(
            feature
            for rule in tree.get_rules(leaf_only=False)
            if int(rule.name.rsplit("_", 1)[1].lstrip("N")) in node_order
            for feature in rule.feature_names_in_
        )
    )
    name = f"{sheet}_特征分箱"
    if name in writer.workbook.sheetnames:
        del writer.workbook[name]
    ws = writer.get_sheet_by_name(name)
    ws.freeze_panes = None
    if not features:
        writer.insert_value2sheet(
            ws, "B2", value="当前报告规则未使用分裂特征", style="header_middle", align={"horizontal": "left"}
        )
        return name
    summaries, details = [], {}
    for dataset, (prepared, bound) in contexts.items():
        parts = []
        for basis, weight in [("订单", None)] + ([("金额", amount)] if amount else []):
            tables, summary = feature_binning_summary(
                prepared,
                features=features,
                methods=methods,
                bin_params=bin_params,
                target=bound,
                del_grey=del_grey,
                overdue_operator=operator,
                amount=weight,
                long_format=True,
                round_digits=False,
                n_jobs=1,
                verbose=0,
            )
            # 保留 feature_binning_summary 的指标 × 标签宽表；KS/AUC 使用原始特征的统一指标入口。
            summary = summary.drop(columns=[c for c in summary if c[0] == "分档KS值"])
            ranking = _feature_ranking_metrics(prepared, bound, features, weight)
            for (metric, label), values in ranking.items():
                if any(pd.notna(value) for value in values.values()):
                    summary[(metric, label)] = summary[("分箱详情", "指标名称")].map(values)
            metadata = [c for c in summary if c[0] == "分箱详情"]
            ordered = metadata + [
                c
                for metric in ["KS", "AUC", "LIFT值", "指标IV值", "坏样本数", "坏样本率"]
                for c in summary
                if c[0] == metric
            ]
            summary = summary[ordered]
            summary.insert(0, ("分箱详情", "数据集"), dataset)
            summary.insert(1, ("分箱详情", "统计口径"), basis)
            summaries.append(summary)
            for feature, methods_tables in tables.items():
                for method, table in methods_tables.items():
                    table = table.copy()
                    table.insert(0, "分箱方法", method)
                    table.insert(0, "统计口径", basis)
                    parts.append(table)
        details[dataset] = pd.concat(parts, ignore_index=True)
    summary = pd.concat(summaries, ignore_index=True)
    metadata_names = ["数据集", "指标名称", "统计口径", "分箱方法"]
    metadata = [("分箱详情", column) for column in metadata_names]
    # 数据集严格保持入参顺序；订单在金额前，其余按名称排序。
    order = pd.DataFrame(
        {
            "dataset": summary[metadata[0]].map({name: i for i, name in enumerate(contexts)}),
            "feature": summary[metadata[1]],
            "basis": summary[metadata[2]].map({"订单": 0, "金额": 1}),
            "method": summary[metadata[3]],
        }
    ).sort_values(["dataset", "feature", "basis", "method"], kind="stable")
    summary = summary.loc[order.index].reset_index(drop=True)
    labels = [column[1] for column in summary if column[0] == "坏样本率"]
    summary = summary[
        metadata
        + [column for column in summary if column[0] == "分箱详情" and column not in metadata]
        + [
            (metric, label)
            for metric in ["KS", "AUC", "LIFT值", "指标IV值", "坏样本数", "坏样本率"]
            for label in labels
            if (metric, label) in summary
        ]
    ]

    def write_filter_table(table, row, col, title=None):
        wide = isinstance(table.columns, pd.MultiIndex)
        end, _ = dataframe2excel(
            table,
            writer,
            sheet_name=name,
            start_row=row,
            start_col=col,
            title=title,
            percent_cols=["坏样本率", "样本占比", "好样本占比", "坏样本占比", "坏账改善"],
            condition_cols=["LIFT值", "指标IV值"],
            decimal=None,
            merge_columns=metadata_names if wide else None,
        )
        # 宽表保留合并分层表头，使用工作表筛选；明细长表使用多个原生 Table。
        if wide:
            header_row = end - len(table) - 1
            writer.add_auto_filter(
                ws, f"{get_column_letter(col)}{header_row}:" f"{get_column_letter(col + len(table.columns)-1)}{end-1}"
            )
            return end
        # Worksheet 的 AutoFilter 只有一个范围；用原生 Table 为各数据集提供筛选按钮。
        identifiers = {key for worksheet in writer.workbook for key in worksheet.tables}
        ordinal = len(identifiers) + 1
        while f"TreeFeatureBins{ordinal}" in identifiers:
            ordinal += 1
        header_row = row + 2 if title else row
        ref = f"{get_column_letter(col)}{header_row}:{get_column_letter(col + len(table.columns)-1)}{end-1}"
        native_table = Table(displayName=f"TreeFeatureBins{ordinal}", ref=ref)
        native_table.tableStyleInfo = TableStyleInfo(
            name=None, showFirstColumn=False, showLastColumn=False, showRowStripes=False, showColumnStripes=False
        )
        ws.add_table(native_table)
        return end

    end = write_filter_table(summary, 2, 2, title="规则特征效果汇总")
    row, col = end + 3, 2
    for dataset, table in details.items():
        writer.insert_value2sheet(
            ws,
            (row, col),
            value=str(dataset),
            style="header_middle",
            align={"horizontal": "left"},
            end_space=(row, col + len(table.columns) - 1),
        )
        write_filter_table(table, row + 2, col)
        col += len(table.columns) + 1
    return name


def _feature_ranking_metrics(data, targets, features, amount):
    """复用 core.metrics.ks/auc；无法形成有效数值评分时不伪造指标。"""
    from ...core.metrics import ks, auc

    result = {}
    weights = pd.to_numeric(data[amount], errors="coerce") if amount is not None else None
    for label, target in targets.y_map(data, style="at").items():
        result[("KS", label)], result[("AUC", label)] = {}, {}
        for feature in features:
            ks_value, auc_value = np.nan, np.nan
            if pd.api.types.is_numeric_dtype(data[feature]):
                scores = pd.to_numeric(data[feature], errors="coerce").to_numpy(dtype=float, na_value=np.nan)
                valid = target.isin([0, 1]).to_numpy() & np.isfinite(scores)
                if weights is not None:
                    valid &= np.isfinite(weights.to_numpy(dtype=float, na_value=np.nan)) & weights.gt(0).to_numpy()
                if valid.any():
                    y = target.to_numpy()[valid]
                    score = scores[valid]
                    sample_weight = weights.to_numpy()[valid] if weights is not None else None
                    ks_value = ks(y, score, sample_weight=sample_weight)
                    if np.unique(y).size == 2:
                        auc_value = auc(y, score, sample_weight=sample_weight)
            result[("KS", label)][feature] = ks_value
            result[("AUC", label)][feature] = auc_value
    return result
