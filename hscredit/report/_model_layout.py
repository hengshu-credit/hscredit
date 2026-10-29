"""模型报告的参数展开与长文本排版，保留数值和既有主题。"""

import copy
import math
import unicodedata

import pandas as pd
import numpy as np
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


def record_identifier(value):
    """定位字段按文本导出，避免长整型标识在 Excel 中被舍入或显示为科学计数。"""
    if pd.isna(value):
        return None
    if isinstance(value, (float, np.floating)) and np.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value)


def configuration_table(mapping):
    """按配置路径逐项展示嵌套信息，避免把整个模型字典写进一个单元格。"""
    rows = []

    def append(name, value):
        if isinstance(value, dict) and value:
            for key, item in value.items():
                append(f"{name}.{key}" if name else str(key), item)
        elif isinstance(value, (list, tuple)) and value:
            for index, item in enumerate(value):
                append(f"{name}[{index}]", item)
        else:
            rows.append({"配置项": name, "配置值": str(value) if isinstance(value, (dict, list, tuple)) else value})

    for name, value in mapping.items():
        append(str(name), value)
    return pd.DataFrame(rows, columns=["配置项", "配置值"])


def process_display_table(report, title, table):
    """公共过程表保持原结构，Excel 对宽记录采用逐条配置展示。"""
    if title == "模型口径":
        return configuration_table(report.model_context_)
    if title.endswith("训练记录"):
        rows = []
        for index, record in enumerate(table.to_dict(orient="records")):
            detail = configuration_table(record)
            detail.insert(0, "记录序号", index + 1)
            rows.append(detail)
        return pd.concat(rows, ignore_index=True) if rows else table
    return table.apply(
        lambda column: column.map(lambda value: str(value) if isinstance(value, (dict, list, tuple)) else value)
    )


def fit_metadata_sheet(worksheet, default_column_width=None):
    """限定参数/过程页宽度并匹配换行高度；不调整图表和业务指标格式。"""
    worksheet.sheet_view.showGridLines = False
    # 模板 C:XFD 共用一条列样式。调整 C 的列宽不能意外把所有空白列一起拉宽。
    last_column = worksheet.max_column
    for key, dimension in list(worksheet.column_dimensions.items()):
        first = dimension.min or worksheet[key + "1"].column
        last = dimension.max or first
        if last <= first or first > last_column:
            continue
        for column in range(first, min(last, last_column) + 1):
            name = get_column_letter(column)
            if name not in worksheet.column_dimensions or name == key:
                part = copy.copy(dimension)
                part.index = name
                part.min = part.max = column
                worksheet.column_dimensions[name] = part
        if last > last_column:
            name = get_column_letter(last_column + 1)
            tail = copy.copy(dimension)
            tail.index, tail.min, tail.max = name, last_column + 1, last
            if default_column_width is not None:
                tail.width = default_column_width
            worksheet.column_dimensions[name] = tail

    def width(value):
        return sum(2 if unicodedata.east_asian_width(char) in {"F", "W"} else 1 for char in str(value))

    merged = {
        cell.coordinate
        for area in worksheet.merged_cells.ranges
        for row in worksheet.iter_rows(
            min_row=area.min_row, max_row=area.max_row, min_col=area.min_col, max_col=area.max_col
        )
        for cell in row
    }
    for column in range(2, worksheet.max_column + 1):
        cells = [
            row[0]
            for row in worksheet.iter_rows(min_col=column, max_col=column)
            if not isinstance(row[0], MergedCell) and row[0].value is not None and row[0].coordinate not in merged
        ]
        desired = max(
            (max(width(line) for line in str(cell.value).splitlines() or [""]) + 2 for cell in cells), default=12
        )
        limit = 60 if column == 3 else 40 if column == 2 else 36
        worksheet.column_dimensions[get_column_letter(column)].width = min(limit, max(14, desired))
    for row in worksheet:
        required = worksheet.row_dimensions[row[0].row].height or 15
        for cell in row:
            if isinstance(cell, MergedCell) or not isinstance(cell.value, str) or cell.coordinate in merged:
                continue
            available = max(8, worksheet.column_dimensions[cell.column_letter].width - 2)
            lines = sum(max(1, math.ceil(width(line) / available)) for line in cell.value.splitlines() or [""])
            if lines > 1:
                alignment = copy.copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "center"
                cell.alignment = alignment
                required = max(required, lines * ((cell.font.sz or 10) + 3) + 3)
        worksheet.row_dimensions[row[0].row].height = min(409, required)
