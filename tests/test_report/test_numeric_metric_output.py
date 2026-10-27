"""结果表只展示数值；全拒绝风险拒绝比为 NaN，不附加说明。"""

import numpy as np
import pandas as pd
import pytest

from hscredit.core.metrics import compute_bin_stats, add_margins
from hscredit.core.rules import Rule
from hscredit.report import feature_bin_stats


@pytest.mark.parametrize("weighted", [False, True])
@pytest.mark.parametrize("target,expected", [([0, 1, 0, 1], 1.0), ([0, 0, 0, 0], 0.0)])
def test_full_rejection_keeps_improvement_but_risk_ratio_is_nan(weighted, target, expected):
    options = dict(target_type="amount_weighted", amount=[1, 2, 3, 4]) if weighted else {}
    table = compute_bin_stats([0] * 4, target, bin_labels=["全体"], round_digits=False, **options)
    row = table.iloc[0]
    assert row["坏账改善"] == row["累积坏账改善"] == expected
    assert pd.isna(row["风险拒绝比"]) and pd.isna(row["累计风险拒绝比"])
    assert not any("状态" in name or "原因" in name or "适用" in name for name in table.columns)
    assert table.drop(columns=["分箱标签"]).select_dtypes(exclude="number").empty
    total = add_margins(table).iloc[-1]
    assert total["坏账改善"] == total["风险拒绝比"] == 0
    assert total["累积坏账改善"] == expected
    assert pd.isna(total["累计风险拒绝比"])


def test_rule_and_feature_outputs_do_not_add_status_columns():
    data = pd.DataFrame({"评分": [0, 1, 2, 3], "FPD": [0, 1, 0, 1]})
    table = Rule("评分 >= 0").report(data, target="FPD").set_index("分箱")
    assert table.loc["命中", "坏账改善"] == 1
    assert pd.isna(table.loc["命中", "风险拒绝比"])
    assert table.loc["未命中", "风险拒绝比"] == 0
    result = feature_bin_stats(data, "评分", target="FPD", method="quantile", max_n_bins=2, margins=True, n_jobs=1)
    assert pd.isna(result.iloc[-1]["累计风险拒绝比"])
    for frame in [table, result]:
        assert not any("状态" in str(name) or "原因" in str(name) for name in frame.columns)


def test_excel_keeps_numbers_and_blank_nan_without_metric_notes(tmp_path):
    from openpyxl import load_workbook
    from hscredit.excel import dataframe2excel

    table = compute_bin_stats([0, 0], [0, 1], bin_labels=["全体"])
    output = tmp_path / "数值指标.xlsx"
    dataframe2excel(table, str(output), sheet_name="指标", start_row=1, start_col=1, index=False)
    book = load_workbook(output, data_only=True)
    sheet = book["指标"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("坏账改善") + 1).value == 1
    assert sheet.cell(2, headers.index("风险拒绝比") + 1).value is None
    assert not any(isinstance(cell.value, str) and "不适用" in cell.value for row in sheet for cell in row)
    book.close()
