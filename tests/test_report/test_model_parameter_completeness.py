"""参数页必须保留可复现配置；包装器、长列表和嵌套配置不能退化成类型摘要。"""

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

from hscredit import LogisticRegression, ProbabilityCalibrator, CalibratedModel, RandomForest
from hscredit.core.rules.rule import Rule
from hscredit.report import ModelReport
from hscredit.report._model_context import _configuration_snapshot, _model_parameter_snapshot
from hscredit.report._model_layout import configuration_table


def test_configuration_retains_long_deep_and_object_parameters_without_training_data():
    config = {"层": {"层": {"层": {"层": {"层": {"learning_rate": 1e-8}}}}}}
    config["列表"] = list(range(40))
    config["数组"] = np.arange(40)
    config["字典"] = {str(i): i for i in range(300)}
    config["规则"] = Rule("余额 > 100", weight=0.25)
    config["空参数"] = None
    nested = RandomForest(n_estimators=7, max_depth=3, random_state=8)
    nested.estimators_ = ["不得导出的训练权重"]
    config["模型"] = nested
    config["循环"] = config
    snapshot = _configuration_snapshot(config)
    rows = configuration_table(snapshot).set_index("配置项")["配置值"]
    assert rows["层.层.层.层.层.learning_rate"] == 1e-8
    assert rows["列表[39]"] == rows["数组[39]"] == 39
    assert rows["字典.299"] == 299
    assert rows["规则.expr"] == "余额 > 100"
    assert rows["规则.weight"] == 0.25
    assert rows["规则.description"] == '""'
    assert rows["模型.n_estimators"] == 7
    assert rows["模型.max_depth"] == 3
    assert rows["空参数"] == "None"
    assert rows["循环"] == "<recursive>"
    assert "不得导出的训练权重" not in str(snapshot)
    assert "已截断" not in str(snapshot) and "未保留逐项明细" not in str(snapshot)
    nested.set_params(n_estimators=10)
    assert snapshot["模型"]["n_estimators"] == 7


@pytest.mark.parametrize("fit_supplied", [False, True])
def test_calibrated_pipeline_parameters_reach_excel_and_status_stays_in_result(tmp_path, fit_supplied):
    X = pd.DataFrame({"余额": np.arange(40), "期限": np.arange(40) % 3})
    y = (X["余额"] > 19).astype(int)
    base = LogisticRegression(C=0.123456789, calculate_stats=False, random_state=4).fit(X, y)
    calibrator = ProbabilityCalibrator(model=None if fit_supplied else base, calib_ratio=None)
    calibrator.fit(X, y, model=base if fit_supplied else None)
    combined = CalibratedModel(base, calibrator)
    model = Pipeline([("标准化", StandardScaler().fit(X)), ("校准模型", combined)])
    report = ModelReport(model, datasets={"验证集": (X, y)}, n_jobs=1)
    path = tmp_path / "完整参数.xlsx"
    result = report.to_excel(path, with_plots=False, mode="strict", return_result=True)
    assert result.complete and all(section.status != "失败" for section in result.sections.values())
    book = load_workbook(path)
    assert book.sheetnames[0] == "目录"
    assert "报告执行状态" not in book.sheetnames
    assert not any(cell.value == "报告执行状态" for row in book["目录"] for cell in row)
    values = {row[1].value: row[2].value for row in book["5-模型参数"] if len(row) > 2}
    assert values["base_model.C"] == pytest.approx(0.123456789)
    if not fit_supplied:
        assert values["calibrator.model.C"] == pytest.approx(0.123456789)
    else:
        assert values["calibrator.model_.C"] == pytest.approx(0.123456789)
    child = _model_parameter_snapshot(calibrator)
    if fit_supplied:
        assert child["model_"]["C"] == pytest.approx(0.123456789)
    assert not any(
        any("\u4e00" <= char <= "\u9fff" for char in key) for key in values if isinstance(key, str) and "." in key
    )
    assert "工作表：目录" in result.sections
    assert book["7-训练与调参过程"]["B2"].value == "七、训练与调参过程"
    assert book["7-训练与调参过程"]["B5"].value == "1、训练概览"
    book.close()


def test_native_effective_parameters_are_preserved_separately_from_public_defaults():
    model = RandomForest(n_estimators=7, max_depth=3)
    model.native_params_ = {"n_estimators": 7, "max_depth": 3, "class_weight": {0: 0.25, 1: 0.75}}
    snapshot = _model_parameter_snapshot(model)
    assert snapshot["n_estimators"] == 7
    assert snapshot["class_weight"] is None
    assert snapshot["native_params_"]["class_weight"] == {"0": 0.25, "1": 0.75}
