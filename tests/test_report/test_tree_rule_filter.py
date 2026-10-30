"""用原始规则命中核验训练集筛选、跨集固定规则与报告布局。"""

import operator

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from hscredit.report.mining import DecisionTreeAnalyzer, ManualTreeExtractor, MultiwayTreeExtractor


@pytest.fixture
def data():
    return pd.DataFrame(
        {
            "评分": np.arange(96),
            "多头": np.tile(np.arange(12), 8),
            "FPD": ([0] * 8 + [1] * 4) * 8,
            "MOB1": [0, 2, 3, 4, 7, 8] * 16,
            "金额": np.arange(96) * 1.13 + 100.03,
            "日期": pd.date_range("2026-01-01", periods=96),
            "品类": ["A", "B"] * 48,
        }
    )


@pytest.fixture(params=[ManualTreeExtractor, MultiwayTreeExtractor, DecisionTreeAnalyzer])
def tree_type(request):
    return request.param


def fitted(tree_type, data, rule_filter=None):
    return tree_type(
        target="FPD",
        features=["评分", "多头"],
        max_depth=3,
        min_samples_leaf=2,
        n_jobs=1,
        rule_filter=rule_filter,
    ).fit(data)


def test_default_does_not_filter(tree_type, data):
    tree = fitted(tree_type, data)
    pd.testing.assert_frame_equal(tree.report(), tree.report(apply_rule_filter=False))
    assert len(tree.report()) == len(tree.get_rules(leaf_only=False))


def test_legacy_multiway_has_disabled_filter(data):
    tree = fitted(MultiwayTreeExtractor, data)
    del tree.rule_filter
    assert tree.get_params()["rule_filter"] is None
    assert len(tree.report()) == len(tree.get_rules(leaf_only=False))


@pytest.mark.parametrize("metric,threshold", [("badrate", 1 / 3), ("lift", 1)])
@pytest.mark.parametrize("op", [">", ">=", "<", "<="])
def test_filter_uses_training_orders_without_changing_tree(tree_type, data, metric, threshold, op):
    config = dict(metric=metric, threshold=threshold, operator=op, leaf_only=False, min_samples=0.05, max_samples=80)
    tree = fitted(tree_type, data, config)
    expected = []
    compare = {">": operator.gt, ">=": operator.ge, "<": operator.lt, "<=": operator.le}[op]
    rules = tree.get_rules(leaf_only=False)
    for rule in rules:
        node = int(rule.name.rsplit("_", 1)[1].lstrip("N"))
        rows = data.loc[rule.predict(data)]
        value = rows.FPD.mean() if metric == "badrate" else rows.FPD.mean() / data.FPD.mean()
        if node != 0 and 0.05 * len(data) <= len(rows) <= 80 and compare(value, threshold):
            expected.append(node)
    validation = data.iloc[::2].assign(FPD=lambda frame: 1 - frame.FPD, 金额=10000)
    for table in tree.report({"验证": validation, "训练": data}, amount="金额", round_digits=False).values():
        assert table["节点编号"].tolist() == expected
    assert len(tree.get_rules(leaf_only=False)) == len(rules)
    assert len(tree.get_rule_table(apply_rule_filter=False)) == len(rules)
    if isinstance(tree, MultiwayTreeExtractor):
        assert len(tree.get_node_table()) == len(rules)
        assert tree.get_node_table()["样本总数"].notna().all()


def test_explicit_defaults_and_sample_bound_types(tree_type, data):
    tree = fitted(tree_type, data, {})
    full = tree.get_rule_table(apply_rule_filter=False, round_digits=False)
    expected = full.loc[full["是否叶子"].eq("是") & full["LIFT值"].ge(1) & full["节点编号"].ne(0), "节点编号"]
    assert tree.report()["节点编号"].tolist() == expected.tolist()
    tree.rule_filter = dict(threshold=0, min_samples=0, max_samples=1)
    assert tree.report().empty  # 整数 1 为 1 个样本，叶节点最小 2 个样本。
    tree.rule_filter = dict(threshold=0, min_samples=0.0, max_samples=1.0)
    assert len(tree.report()) == len(tree.get_rules())  # 浮点数 1.0 为全部样本。
    tree.rule_filter = dict(min_samples=50, max_samples=0.25)
    with pytest.raises(ValueError, match="不能大于"):
        tree.report()


@pytest.mark.parametrize(
    "config",
    [
        False,
        {"unknown": 1},
        {"operator": "=="},
        {"operator": []},
        {"metric": "AUC"},
        {"threshold": np.nan},
        {"threshold": True},
        {"leaf_only": 1},
        {"min_samples": -1},
        {"max_samples": 1.1},
        {"max_samples": True},
    ],
)
def test_invalid_filter_rejected_at_initialization(tree_type, config):
    with pytest.raises(ValueError, match="rule_filter"):
        tree_type(rule_filter=config)


@pytest.mark.parametrize("empty", [False, True])
def test_filtered_excel_uses_same_nodes_for_all_labels_and_bases(tree_type, data, tmp_path, empty):
    tree = fitted(tree_type, data, dict(threshold=100 if empty else 1))
    path = tmp_path / "筛选报告.xlsx"
    expected = tree.get_rule_table()["节点编号"].tolist()
    actual = tree.report(
        {"验证": data.iloc[::2], "训练": data},
        overdue="MOB1",
        dpds=[7, 0],
        amount="金额",
        methods=["quantile"],
        save=path,
    )
    for table in actual.values():
        assert table["节点编号"].tolist() == np.repeat(expected, 2).tolist()
    book = load_workbook(path)
    sheet = book.active
    headers = [
        (c.row, c.column)
        for row in sheet
        for c in row
        if c.value == "节点编号" and sheet.cell(c.row, c.column + 1).value == "指标含义"
    ]
    assert len(headers) == 6
    for i, (row, col) in enumerate(headers):
        nodes = np.repeat(expected, 1 if i < 2 else 2).tolist()
        assert [sheet.cell(row + j + 1, col).value for j in range(len(nodes))] == nodes
        assert not isinstance(sheet.cell(row + len(nodes) + 1, col).value, int)
    assert len(sheet._images) == 6  # 即使筛选为空也展示完整树。
    assert not any(c.data_type == "e" for ws in book for row in ws for c in row)
    book.close()


def test_group_merges_feature_order_width_and_statistics(tree_type, data, tmp_path):
    tree = fitted(tree_type, data)
    path = tmp_path / "版式.xlsx"
    datasets = {"后传名称A": data, "先字典排序B": data.iloc[::2]}
    tree.report(
        datasets,
        overdue="MOB1",
        dpds=[7, 0],
        del_grey=True,
        date="日期",
        group="品类",
        amount="金额",
        methods=["uniform", "quantile", "tree"],
        save=path,
    )
    book = load_workbook(path)
    feature = book.worksheets[1]
    assert all(d.width <= 25 for d in feature.column_dimensions.values())
    left, top, right, bottom = range_boundaries(feature.auto_filter.ref)
    assert [feature.cell(top, c).value for c in range(left, left + 4)] == ["数据集", "指标名称", "统计口径", "分箱方法"]
    rows = pd.DataFrame(
        list(feature.iter_rows(min_row=top + 1, max_row=bottom, min_col=left, max_col=left + 3, values_only=True))
    ).ffill()
    assert rows[0].drop_duplicates().tolist() == list(datasets)
    for _, subset in rows.groupby(0, sort=False):
        assert subset[1].drop_duplicates().tolist() == sorted(subset[1].unique())
    for _, subset in rows.groupby([0, 1], sort=False):
        assert subset[2].drop_duplicates().tolist() == ["订单", "金额"]
        assert list(subset[[2, 3]].itertuples(index=False, name=None)) == sorted(
            subset[[2, 3]].itertuples(index=False, name=None)
        )
    assert any(r.min_col == left + 1 and r.max_row > r.min_row for r in feature.merged_cells.ranges)
    main = book.active
    features_cell = next(c for row in main for c in row if c.value == "features")
    assert main.cell(features_cell.row, features_cell.column + 1).alignment.horizontal == "left"
    titles = {
        c.value: c.row
        for row in main
        for c in row
        if c.value in ["3、数据样本统计", "4、样本分布情况", "二、规则挖掘摘要"]
    }
    first, last = titles["4、样本分布情况"], titles["二、规则挖掘摘要"]
    dataset_headers = [c for row in main.iter_rows(min_row=first, max_row=last) for c in row if c.value == "数据集"]
    assert len(dataset_headers) == 2
    for header in dataset_headers:
        dataset_merges = [
            r
            for r in main.merged_cells.ranges
            if r.min_col == r.max_col == header.column and header.row < r.min_row < last
        ]
        assert [main.cell(r.min_row, r.min_col).value for r in sorted(dataset_merges, key=lambda r: r.min_row)] == list(
            datasets
        )
        # 月份只在各自数据集内合并；不同品类属于同一月份。
        for merged in main.merged_cells.ranges:
            if merged.min_col == merged.max_col == header.column + 1 and header.row < merged.min_row < last:
                assert any(
                    parent.min_row <= merged.min_row <= merged.max_row <= parent.max_row for parent in dataset_merges
                )
    stat_row = titles["3、数据样本统计"] + 5
    order_metrics = [c.value for c in main[stat_row] if c.value is not None]
    expected = ["统计详情", "样本总数", "好样本数", "坏样本数", "坏样本率", "AUC", "KS"]
    assert order_metrics == expected * 2
    for c in main[stat_row]:
        if c.value in {"坏样本率", "KS", "AUC"}:
            assert main.cell(stat_row + 2, c.column).number_format == ("0.0000" if c.value == "AUC" else "0.00%")
    book.close()


@pytest.mark.parametrize("group", ["品类", ["品类", "层次"]])
def test_group_only_report_matches_raw_aggregation(tree_type, data, tmp_path, group):
    data = data.assign(层次=np.arange(len(data)) % 3)
    tree = fitted(tree_type, data)
    path = tmp_path / "按业务分组.xlsx"
    tree.report({"样本": data}, group=group, save=path, methods=["quantile"])
    book = load_workbook(path)
    sheet = book.active
    title = next(c.row for row in sheet for c in row if c.value == "4、样本分布情况")
    columns = [group] if isinstance(group, str) else group
    expected = data.groupby(columns, sort=True).agg(样本总数=("FPD", "size"), 坏样本数=("FPD", "sum"))
    count_col = 3 + len(columns)
    first_row = title + 6
    for offset, row in enumerate(expected.itertuples()):
        assert sheet.cell(first_row + offset, count_col).value == row.样本总数
        assert sheet.cell(first_row + offset, count_col + 2).value == row.坏样本数
        assert sheet.cell(first_row + offset, count_col + 3).value == pytest.approx(row.坏样本数 / row.样本总数)
    book.close()
