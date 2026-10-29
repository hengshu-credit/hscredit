"""样本分布的分组值逐行显示，只合并数据集列。"""

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression

from hscredit.report import ModelReport


@pytest.mark.parametrize("multi_label", [False, True])
def test_distribution_groups_never_merge_across_dataset_boundaries(tmp_path, multi_label):
    frame = pd.DataFrame(
        {
            "甲": np.arange(12),
            "目标": [0, 1] * 6,
            "MOB1": [0, 5] * 6,
            "日期": pd.to_datetime(
                np.repeat(["2025-11-01", "2025-12-01", "2026-01-01", "2026-01-15", "2026-02-01", "2026-02-15"], 2)
            ),
            "类别": ["同一类别"] * 12,
        }
    )
    model = LogisticRegression().fit(frame[["甲"]], frame["目标"])
    report = ModelReport(
        model,
        datasets={"训练集": frame.iloc[:6], "选参验证集": frame.iloc[6:10], "独立测试集": frame.iloc[10:]},
        features=["甲"],
        n_jobs=1,
        **({"overdue": "MOB1", "dpds": [0, 3]} if multi_label else {"target": "目标"}),
    )
    output = tmp_path / "样本分布.xlsx"
    report.to_excel(output, with_plots=False, date="日期", group="类别", mode="strict")
    book = load_workbook(output)
    sheet = book["1-基本信息"]
    headers = [cell for row in sheet for cell in row if cell.value == "数据分组"]
    assert len(headers) == 2
    expected = [
        (["2025-11", "2025-12", "2026-01", "2026-01", "2026-02", "2026-02"], [2] * 6),
        (["同一类别"] * 3, [6, 4, 2]),
    ]
    for header, (groups, counts) in zip(headers, expected):
        for offset, (group, count) in enumerate(zip(groups, counts), 1):
            cell = sheet.cell(header.row + offset, header.column)
            assert cell.value == group
            assert not any(cell.coordinate in area for area in sheet.merged_cells.ranges)
            assert sheet.cell(cell.row, cell.column + 1).value == count
    month_header = headers[0]
    dataset_merges = [
        area
        for area in sheet.merged_cells.ranges
        if area.min_col == area.max_col == month_header.column - 1
        and month_header.row < area.min_row <= month_header.row + 6
    ]
    assert {(area.min_row - month_header.row, area.max_row - month_header.row) for area in dataset_merges} == {
        (1, 3),
        (4, 5),
    }
    book.close()
