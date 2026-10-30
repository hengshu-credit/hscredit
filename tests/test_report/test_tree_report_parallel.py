"""树报告并行计算、预测复用与结果不变契约。"""

import copy

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from hscredit.report.mining import DecisionTreeAnalyzer, ManualTreeExtractor, MultiwayTreeExtractor


@pytest.fixture(params=[ManualTreeExtractor, MultiwayTreeExtractor, DecisionTreeAnalyzer])
def model_and_data(request):
    frame = pd.DataFrame(
        {
            "评分": np.arange(72),
            "FPD": np.repeat([0, 1, 0], 24),
            "MOB1": np.tile([0, 2, 4, 8, 0, 10], 12),
            "金额": np.arange(72) + 100.25,
            "日期": pd.date_range("2026-01-01", periods=72),
            "品类": ["A", "B"] * 36,
        },
        index=np.arange(72) // 2,
    )
    tree = request.param(target="FPD", features=["评分"], max_depth=2, min_samples_leaf=2, n_jobs=1).fit(frame)
    return tree, frame


@pytest.mark.parametrize("backend", ["threading", "loky"])
def test_parallel_excel_matches_serial_with_duplicate_indexes(model_and_data, backend, tmp_path):
    tree, frame = model_and_data
    frames = {"验证": frame.iloc[::2], "训练": frame}
    original = frame.copy(deep=True)
    options = dict(
        overdue="MOB1",
        dpds=[7, 3, 0],
        del_grey=True,
        amount="金额",
        date="日期",
        group="品类",
        methods=["quantile", "uniform"],
        round_digits=False,
    )
    serial = tree.report(frames, save=tmp_path / "串行.xlsx", **options)
    tree.n_jobs, tree.parallel_backend = 2, backend
    tree.parallel_config = {"result_retention": "iterator", "max_inflight_tasks": 2}
    parallel = tree.report(frames, save=tmp_path / "并行.xlsx", **options)
    for name in frames:
        pd.testing.assert_frame_equal(serial[name], parallel[name], check_exact=True)
    assert tree.n_jobs == 2 and tree.parallel_backend == backend
    pd.testing.assert_frame_equal(frame, original)
    before, after = load_workbook(tmp_path / "串行.xlsx"), load_workbook(tmp_path / "并行.xlsx")
    for old, new in zip(before, after):
        assert old.max_row == new.max_row and old.max_column == new.max_column
        assert set(map(str, old.merged_cells.ranges)) == set(map(str, new.merged_cells.ranges))
        assert len(old._images) == len(new._images)
        for row in old:
            for cell in row:
                if cell.column == 3 and old.cell(cell.row, 2).value in {
                    "n_jobs",
                    "parallel_backend",
                    "parallel_config",
                }:
                    continue
                actual = new[cell.coordinate]
                assert cell.value == actual.value, (old.title, cell.coordinate)
                assert cell.number_format == actual.number_format
                assert copy.copy(cell.alignment) == copy.copy(actual.alignment)
    before.close()
    after.close()


def test_prediction_reused_across_labels_weights_and_groups_but_not_reports(model_and_data, tmp_path, monkeypatch):
    import hscredit.report.mining._tree_report as module

    tree, frame = model_and_data
    calls = []
    original = module._tree_score

    def count(model, data, class_index=1):
        calls.append((len(data), class_index))
        return original(model, data, class_index)

    monkeypatch.setattr(module, "_tree_score", count)
    options = dict(
        overdue="MOB1", dpds=[7, 3, 0], del_grey=True, amount="金额", date="日期", group="品类", methods=["quantile"]
    )
    tree.report({"样本": frame}, save=tmp_path / "首次.xlsx", **options)
    assert calls == [(len(frame), 1)]
    changed = frame.assign(评分=frame["评分"].to_numpy()[::-1])
    tree.report({"样本": changed}, save=tmp_path / "再次.xlsx", **options)
    assert calls == [(len(frame), 1), (len(frame), 1)]
