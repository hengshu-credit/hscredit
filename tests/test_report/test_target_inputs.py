"""统一标签输入、作用数据、灰客户及跨入口统计的语义契约。"""

import pickle
from functools import partial
from inspect import signature

import matplotlib
import numpy as np
import pandas as pd
import pytest
from pandas.testing import assert_frame_equal

from hscredit.core.eda.target import bad_rate_overall, bad_rate_by_dimension, bad_rate_trend, bad_rate_by_bins
from hscredit.core.rules import Rule
from hscredit.core.viz import bin_overdues_plot, distribution_plot, bad_rate_trend_plot
from hscredit.report import (
    feature_bin_stats,
    feature_bin_stats_2d,
    feature_binning_summary,
    feature_group_binning_summary,
)
from hscredit.report.feature_analyzer import (
    benchmark_binning_methods,
    feature_efficiency_analysis,
    auto_feature_analysis,
)
from hscredit.report.model_report import ModelReport
from hscredit.report.overdue_predictor import OverduePredictor
from hscredit.report.rule_analysis import ruleset_analysis, rule_swap_analysis
from hscredit.report.rule_strategy import rule_group_compare, swap_out_report
from hscredit.utils.targets import bind_targets

matplotlib.use("Agg")


@pytest.fixture
def frame():
    return pd.DataFrame(
        {
            "score": np.arange(12, dtype=float),
            "feature": np.tile([0, 1, 2], 4),
            "target": [0, 1, 2, -1, 0, 1, 0, 1, np.nan, 2, 0, 1],
            "MOB1": [0, 8, 2, 3, 0, 8, 0, 8, 3, 1, 0, 8],
            "group": np.repeat(["A", "B", "C"], 4),
            "date": pd.to_datetime(np.repeat(["2026-01-01", "2026-02-01", "2026-03-01"], 4)),
            "amount": np.arange(1, 13, dtype=float) * 10,
        },
        index=np.arange(12)[::-1] + 20,
    )


def definitions(frame):
    return {"自定义_7+": frame["target"], "另一个@7": lambda data: data["MOB1"].gt(0)}


@pytest.mark.parametrize("kind", ["column", "series", "array", "list", "lambda", "partial"])
@pytest.mark.parametrize("grey", [False, True])
def test_all_single_input_forms_share_grey_semantics(frame, kind, grey):
    def derive(data, column):
        return data[column]

    inputs = {
        "column": "target",
        "series": frame.target.sort_index(),
        "array": frame.target.to_numpy(),
        "list": frame.target.tolist(),
        "lambda": lambda data: data.target,
        "partial": partial(derive, column="target"),
    }
    original = frame.copy(deep=True)
    table = feature_bin_stats(
        frame, "score", target=inputs[kind], del_grey=grey, rules=[3.5, 7.5], margins=True, n_jobs=1
    )
    valid = frame.target.isin([0, 1]) if grey else pd.Series(True, index=frame.index)
    assert table.iloc[-1]["样本总数"] == valid.sum()
    assert table.iloc[-1]["坏样本数"] == (valid & frame.target.eq(1)).sum()
    assert_frame_equal(frame, original)


@pytest.mark.parametrize("operator", [">", ">=", "<", "<="])
def test_explicit_grey_is_independent_of_overdue_operator(frame, operator):
    table = bad_rate_overall(frame, target="target", del_grey=True, overdue_operator=operator)
    assert table["样本总数"] == 8
    assert table["坏样本数"] == 4


def test_invalid_shape_index_and_callable_errors_are_explicit(frame):
    for target in [np.ones((12, 1)), [0, 1], pd.Series([0, 1]), lambda data: 1, lambda: [0] * 12]:
        with pytest.raises((ValueError, TypeError)):
            bind_targets(frame, target=target)
    calls = []

    def broken(data):
        calls.append(1)
        raise TypeError("函数内部错误")

    with pytest.raises(ValueError, match="计算失败.*函数内部错误"):
        bind_targets(frame, target={"风险": broken}, dataset_name="测试集")
    assert calls == [1]


def test_duplicate_index_requires_exact_order_and_uses_positional_slicing(frame):
    data = frame.set_axis([0, 0, 1, 1, 2, 2, 3, 3, 4, 4, 5, 5])
    result = feature_bin_stats(data, "score", target=lambda df: df.target, del_grey=True, rules=[3.5], n_jobs=1)
    assert result["样本总数"].sum() == 8
    with pytest.raises(ValueError, match="重复索引"):
        bind_targets(data, target=data.target.iloc[::-1])


@pytest.mark.parametrize("entry", ["bin", "2d", "summary", "group", "rule", "ruleset", "rule_group"])
def test_callables_run_once_on_full_acting_data(frame, entry):
    seen = []

    def label(data):
        assert_frame_equal(data, frame)
        seen.append(data.index.tolist())
        return data.score.gt(data.score.median())

    kwargs = dict(target={"分位标签": label, "灰标签": frame.target}, del_grey=True, n_jobs=1)
    if entry == "bin":
        feature_bin_stats(frame, ["score", "feature"], method="quantile", **kwargs)
    elif entry == "2d":
        feature_bin_stats_2d(frame, ["score", "feature"], max_n_bins=2, **kwargs)
    elif entry == "summary":
        feature_binning_summary(frame, ["score", "feature"], methods=["quantile", "uniform"], **kwargs)
    elif entry == "group":
        feature_group_binning_summary(frame, "score", methods="quantile", group="group", **kwargs)
    elif entry == "rule":
        Rule("score >= 4").report(frame, **kwargs)
    elif entry == "ruleset":
        ruleset_analysis(frame, [Rule("score >= 8"), Rule("score >= 4")], **kwargs)
    else:
        rule_group_compare(frame, Rule("score >= 4"), group="group", **kwargs)
    assert seen == [frame.index.tolist()]


@pytest.mark.parametrize("amount", [None, "amount"])
@pytest.mark.parametrize("grey", [False, True])
def test_multi_labels_keep_names_counts_amounts_and_shared_bins(frame, amount, grey):
    targets = definitions(frame)
    wide = feature_bin_stats(
        frame, "score", target=targets, del_grey=grey, amount=amount, rules=[3.5, 7.5], margins=True, n_jobs=1
    )
    assert list(wide.columns.get_level_values(0).unique()) == ["分箱详情", *targets]
    for name, raw in targets.items():
        raw = raw(frame) if callable(raw) else raw
        single = feature_bin_stats(
            frame, "score", target=raw, del_grey=grey, amount=amount, rules=[3.5, 7.5], margins=True, n_jobs=1
        )
        total = wide.iloc[-1]
        count_group = name if grey else "分箱详情"
        assert total[(count_group, "样本总数")] == single.iloc[-1]["样本总数"]
        assert total[(name, "坏样本数")] == single.iloc[-1]["坏样本数"]
    _, summary = feature_binning_summary(frame, "score", target=targets, del_grey=grey, methods="quantile", n_jobs=1)
    assert set(summary.columns.get_level_values(1)) >= set(targets)


@pytest.mark.parametrize(
    "fn,args", [(bad_rate_by_dimension, ("group",)), (bad_rate_trend, ("date",)), (bad_rate_by_bins, ("score",))]
)
def test_eda_multi_labels_exclude_grey_independently(frame, fn, args):
    results = fn(frame, *args, target=definitions(frame), del_grey=True)
    assert results["自定义_7+"]["样本数"].sum() == 8
    assert results["自定义_7+"]["坏样本数"].sum() == 4
    assert results["另一个@7"]["样本数"].sum() == 12


class ProbabilityModel:
    def predict_proba(self, data):
        probability = (data.score.to_numpy() + 1) / 20
        return np.column_stack([1 - probability, probability])


def test_model_multiple_datasets_cache_serialization_and_feature_views(frame):
    seen = []

    def label(data):
        seen.append(len(data))
        return data.target

    report = ModelReport(
        ProbabilityModel(),
        datasets={"训练": frame, "测试": frame.iloc[:8]},
        features=["score"],
        target={"自定义_7+": label, "另一个@7": lambda df: df.MOB1.gt(0)},
        del_grey=True,
        n_jobs=1,
    )
    assert seen == [12, 8]
    assert report.summary().loc["自定义_7+", ("样本数", "训练")] == 8
    for table in [
        report.get_bin_table("训练", label="自定义_7+"),
        report.get_feature_bin_table("score", "训练", label="自定义_7+"),
    ]:
        assert table.iloc[-1]["样本总数"] == 8
    assert seen == [12, 8]
    restored = pickle.loads(pickle.dumps(report))
    assert_frame_equal(restored.summary(), report.summary())
    restored.add_dataset("新增", "新增", frame.iloc[:4])
    assert restored.summary().loc["自定义_7+", ("样本数", "新增")] == 2
    with pytest.raises(ValueError, match="不存在标签"):
        report.get_metrics("未定义")


def test_predictor_and_callable_reference_datasets(frame):
    targets = definitions(frame)
    predictor = OverduePredictor("score", target=targets, del_grey=True, rules=[3.5, 7.5], n_jobs=1).fit(frame)
    assert predictor.target_names_ == list(targets)
    assert predictor.bin_table_.iloc[-1][("自定义_7+", "样本总数")] == 8
    seen = []

    def label(data):
        seen.append(len(data))
        return data.target

    result = rule_swap_analysis(
        frame,
        score={"分数一": "score", "分数二": "feature"},
        reference_data=frame.iloc[:8],
        target={"自定义": label},
        del_grey=True,
        rules_out=[Rule("score > 7")],
        n_jobs=1,
    )
    assert sorted(seen) == [8, 12]
    assert result["swap_pipeline"].iloc[0]["样本总数"] == 8


def test_plots_use_same_labels_and_counts(frame):
    import matplotlib.pyplot as plt

    targets = definitions(frame)
    table = distribution_plot(frame, date="date", target=targets, del_grey=True, result=True)
    np.testing.assert_allclose(table["自定义_7+_坏样本率"], [0.5, 0.5, 0.5])
    fig = bad_rate_trend_plot(frame, "date", target=targets, del_grey=True)
    np.testing.assert_allclose(fig.axes[0].lines[0].get_ydata(), [0.5, 0.5, 0.5])
    bins = bin_overdues_plot(frame, feature="score", target=targets, del_grey=True, n_jobs=1)
    assert not any("Error:" in text.get_text() for axis in bins.axes for text in axis.texts)
    plt.close("all")


def test_auto_report_uses_full_data_once_and_preserves_custom_names(frame, tmp_path):
    from hscredit.excel import ExcelWriter

    seen = []

    def target(data):
        assert_frame_equal(data, frame)
        seen.append(1)
        return data.target

    writer = ExcelWriter()
    result = auto_feature_analysis(
        frame,
        features=["score", "feature"],
        target={"自定义_7+": target, "另一个@7": frame.MOB1.gt(0)},
        del_grey=True,
        date="date",
        pictures=[],
        show_progress=False,
        excel_writer=writer,
        output_dir=str(tmp_path),
        n_jobs=1,
        return_result=True,
    )
    assert result.complete
    assert seen == [1]
    writer.workbook.close()


@pytest.mark.parametrize(
    "call",
    [
        lambda d: feature_bin_stats(d, "score", target="target", overdue="MOB1", dpds=[3]),
        lambda d: Rule("score>1").report(d, target="target", overdue="MOB1", dpds=[3]),
        lambda d: ModelReport(
            ProbabilityModel(), datasets={"样本": d}, features=["score"], target="target", overdue="MOB1", dpds=[3]
        ),
        lambda d: bad_rate_overall(d, target="target", overdue="MOB1", dpds=[3]),
    ],
)
def test_two_modes_are_mutually_exclusive(frame, call):
    with pytest.raises(ValueError, match="不能同时指定"):
        call(frame)


def test_old_configuration_and_parameter_names_are_not_accepted(frame):
    with pytest.raises(TypeError):
        bind_targets(frame, target={"overdue": "MOB1", "dpds": [7, 30]})
    for function in [
        bad_rate_overall,
        bad_rate_by_dimension,
        bad_rate_trend,
        bad_rate_by_bins,
        feature_efficiency_analysis,
    ]:
        assert "target" in signature(function).parameters
        assert "target_col" not in signature(function).parameters
        assert "dpd" not in signature(function).parameters
    with pytest.raises(TypeError, match="dpd.*dpds"):
        feature_bin_stats(frame, "score", target="target", dpd=3)


def test_parallel_features_reuse_parent_evaluated_targets(frame):
    calls = []

    def label(data):
        calls.append(1)
        return data.target

    kwargs = dict(
        feature=["score", "feature"],
        target={"灰标签": label, "正常": frame.MOB1.gt(0)},
        del_grey=True,
        method="quantile",
        max_n_bins=2,
    )
    serial = feature_bin_stats(frame, n_jobs=1, **kwargs)
    parallel = feature_bin_stats(
        frame, n_jobs=2, parallel_backend="loky", parallel_config={"adaptive": False}, **kwargs
    )
    assert_frame_equal(serial, parallel)
    assert calls == [1, 1]


def test_tree_reports_parse_target_once_and_keep_explicit_names(frame):
    from hscredit.report.mining import DecisionTreeAnalyzer

    data = frame.assign(fit_target=frame.MOB1.gt(0).astype(int))
    tree = DecisionTreeAnalyzer(target="fit_target", features=["score"], tree_params={"max_depth": 2}, n_jobs=1).fit(
        data
    )
    calls = []

    def label(current):
        assert_frame_equal(current, data)
        calls.append(1)
        return current.target

    result = tree.report(data, target={"自定义_7+": label, "另一个@7": data.fit_target}, del_grey=True)
    assert "自定义_7+" in result.columns.get_level_values(0)
    assert calls == [1]


def test_skill_resolves_label_arrays_and_callable_object_references(frame, tmp_path):
    from hscredit.skills_runtime import execute_skill
    from openpyxl import load_workbook

    calls = []

    def label(data):
        assert_frame_equal(data, frame)
        calls.append(1)
        return data.MOB1.gt(0)

    result = execute_skill(
        "hsbin",
        {
            "version": "1",
            "operation": "feature_bin_stats",
            "inputs": {name: {"kind": "object_ref", "ref": name} for name in ("data", "label", "factory")},
            "parameters": {
                "feature": "score",
                "target": {"灰标签": {"input": "label"}, "正常": {"input": "factory"}},
                "del_grey": True,
                "method": "quantile",
                "n_jobs": 1,
            },
            "output": {"directory": str(tmp_path), "name": "labels", "overwrite": False},
            "environment": {"mode": "current", "install_missing": False},
        },
        objects={"data": frame, "label": frame.target, "factory": label},
    )
    assert result["status"] == "success"
    assert result["summary"]["target_mode"] == "target"
    assert result["summary"]["target_names"] == ["灰标签", "正常"]
    assert calls == [1]
    workbook = load_workbook(next(a["path"] for a in result["artifacts"] if a["type"] == "excel"))
    assert "灰标签" in [cell.value for row in workbook.active for cell in row]
    workbook.close()


def test_explicit_labels_match_overdue_count_and_amount_on_real_loans():
    from pathlib import Path

    data = pd.read_excel(Path(__file__).resolve().parents[2] / "examples" / "hscredit_yyp.xlsx")
    from hscredit.utils.overdue import make_overdue_target

    for operator in (">", ">=", "<", "<="):
        targets = {
            f"MOB1{operator}{threshold}": make_overdue_target(data.MOB1, threshold, True, operator)
            for threshold in (7, 3, 0)
        }
        for amount in (None, "放款金额"):
            manual = feature_bin_stats(
                data,
                "衡枢鉴真分老客版",
                target=targets,
                del_grey=True,
                rules=[0.3, 0.6],
                amount=amount,
                margins=True,
                n_jobs=1,
            )
            overdue = feature_bin_stats(
                data,
                "衡枢鉴真分老客版",
                overdue="MOB1",
                dpds=[7, 3, 0],
                overdue_operator=operator,
                del_grey=True,
                rules=[0.3, 0.6],
                amount=amount,
                margins=True,
                n_jobs=1,
            )
            for threshold in (7, 3, 0):
                label = f"MOB1{operator}{threshold}"
                legacy = f"MOB1_{threshold}+" if operator == ">" else label
                count_group = legacy if operator in (">", ">=") else "分箱详情"
                assert manual.iloc[-1][(label, "样本总数")] == overdue.iloc[-1][(count_group, "样本总数")]
                assert manual.iloc[-1][(label, "坏样本数")] == overdue.iloc[-1][(legacy, "坏样本数")]


def test_benchmark_efficiency_and_strategy_preserve_target_definitions(frame, tmp_path):
    import matplotlib.pyplot as plt
    from openpyxl import load_workbook

    targets = definitions(frame)
    benchmark = benchmark_binning_methods(
        frame, "score", target=targets, del_grey=True, methods=["quantile"], n_jobs=1
    )
    assert set(benchmark.columns.get_level_values(0)) == {"分箱详情", *targets}
    assert all(pd.isna(benchmark.iloc[0][(name, "错误信息")]) for name in targets)
    results = feature_efficiency_analysis(
        frame,
        "score",
        target=targets,
        del_grey=True,
        rules=[3.5, 7.5],
        method="quantile",
        margins=True,
        n_jobs=1,
    )
    assert list(results["targets"]) == list(targets)
    assert results["targets"]["自定义_7+"]["manual_table"].iloc[-1]["样本总数"] == 8
    assert results["targets"]["另一个@7"]["manual_table"].iloc[-1]["样本总数"] == 12
    path = tmp_path / "自定义标签策略.xlsx"
    writer = swap_out_report(
        frame,
        Rule("score > 4"),
        features=["score"],
        target=targets,
        del_grey=True,
        methods="quantile",
        group="group",
        amount="amount",
        n_jobs=1,
        save=str(path),
    )
    workbook = load_workbook(path)
    values = [cell.value for sheet in workbook for row in sheet for cell in row]
    assert "自定义_7+" in values and "另一个@7" in values
    assert "自定义 7+" not in values and "另一个 7+" not in values
    workbook.close()
    writer.workbook.close()
    plt.close("all")


def test_model_report_wrappers_expose_target_and_grey_contract(frame):
    from hscredit.core.models import LogisticRegression

    data = frame.assign(binary=frame.MOB1.gt(0).astype(int))
    model = LogisticRegression(max_iter=100).fit(data[["score"]], data.binary)
    assert "target" in signature(model.report).parameters
    report = model.report(
        datasets={"样本": data},
        target=definitions(data),
        del_grey=True,
        features=["score"],
        verbose=False,
        with_plots=False,
        n_jobs=1,
    )
    assert report.summary().loc["自定义_7+", ("样本数", "样本")] == 8


def test_swap_unobserved_labels_have_same_semantics_for_arrays_and_callable(frame):
    from hscredit.report.rule_analysis import _resolve_target_series

    expected = _resolve_target_series(frame, "target", None, None)["target"]
    for value in (frame.target.to_numpy(), frame.target.tolist(), lambda df: df.target):
        actual = _resolve_target_series(frame, {"target": value}, None, None)["target"]
        np.testing.assert_equal(actual.to_numpy(), expected.to_numpy())
    assert np.isnan(expected.iloc[8])
    assert expected.iloc[2] == 0


def test_report_callable_error_identifies_dataset(frame):
    with pytest.raises(ValueError, match="测试集.*风险.*计算失败"):
        ModelReport(
            ProbabilityModel(),
            datasets={"测试集": frame},
            features=["score"],
            target={"风险": lambda df: df["缺失字段"]},
            n_jobs=1,
        )


@pytest.mark.parametrize("whole_definition", [False, True])
def test_visualization_skill_accepts_target_object_reference(frame, tmp_path, whole_definition):
    from hscredit.skills_runtime import execute_skill

    inputs = {"data": {"kind": "object_ref", "ref": "data"}}
    objects = {"data": frame}
    kwargs = {"feature": "score", "del_grey": True, "n_jobs": 1}
    if whole_definition:
        inputs["target"] = {"kind": "object_ref", "ref": "definition"}
        objects["definition"] = definitions(frame)
    else:
        inputs["label_values"] = {"kind": "object_ref", "ref": "labels"}
        objects["labels"] = frame.target
        kwargs["target"] = {"自定义_7+": {"input": "label_values"}}
    result = execute_skill(
        "hsviz",
        {
            "version": "1",
            "operation": "plot",
            "inputs": inputs,
            "parameters": {"plot": "bin_overdues_plot", "kwargs": kwargs},
            "output": {"directory": str(tmp_path), "name": "target_plot", "format": "png"},
            "environment": {"mode": "current", "install_missing": False},
        },
        objects=objects,
    )
    assert result["status"] == "success"
    assert "自定义_7+" in result["summary"]["target_names"]
    assert any(item["type"] == "image" for item in result["artifacts"])


def test_primary_model_summary_months_and_plots_share_filtered_samples(frame, tmp_path, monkeypatch):
    import hscredit.core.viz as viz

    report = ModelReport(
        ProbabilityModel(),
        datasets={"样本": frame},
        features=["score"],
        target=definitions(frame),
        del_grey=True,
        n_jobs=1,
    )
    summary = report._get_features_summary()
    assert summary["样本数"].tolist() == [8]
    monthly = report._get_monthly_metrics("date")
    assert monthly["样本数"].sum() == 8
    assert monthly["KS"].notna().all()
    calls = []
    original = viz.ks_plot

    def capture(score, target, **kwargs):
        calls.append((len(score), len(target), pd.isna(target).sum()))
        return original(score, target, **kwargs)

    monkeypatch.setattr(viz, "ks_plot", capture)
    paths, _ = report._export_plots(tmp_path, max_n_bins=2, show_lift=False)
    assert paths["model_样本"]
    assert calls and all(count == target_count == 8 and missing == 0 for count, target_count, missing in calls)


def test_rule_skill_evaluates_labels_once_for_multiple_rules(frame, tmp_path):
    from hscredit.skills_runtime import execute_skill

    calls = []

    def label(data):
        calls.append(len(data))
        return data.target

    result = execute_skill(
        "hsrule",
        {
            "version": "1",
            "operation": "rule_report",
            "inputs": {
                "data": {"kind": "object_ref", "ref": "frame"},
                "target": {"kind": "object_ref", "ref": "target"},
            },
            "parameters": {"rules": ["score > 3", "score > 7"], "del_grey": True, "n_jobs": 1},
            "output": {"directory": str(tmp_path), "name": "rule_targets"},
            "environment": {"mode": "current", "install_missing": False},
        },
        objects={"frame": frame, "target": {"灰标签": label}},
    )
    assert result["status"] == "success"
    assert calls == [len(frame)]


def test_custom_label_name_does_not_overwrite_model_feature(frame):
    summaries = []
    for name in ("风险", "score"):
        report = ModelReport(
            ProbabilityModel(),
            datasets={"样本": frame},
            features=["score"],
            target={name: lambda df: df.target},
            del_grey=True,
            n_jobs=1,
        )
        summaries.append(report._get_features_summary())
    assert_frame_equal(*summaries)
