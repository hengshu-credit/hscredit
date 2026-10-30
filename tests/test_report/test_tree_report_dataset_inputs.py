"""按 auto_model_report 的 dict/list DataFrame 约定验证单标签树报告。"""

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from sklearn.tree import DecisionTreeClassifier

from hscredit import MultiwayTreeExtractor
from hscredit.report import auto_model_report
from hscredit.report.mining import DecisionTreeAnalyzer, ManualTreeExtractor


@pytest.fixture
def samples():
    return pd.DataFrame(
        {
            "评分": np.arange(96),
            "多头": np.tile(np.arange(12), 8),
            "FPD": np.tile([0, 0, 1, 1], 24),
            "MOB1": np.tile([0, 2, 3, 4, 7, 8], 16),
            "金额": np.arange(96) + 101.25,
            "日期": pd.date_range("2026-01-01", periods=96),
        },
        index=np.arange(96)[::-1] + 100,
    )


@pytest.fixture(params=[MultiwayTreeExtractor, ManualTreeExtractor, DecisionTreeAnalyzer])
def tree(request, samples):
    if request.param is MultiwayTreeExtractor:
        result = request.param(target="FPD", features=["评分", "多头"], max_depth=2, n_jobs=1).fit(samples)
        result.manual_split(feature="评分", thresholds=[30, 60])
    else:
        result = request.param(target="FPD", n_jobs=1).fit(samples, features=["评分", "多头"], max_depth=2)
        if request.param is ManualTreeExtractor:
            result.manual_split(samples, feature="评分", threshold=45)
    return result


def _datasets(data, count, kind):
    frames = [data.iloc[i::count].copy() for i in range(count)]
    names = ["建模集", "独立验证", "测试集", "OOT 1", "OOT 2"][:count]
    return dict(zip(names, frames)) if kind == "dict" else frames


@pytest.mark.parametrize("kind", ["dict", "list"])
@pytest.mark.parametrize("count", [1, 2, 5])
@pytest.mark.parametrize("label_mode", ["target", "overdue"])
def test_single_label_frames_match_auto_model_report_naming_and_raw_rows(tree, samples, kind, count, label_mode):
    datasets = _datasets(samples, count, kind)
    options = {"target": "FPD"} if label_mode == "target" else {"overdue": ["MOB1"], "dpds": [7]}
    baseline = DecisionTreeClassifier(max_depth=2, random_state=42).fit(samples[["评分", "多头"]], samples.FPD)
    reference = auto_model_report(
        baseline, datasets=datasets, features=["评分", "多头"], verbose=False, with_plots=False, n_jobs=1, **options
    )
    names = list(reference._datasets)
    frames = list(datasets.values()) if kind == "dict" else datasets
    originals = [frame.copy(deep=True) for frame in frames]
    assert names == (list(datasets) if kind == "dict" else [f"数据集{i+1}" for i in range(count)])
    result = tree.report(datasets=datasets, round_digits=False, **options)
    convenient = tree.get_rule_table(datasets=datasets, round_digits=False, **options)
    assert isinstance(result, dict if kind == "dict" else list)
    tables = list(result.values()) if kind == "dict" else result
    other = list(convenient.values()) if kind == "dict" else convenient
    assert len(tables) == count
    if kind == "dict":
        assert list(result) == names
    for frame, original, table, same in zip(frames, originals, tables, other):
        pd.testing.assert_frame_equal(frame, original)
        pd.testing.assert_frame_equal(table, same)
        labels = frame.FPD.eq(1) if label_mode == "target" else frame.MOB1.gt(7)
        for row, rule in zip(table.itertuples(index=False), tree.get_rules(leaf_only=False)):
            mask = rule.predict(frame)
            assert row.样本总数 == mask.sum()
            assert row.坏样本数 == (labels & mask).sum()
            assert row.坏样本率 == pytest.approx(labels[mask].mean() if mask.any() else 0)
        assert table.loc[table.是否叶子.eq("是"), "样本总数"].sum() == len(frame)


@pytest.mark.parametrize("kind,count", [("dict", 1), ("list", 1), ("dict", 5), ("list", 5)])
def test_single_label_excel_handles_one_and_many_frames(tree, samples, tmp_path, kind, count):
    datasets = _datasets(samples, count, kind)
    frames = list(datasets.values()) if kind == "dict" else datasets
    names = list(datasets) if kind == "dict" else [f"数据集{i+1}" for i in range(count)]
    path = tmp_path / f"{kind}-{count}.xlsx"
    returned = tree.report(
        datasets=datasets, target="FPD", amount="金额", date="日期", methods=["quantile"], save=path, round_digits=False
    )
    assert isinstance(returned, dict if kind == "dict" else list)
    book = load_workbook(path)
    assert len(book.sheetnames) == 2
    main, features = book.worksheets
    assert len(main._images) == 2 * (count + 1)
    title_row = next(c.row for row in main for c in row if c.value == "3、数据样本统计")
    header = title_row + 5
    assert main.cell(header, 2).value == "数据集"
    assert [main.cell(header + 1 + i, 2).value for i in range(count)] == names
    for i, frame in enumerate(frames):
        assert main.cell(header + 1 + i, 3).value == len(frame)
        assert main.cell(header + 1 + i, 5).value == frame.FPD.sum()
    headers = [
        (r, c)
        for r in range(1, main.max_row + 1)
        for c in range(2, main.max_column - 3)
        if [main.cell(r, c + j).value for j in range(5)] == ["节点编号", "指标含义", "组合策略", "是否叶子", "是否调整"]
    ]
    assert len(headers) == 2 * (count + 1)
    rules = tree.get_rules(leaf_only=False)
    for i, frame in enumerate(frames):
        for basis in [0, 1]:
            row, col = headers[2 * (i + 1) + basis]
            columns = [main.cell(row, c).value for c in range(col, main.max_column + 1)]
            for ordinal, rule in enumerate(rules):
                values = [main.cell(row + 1 + ordinal, c).value for c in range(col, main.max_column + 1)]
                mask = rule.predict(frame)
                expected = frame.loc[mask, "金额"].sum() if basis else mask.sum()
                assert values[2] == rule.expr
                assert values[columns.index("命中数")] == pytest.approx(expected)
    assert len(features.tables) == count
    assert features.auto_filter.ref
    assert all(c.data_type != "e" for sheet in book for row in sheet for c in row)
    book.close()
