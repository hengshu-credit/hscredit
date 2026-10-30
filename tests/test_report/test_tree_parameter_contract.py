"""手工树与多叉树的规范参数、字段集合及已移除入口契约。"""

import inspect

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from hscredit.report.mining import DecisionTreeAnalyzer, ManualTreeExtractor, MultiwayTreeExtractor

CLASSES = [ManualTreeExtractor, MultiwayTreeExtractor, DecisionTreeAnalyzer]


@pytest.fixture(params=CLASSES)
def tree_class(request):
    return request.param


@pytest.fixture
def sample():
    return pd.DataFrame(
        {
            "评分": range(48),
            "多头": np.arange(48) % 7,
            "FPD": [0] * 24 + [1] * 24,
            "金额": np.arange(48) + 100,
            "MOB1": [0, 2, 5, 8] * 12,
        }
    )


def test_report_and_rule_table_have_same_public_parameters():
    for method in ["report", "get_rule_table"]:
        signatures = [inspect.signature(getattr(cls, method)) for cls in CLASSES]
        expected = [(name, p.kind, p.default) for name, p in signatures[0].parameters.items()]
        for signature in signatures:
            assert [(name, p.kind, p.default) for name, p in signature.parameters.items()] == expected
            assert "amount" in signature.parameters and "feature_map" in signature.parameters
            assert not {"data", "sheet", "max_branches"}.intersection(signature.parameters)
    assert inspect.signature(MultiwayTreeExtractor).parameters["max_n_bins"].default == 4
    assert "max_branches" not in inspect.signature(MultiwayTreeExtractor).parameters


@pytest.mark.parametrize("factory", [list, tuple, np.array, pd.Index])
@pytest.mark.parametrize("fit_override", [False, True])
def test_feature_containers_and_explicit_fit_precedence(tree_class, sample, factory, fit_override):
    features = factory(["评分", "多头"])
    tree = tree_class(target="FPD", features=["缺少列"] if fit_override else features, n_jobs=1)
    tree.fit(sample, **({"features": features} if fit_override else {}))
    actual = tree.features_ if hasattr(tree, "features_") else tree._feature_list
    assert actual == ["评分", "多头"]
    assert list(features) == ["评分", "多头"]
    assert tree.features is features if not fit_override else tree.features == ["缺少列"]


@pytest.mark.parametrize(
    "kwargs",
    [
        {"sheet": "旧页"},
        {"data": None},
        {"max_branches": 3},
        {"excel_path": "旧.xlsx"},
        {"amount_col": "金额"},
        {"n_jobs": 2},
        {"unknown_option": True},
    ],
)
def test_removed_or_misplaced_parameters_are_not_silently_ignored(tree_class, sample, kwargs):
    tree = tree_class(target="FPD", features="评分", n_jobs=1).fit(sample)
    for method in [tree.report, tree.get_rule_table]:
        with pytest.raises(TypeError):
            method(sample, **kwargs)


@pytest.mark.parametrize("metric", ["LIFT", "LIFT值", "坏样本率", "bad_rate"])
def test_only_canonical_filter_metric_names_are_accepted(tree_class, metric):
    with pytest.raises(ValueError, match="lift.*badrate"):
        tree_class(rule_filter={"metric": metric})


@pytest.mark.parametrize("metric,threshold", [("lift", 1.0), ("badrate", 0.5)])
def test_canonical_report_uses_amount_feature_map_and_training_filter(tree_class, sample, tmp_path, metric, threshold):
    tree = tree_class(
        target="FPD",
        features=("评分",),
        feature_map={"评分": "构造说明"},
        n_jobs=1,
        rule_filter={"metric": metric, "threshold": threshold},
    ).fit(sample)
    datasets = {"验证集": sample.iloc[::2], "训练集": sample}
    path = tmp_path / f"{metric}.xlsx"
    result = tree.report(
        datasets,
        target="FPD",
        amount="金额",
        sheet_name="规则挖掘报告",
        save=path,
        methods=["quantile"],
        feature_map={"评分": "本次说明"},
    )
    expected = tree.get_rule_table(
        datasets,
        amount="金额",
        feature_map={"评分": "本次说明"},
        methods=["quantile"],
        sheet_name="规则挖掘报告",
        save=tmp_path / "规则表导出.xlsx",
    )
    for name in datasets:
        pd.testing.assert_frame_equal(result[name], expected[name])
        assert all("本次说明" in value for value in result[name]["字段含义"])
    assert tree.feature_map == {"评分": "构造说明"}
    book = load_workbook(path)
    assert book.sheetnames == ["规则挖掘报告", "规则挖掘报告_特征分箱"]
    book.close()


def test_single_frame_external_amount_matches_column(tree_class, sample):
    tree = tree_class(target="FPD", features="评分", n_jobs=1).fit(sample)
    expected = tree.get_rule_table(sample, amount="金额")
    actual = tree.get_rule_table(sample, amount=sample["金额"].iloc[::-1])
    pd.testing.assert_frame_equal(expected, actual)
