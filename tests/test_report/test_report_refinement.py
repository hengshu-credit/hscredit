"""精简报告展示，并按首标签独立核验分月指标和权重对齐。"""

import pickle

import numpy as np
import pandas as pd
import pytest
from numpy.testing import assert_allclose
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import Pipeline

from hscredit import ScorecardFeatureSelection
from hscredit.report import ModelReport
from hscredit.report._model_protocol import label_configuration_table
from hscredit.report._training_process_report import compact_selection_details, process_tables


class ProbabilityModel:
    feature_names_in_ = np.array(["概率"])
    classes_ = np.array([0, 1])

    def predict_proba(self, X):
        p = np.asarray(X["概率"])
        return np.column_stack((1 - p, p))


@pytest.fixture
def data():
    return pd.DataFrame(
        {
            "概率": [0.2, 0.99, 0.01, 0.8, 0.4, 0.3, 0.9, 0.55, 0.05, 0.1, 0.7, 0.3],
            "MOB1": [0, 1, 4, 8, 10, 0, 0, 2, 4, 8, 10, 0],
            "日期": pd.to_datetime(["2026-01-01"] * 6 + ["2026-02-01"] * 6),
            "权重": [1, 20, 30, 2, 4, 1, 3, 40, 50, 1, 2, 1],
        },
        index=np.repeat(np.arange(6), 2),
    )


@pytest.mark.parametrize("mode", ["overdue", "target"])
@pytest.mark.parametrize("weighted", [False, True])
def test_monthly_metrics_use_first_label_and_its_valid_rows(data, mode, weighted):
    first = data.MOB1.gt(7).astype(float).mask(data.MOB1.gt(0) & data.MOB1.le(7))
    options = (
        {"overdue": "MOB1", "dpds": [7, 0]}
        if mode == "overdue"
        else {"target": {"首标签": first, "次标签": data.MOB1.gt(0)}}
    )
    report = ModelReport(
        ProbabilityModel(),
        datasets={"样本": data},
        del_grey=True,
        sample_weight="权重" if weighted else None,
        n_jobs=1,
        **options,
    )
    monthly = report._get_monthly_metrics("日期").set_index("月份")
    for month in ["2026-01", "2026-02"]:
        mask = (data.日期.dt.strftime("%Y-%m") == month).to_numpy() & first.notna().to_numpy()
        y, p = first.to_numpy()[mask], data.概率.to_numpy()[mask]
        w = data.权重.to_numpy()[mask] if weighted else np.ones(len(y))
        bad, good = y == 1, y == 0
        difference = p[bad, None] - p[good][None, :]
        pairs = w[bad, None] * w[good][None, :]
        auc = np.sum(((difference > 0) + 0.5 * (difference == 0)) * pairs) / pairs.sum()
        order = np.argsort(p, kind="stable")
        ks = np.max(np.abs(np.cumsum((w * y)[order]) / w[bad].sum() - np.cumsum((w * (1 - y))[order]) / w[good].sum()))
        row = monthly.loc[month]
        assert row["样本数"] == 4
        assert_allclose(row[["坏样本率", "AUC", "KS"]].astype(float), [np.average(y, weights=w), auc, ks])
        assert_allclose(row[["均值", "最小值", "最大值"]].astype(float), [p.mean(), p.min(), p.max()])
    if mode == "overdue":
        reverse = ModelReport(
            ProbabilityModel(), datasets={"样本": data}, overdue="MOB1", dpds=[0, 7], del_grey=True, n_jobs=1
        )
        assert reverse._get_monthly_metrics("日期")["样本数"].tolist() == [6, 6]
    report._monthly_metrics_cache["日期"] = pd.DataFrame({"旧联合口径": [1]})
    restored = pickle.loads(pickle.dumps(report))
    pd.testing.assert_frame_equal(restored._get_monthly_metrics("日期").set_index("月份"), monthly)


def test_description_four_configs_and_monthly_psi_percentage(data, tmp_path):
    report = ModelReport(
        ProbabilityModel(), datasets={"样本": data}, overdue="MOB1", dpds=[7, 0], del_grey=True, n_jobs=1
    )
    config = label_configuration_table(report)
    assert config["配置项"].tolist() == ["标签来源", "操作符", "灰样本定义", "剔灰配置"]
    values = config.set_index("配置项")["配置值"]
    assert values["标签来源"] == "overdue+dpds" and values["操作符"] == ">" and values["剔灰配置"] == "是"
    assert "MOB1@7：0 < MOB1 <= 7" in values["灰样本定义"]
    expected_psi = report._get_monthly_psi_matrix("日期")
    path = tmp_path / "报告.xlsx"
    report.to_excel(path, date="日期", with_plots=False, include_sample_records=False, mode="strict")
    book = load_workbook(path)
    basic = book["1-基本信息"]
    basic_values = [c.value for row in basic for c in row]
    assert "配置项" in basic_values and "配置值" in basic_values
    assert "标签定义与剔灰规则" not in basic_values
    for field in config["配置项"]:
        assert basic_values.count(field) == 1
    assert not {"坏样本定义", "实际剔除规则", "缺失值处理", "原始样本数", "有效样本数", "剔除样本数"} & set(
        basic_values
    )
    for row in basic:
        if row[1].value in values:
            assert row[2].value == values[row[1].value]
    sheet = book["2-模型性能"]
    assert any(c.value == "2、分月模型效果（MOB1@7）" for row in sheet for c in row)
    title = next(c.row for row in sheet for c in row if isinstance(c.value, str) and c.value.endswith("、分月对比PSI"))
    header = next(
        r
        for r in range(title + 1, sheet.max_row)
        if sheet.cell(r, 3).value == "2026-01" and sheet.cell(r, 4).value == "2026-02"
    )
    for i, month in enumerate(expected_psi.index, header + 1):
        assert sheet.cell(i, 2).value == month
        for j, column in enumerate(expected_psi.columns, 3):
            assert_allclose(sheet.cell(i, j).value, expected_psi.loc[month, column])
            assert sheet.cell(i, j).number_format == "0.00%"
    book.close()


def test_label_configuration_preserves_mixed_external_y_source(data):
    report = ModelReport(
        ProbabilityModel(),
        datasets={"派生": data, "外部": (data, data.MOB1.gt(3).astype(int))},
        overdue="MOB1",
        dpds=[7, 0],
        del_grey=True,
        n_jobs=1,
    )
    values = label_configuration_table(report).set_index("配置项")["配置值"]
    assert values["标签来源"] == "派生：overdue+dpds；外部：显式y"
    assert "外部：不适用（直接标签）" in values["操作符"]
    assert "不应用 target/overdue 的灰样本规则" in values["灰样本定义"]


def test_selection_details_keep_leaf_decisions_and_preserve_full_process():
    rng = np.random.default_rng(9)
    X = pd.DataFrame({"甲": rng.normal(size=80), "乙": rng.normal(size=80), "空": np.nan, "常量": 1.0})
    X["甲副本"] = X["甲"]
    y = (X["甲"] + X["乙"] > 0).astype(int)
    selector = ScorecardFeatureSelection(
        iv_threshold=0.01,
        corr_threshold=0.9,
        n_jobs=1,
        bin_params={"method": "quantile", "max_n_bins": 3, "monotonic": False},
    )
    model = Pipeline([("筛选", selector), ("模型", LogisticRegression())]).fit(X, y)
    report = ModelReport(model, datasets={"样本": (X, y)}, n_jobs=1)
    raw = report.get_training_process()["筛选明细"]
    saved = raw.copy(deep=True)
    display = process_tables(report)["筛选明细"]
    expected = raw[(raw["记录类型"] == "特征决策") & (raw["阶段路径"] != "selector")]
    assert len(display) == len(expected) < len(raw)
    assert len(display.columns) <= 9
    assert not {"记录类型", "阶段路径", "筛选器", "计算域", "补充信息", "有效性", "输入顺序"} & set(display.columns)
    assert display["特征"].tolist() == expected["特征"].tolist()
    assert display["筛选结果"].tolist() == expected["筛选结果"].tolist()
    assert_allclose(display["指标值"].astype(float), expected["指标值"].astype(float), equal_nan=True)
    assert "剔除" in set(display["筛选结果"])
    pd.testing.assert_frame_equal(raw, saved)
    pd.testing.assert_frame_equal(report.get_training_process()["筛选明细"], saved)


def test_parallel_selection_root_decisions_are_not_discarded():
    details = pd.DataFrame(
        {
            "阶段路径": ["root", "root/a"],
            "阶段名称": ["并集", "分支"],
            "特征": ["甲", "甲"],
            "记录类型": ["特征决策"] * 2,
            "筛选结果": ["保留", "剔除"],
        }
    )
    summary = pd.DataFrame({"阶段路径": ["root", "root/a"], "关系类型": ["union", "single"]})
    result = compact_selection_details(details, summary)
    assert result["阶段名称"].tolist() == ["并集", "分支"]
    assert result["筛选结果"].tolist() == ["保留", "剔除"]


def test_same_named_selection_stages_remain_distinguishable():
    details = pd.DataFrame(
        {
            "阶段路径": ["root/a", "root/b"],
            "阶段编号": [2, 3],
            "阶段名称": ["缺失筛选"] * 2,
            "记录类型": ["特征决策"] * 2,
            "特征": ["甲"] * 2,
            "筛选结果": ["保留", "剔除"],
            "比较符": ["<"] * 2,
            "有效阈值": [0.9, 0.1],
        }
    )
    result = compact_selection_details(details, pd.DataFrame())
    assert result["阶段编号"].tolist() == [2, 3]
    assert result["筛选条件"].tolist() == ["< 0.9", "< 0.1"]
