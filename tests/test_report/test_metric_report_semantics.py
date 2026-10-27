"""模型TOP表、规则指标与报告故障状态的独立核算。"""

import numpy as np
import pandas as pd
import pytest

from hscredit.core.metrics import rule_lift
from hscredit.report import ModelReport, auto_feature_analysis
from hscredit.report.result import ReportGenerationError, ReportResult


class _ProbabilityModel:
    classes_ = np.array([0, 1])
    feature_names_in_ = np.array(["分数"])

    def predict_proba(self, data):
        p = data["分数"].to_numpy()
        return np.column_stack([1 - p, p])


def _report():
    data = pd.DataFrame({"分数": np.arange(8) / 10, "金额": [1, 2, 1, 2, 1, 2, 4, 4]})
    y = np.array([0, 0, 0, 0, 0, 0, 1, 1])
    return ModelReport(_ProbabilityModel(), X_train=data, y_train=y, features=["分数"], n_jobs=1), data, y


@pytest.mark.parametrize("amount", [None, "金额"])
def test_top_table_improvement_is_remaining_population_reduction(amount):
    report, data, y = _report()
    table = report._compute_top_n_lift_table(percentiles=[0.25, 0.5, 1.0], amount=amount)
    improvement = table.loc[table["统计项"] == "坏账改善"].iloc[0]
    assert improvement["TOP 25%"] == pytest.approx(1)
    assert improvement["TOP 50%"] == pytest.approx(1)
    assert improvement["TOP 100%"] == 1
    assert improvement["TOTAL"] == 0


def test_zero_amount_top_table_does_not_substitute_count_bad_rate():
    report, data, y = _report()
    data["金额"] = 0.0
    report = ModelReport(_ProbabilityModel(), X_train=data, y_train=y, features=["分数"], n_jobs=1)
    table = report._compute_top_n_lift_table(percentiles=[0.25], amount="金额")
    numeric = table.set_index("统计项")
    assert numeric.loc[["坏样本率", "LIFT值"], ["TOP 25%", "TOTAL"]].isna().all().all()
    assert numeric.loc["坏账改善", "TOP 25%"] == numeric.loc["坏账改善", "TOTAL"] == 0


def test_rule_lift_binary_integer_mask_is_a_mask_and_all_rejection_is_undefined():
    values = rule_lift([0, 1, 0, 1], [0, 0, 1, 1])
    assert values["hit_count"] == 2
    assert values["bad_count"] == 1
    assert values["bad_improve"] == 0
    assert rule_lift([0, 1, 0, 1], np.ones(4, dtype=bool))["bad_improve"] == 1
    assert rule_lift([0, 0], np.zeros(2, dtype=bool))["bad_improve"] == 0


def test_scorecard_summary_exception_is_visible_to_strict_report(monkeypatch):
    report, _, _ = _report()

    class BrokenSummary:
        coef_ = np.array([[1.0]])
        intercept_ = np.array([0.0])
        feature_names_in_ = np.array(["分数"])

        def summary(self):
            raise RuntimeError("审计注入：统计摘要失败")

    monkeypatch.setattr(report, "_scorecard_lr_model", lambda: BrokenSummary())
    report._execution_result = ReportResult(mode="strict")
    fallback = report._scorecard_lr_summary()
    assert not fallback.empty
    assert not report._execution_result.complete
    assert "统计摘要失败" in report._execution_result.status_table()["说明"].str.cat()
    with pytest.raises(ReportGenerationError):
        report._execution_result.ensure_publishable()


def test_non_linear_probability_scorecard_has_no_required_lr_coefficients(monkeypatch):
    report, _, _ = _report()
    monkeypatch.setattr(report, "_scorecard_lr_model", lambda: _ProbabilityModel())
    report._execution_result = ReportResult(mode="strict")
    assert report._scorecard_lr_summary().empty
    assert report._execution_result.complete
    assert report._execution_result.status_table()["状态"].eq("不适用").all()


@pytest.mark.parametrize("kind", ["model", "feature"])
def test_exported_workbook_does_not_append_metric_explanations(tmp_path, kind):
    from openpyxl import load_workbook

    report, data, y = _report()
    destination = tmp_path / f"{kind}.xlsx"
    if kind == "model":
        result = report.to_excel(
            str(destination),
            mode="strict",
            return_result=True,
            with_plots=False,
            show_importance=False,
            include_sample_records=False,
            method="quantile",
            max_n_bins=2,
        )
    else:
        data = data.assign(标签=y)
        result = auto_feature_analysis(
            data,
            features=["分数"],
            target="标签",
            excel_writer=str(destination),
            output_dir=str(tmp_path),
            pictures=[],
            bin_params={"method": "quantile", "max_n_bins": 2},
            margins=True,
            n_jobs=1,
            show_progress=False,
            mode="strict",
            return_result=True,
        )
    assert result.complete
    book = load_workbook(destination, read_only=True, data_only=True)
    cells = [cell for sheet in book for row in sheet.values for cell in row]
    assert "拒绝改善未定义" not in cells
    assert not any(isinstance(value, str) and "不适用（NaN）" in value for value in cells)
    book.close()


def test_undefined_metric_is_blank_in_excel_not_nan_text(tmp_path):
    from openpyxl import load_workbook
    from hscredit.core.metrics import compute_bin_stats
    from hscredit.excel import dataframe2excel

    data = compute_bin_stats([0, 0, 1, 1], [0, 0, 1, 1])[["累计风险拒绝比"]]
    destination = tmp_path / "边界.xlsx"
    dataframe2excel(data, str(destination), sheet_name="边界", start_row=1, start_col=1, index=False)
    book = load_workbook(destination, data_only=True)
    sheet = book["边界"]
    assert sheet.cell(1, 1).value == "累计风险拒绝比"
    assert sheet.cell(2, 1).value is None
    assert sheet.cell(3, 1).value == 2
    book.close()


@pytest.mark.parametrize("speed", ["normal", "fast"])
def test_missing_cells_are_blank_before_workbook_serialization(speed):
    from hscredit.excel import ExcelWriter, dataframe2excel

    writer = ExcelWriter()
    data = pd.DataFrame({"指标": [np.nan, pd.NA, pd.NaT, 0.25]})
    dataframe2excel(data, writer, sheet_name="边界", start_row=1, start_col=1, speed=speed, index=False)
    sheet = writer.workbook["边界"]
    assert [sheet.cell(row, 1).value for row in range(2, 6)] == [None, None, None, 0.25]
