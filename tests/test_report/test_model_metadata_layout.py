"""参数和部署页必须可读，调整模板列宽不能扩散到所有空白列。"""

import unicodedata

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from hscredit import LogisticRegression
from sklearn.base import BaseEstimator, ClassifierMixin
from hscredit.report import ModelReport


class NoImportanceClassifier(ClassifierMixin, BaseEstimator):
    def fit(self, X, y):
        self.feature_names_in_ = X.columns.to_numpy()
        self.classes_ = np.array([0, 1])
        return self

    def predict_proba(self, X):
        p = 0.2 + 0.5 * (np.arange(len(X)) % 2)
        return np.column_stack([1 - p, p])


def test_missing_native_importance_preserves_numeric_and_categorical_statistics(tmp_path):
    X = pd.DataFrame({"数值": [1.0, 2.0, 3.0, 4.0], "类别": ["甲", "甲", "乙", "甲"]})
    y = [0, 1, 0, 1]
    model = NoImportanceClassifier().fit(X, y)
    report = ModelReport(model, datasets={"训练": (X, y)}, n_jobs=1)
    table = report.get_features_describe()
    assert set(table.index) == set(X.columns)
    assert table.loc["数值", "平均值"] == 2.5
    assert table.loc["类别", "众数"] == "甲"
    assert table.loc["类别", "枚举数"] == 2
    assert table["特征重要性"].isna().all()
    correlation = report.get_features_corr()
    assert correlation.columns.tolist() == ["数值"]
    assert correlation.iloc[0, 0] == 1.0
    result = report.to_excel(tmp_path / "无原生重要性.xlsx", with_plots=False, mode="strict", return_result=True)
    assert result.complete


def test_long_feature_names_and_parameters_do_not_clip_or_resize_unused_columns(tmp_path):
    feature = "近六个月非银多头机构数量与真实完整字段名称"
    X = pd.DataFrame({feature: np.arange(30), "另一变量": np.arange(30) % 3})
    model = LogisticRegression(calculate_stats=False).fit(X, (X[feature] > 12).astype(int))
    output = tmp_path / "完整报告.xlsx"
    result = model.model_report(
        datasets={"验证集": (X, (X[feature] > 12).astype(int))},
        excel_path=output,
        mode="strict",
        return_result=True,
        verbose=False,
        with_plots=False,
        n_jobs=1,
    )
    assert result.complete
    workbook = load_workbook(output)
    for name in ("5-模型参数", "6-模型部署需求", "7-训练与调参过程"):
        sheet = workbook[name]
        assert sheet.sheet_view.showGridLines is False
        assert not any(d.min <= sheet.max_column < d.max for d in sheet.column_dimensions.values())
        for row in sheet:
            for cell in row:
                if cell.value == feature:
                    units = sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in feature)
                    assert cell.alignment.wrap_text or sheet.column_dimensions[cell.column_letter].width >= units
        assert not any(isinstance(c.value, str) and len(c.value) > 2000 for row in sheet for c in row)
    workbook.close()


@pytest.mark.parametrize("workers", [1, 2])
def test_categorical_model_report_plots_only_numeric_correlations(tmp_path, workers):
    from hscredit import CatBoost

    rng = np.random.default_rng(7)
    X = pd.DataFrame({"甲": rng.normal(size=60), "乙": rng.normal(size=60)})
    X["类别"] = np.where(X["乙"] > 0, "高", "低")
    y = (X["甲"] > 0).astype(int)
    model = CatBoost(iterations=3, depth=2, n_jobs=1, allow_writing_files=False).fit(X, y, cat_features=["类别"])
    report = ModelReport(model, datasets={"验证": (X, y)}, n_jobs=workers, parallel_backend="threading")
    result = report.to_excel(tmp_path / "类别输入.xlsx", with_plots=True, mode="strict", return_result=True)
    assert result.complete
    assert (tmp_path / "类别输入_assets/feature_corr.png").is_file()
    assert set(report.get_features_corr()) == {"甲", "乙"}
