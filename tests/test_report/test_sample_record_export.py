"""订单样例明细开关真实工作簿回归，不以monkeypatch跳过导出。"""

import inspect
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression

from hscredit.report import ModelReport, auto_model_report


def _data_and_model():
    data = pd.DataFrame({"x": [1.25, 2.5, 3.75, 5.0], "y": [0, 1, 0, 1],
                         "订单标识": ["PRIVATE_ORDER_A", "PRIVATE_ORDER_B", "PRIVATE_ORDER_C", "PRIVATE_ORDER_D"]})
    model = LogisticRegression().fit(data[["x"]], data.y)
    return data, model


@pytest.mark.parametrize("entry", ["direct", "auto", "transactional"])
def test_disabling_sample_records_removes_row_level_table_and_marks_not_applicable(tmp_path, entry):
    data, model = _data_and_model()
    output = tmp_path / "without_records.xlsx"
    kwargs = dict(with_plots=False, loc_cols="订单标识", include_sample_records=False, mode="strict", return_result=True)
    if entry == "auto":
        result = auto_model_report(model, X_train=data[["x", "订单标识"]], y_train=data.y,
                                   excel_path=str(output), n_jobs=1, verbose=False, **kwargs)
    else:
        report = ModelReport(model, X_train=data[["x", "订单标识"]], y_train=data.y, n_jobs=1)
        result = report.to_excel(str(output), transactional=entry == "transactional", **kwargs)
    assert result.complete
    assert result.sections["生产订单测试用例"].status == "不适用"
    assert result.metadata["包含订单样例明细"] is False
    workbook_path = Path(result.artifacts[0]["路径"])
    workbook = load_workbook(workbook_path)
    all_values = [cell.value for sheet in workbook.worksheets for row in sheet for cell in row]
    assert not any(identifier in all_values for identifier in data["订单标识"])
    deployment_values = [cell.value for row in workbook["6-模型部署需求"] for cell in row]
    assert "订单标识" not in deployment_values
    assert "模型分数" not in deployment_values
    assert "2、生产订单测试用例（不适用：已关闭明细输出）" in deployment_values
    workbook.close()


def test_sample_records_default_remains_true_and_keyword_only(tmp_path):
    for entry in (ModelReport.to_excel, auto_model_report):
        parameter = inspect.signature(entry).parameters["include_sample_records"]
        assert parameter.default is True
        assert parameter.kind is inspect.Parameter.KEYWORD_ONLY
    data, model = _data_and_model()
    output = tmp_path / "with_records.xlsx"
    report = ModelReport(model, X_train=data[["x", "订单标识"]], y_train=data.y, n_jobs=1)
    report.to_excel(str(output), with_plots=False, loc_cols="订单标识")
    workbook = load_workbook(output)
    values = [cell.value for row in workbook["6-模型部署需求"] for cell in row]
    assert "PRIVATE_ORDER_A" in values
    assert "模型分数" in values
    workbook.close()


def test_sample_record_switch_rejects_truthy_strings(tmp_path):
    data, model = _data_and_model()
    report = ModelReport(model, X_train=data[["x"]], y_train=data.y, n_jobs=1)
    with pytest.raises(ValueError, match="必须为布尔值"):
        report.to_excel(str(tmp_path / "bad.xlsx"), include_sample_records="False")
