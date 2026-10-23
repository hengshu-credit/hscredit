"""训练、指标与报告采用同一预测语义和训练版本。"""

import numpy as np
import pandas as pd
import pytest
from sklearn.datasets import make_classification

from hscredit import LightGBM, XGBoost, RandomForest, get_loss, get_metric
from hscredit.core.models.tuning import Metric
from hscredit.report import ModelReport


class Predictions:
    classes_ = np.array([0, 1])
    feature_names_in_ = ["概率"]

    def predict_proba(self, X):
        p = X["概率"].to_numpy()
        return np.column_stack([1 - p, p])

    def predict_score(self, X):
        return 1000 * (1 - X["概率"].to_numpy())


def test_report_uses_explicit_probability_direction_and_common_tail_rules():
    X = pd.DataFrame({"概率": [0.9, 0.8, 0.3, 0.1]})
    y = np.array([0, 0, 1, 1])
    model = Predictions()
    probability = ModelReport(model, X_train=X, y_train=y, n_jobs=1)
    score = ModelReport(model, X_train=X, y_train=y, prediction_method="predict_score", n_jobs=1)
    assert probability.get_metrics().set_index("统计项").loc["AUC", "训练集"] == 0
    assert score.get_metrics().set_index("统计项").loc["AUC", "训练集"] == 0
    assert Metric("auc")(y, X["概率"]) == get_metric("auc")(y, X["概率"]) == 0
    pd.testing.assert_frame_equal(
        probability._get_top_n_lift_table(percentiles=(0.5,)), score._get_top_n_lift_table(percentiles=(0.5,))
    )
    np.testing.assert_array_equal(score._datasets["训练集"].prediction, model.predict_score(X))


@pytest.mark.parametrize("weighted", [False, True])
def test_tied_top_lift_matches_common_metric_and_is_row_order_independent(weighted):
    X = pd.DataFrame({"概率": [0.5] * 10, "金额": np.arange(1, 11)})
    y = np.array([1, 0, 0, 0, 0, 0, 0, 0, 0, 1])
    for order in (np.arange(10), np.arange(10)[::-1]):
        data = X.iloc[order]
        report = ModelReport(
            Predictions(), X_train=data, y_train=y[order], sample_weight="金额" if weighted else None, n_jobs=1
        )
        table = report._get_top_n_lift_table(percentiles=(0.1,)).set_index("统计项")
        metric = get_metric("topk_lift", top_ratio=0.1)
        expected = metric(y[order], data["概率"], sample_weight=data["金额"] if weighted else None)
        assert table.loc["LIFT值", "TOP 10%"] == pytest.approx(expected)
        assert expected == pytest.approx(1)


def test_report_financial_metric_and_weight_match_model_evaluation():
    X, y = make_classification(n_samples=80, n_features=4, random_state=4)
    X = pd.DataFrame(X, columns=list("abcd"))
    model = RandomForest(n_estimators=4, random_state=4, n_jobs=1).fit(X, y)
    metric = get_loss("cost_sensitive", fn_cost=7).metric(kind="business")
    weights = np.arange(1, 81)
    expected = model.evaluate(X, y, metrics=[metric, "auc", "ks", "lift@10%"], sample_weight=weights)
    report = ModelReport(model, X_train=X, y_train=y, metrics=[metric], sample_weight=weights, n_jobs=1)
    actual = report.get_metrics().set_index("统计项")["训练集"]
    for key in [metric.name, "AUC", "KS"]:
        assert actual[key] == pytest.approx(expected[key])
    assert report._get_top_n_lift_table(percentiles=(0.1,)).set_index("统计项").loc[
        "LIFT值", "TOP 10%"
    ] == pytest.approx(expected["LIFT@10%"])


def test_report_refuses_model_revision_mix_and_snapshots_configuration(tmp_path):
    X, y = make_classification(n_samples=60, n_features=4, random_state=7)
    model = RandomForest(n_estimators=3, random_state=7, n_jobs=1).fit(X, y)
    report = ModelReport(model, X_train=X, y_train=y, n_jobs=1)
    report.get_metrics()
    model.set_params(n_estimators=8)
    assert report.model_context_["构造参数"]["n_estimators"] == 3
    model.fit(X, y)
    with pytest.raises(ValueError, match="重新训练"):
        report.get_metrics()
    with pytest.raises(ValueError, match="重新训练"):
        report.to_excel(str(tmp_path / "stale.xlsx"), with_plots=False)
    assert not (tmp_path / "stale.xlsx").exists()


def test_long_curves_survive_repeated_fit_and_report_export(tmp_path):
    X, y = make_classification(n_samples=90, n_features=4, random_state=9)
    X = pd.DataFrame(X, columns=list("abcd"))
    model = LightGBM(n_estimators=40, n_jobs=1, random_state=9).fit(X, y, eval_set={"验证集": (X, y)})
    expected = model.get_learning_curve().copy()
    assert len(expected) == 40
    model.set_params(n_estimators=3).fit(X, y, eval_set={"验证集": (X, y)})
    pd.testing.assert_frame_equal(model.get_learning_curve(revision=1), expected)
    assert len(model.get_training_history()) == 2
    report = ModelReport(model, X_train=X, y_train=y, metrics=[get_metric("brier")], n_jobs=1)
    report.to_excel(str(tmp_path / "process.xlsx"), with_plots=False, include_sample_records=False, mode="strict")
    from openpyxl import load_workbook

    book = load_workbook(tmp_path / "process.xlsx", read_only=True)
    assert "8-训练与调参过程" in book.sheetnames
    assert any(cell.value == "概率均方误差" for row in book["2-模型性能"] for cell in row)
    book.close()


@pytest.mark.parametrize("model_class", [LightGBM, XGBoost])
def test_direct_custom_loss_rejects_ignored_native_weight(model_class):
    X, y = make_classification(n_samples=60, n_features=4, random_state=5)
    model = model_class(objective=get_loss("weighted_bce"), scale_pos_weight=9, n_estimators=3, n_jobs=1)
    with pytest.raises(ValueError, match="自定义 loss"):
        model.fit(X, y)


def test_report_rejects_probability_metrics_for_credit_scores():
    with pytest.raises(ValueError, match="概率"):
        ModelReport(
            Predictions(),
            X_train=pd.DataFrame({"概率": [0.1, 0.9]}),
            y_train=[0, 1],
            metrics=[get_metric("profit")],
            prediction_method="predict_score",
            n_jobs=1,
        )


@pytest.mark.parametrize("model_class", [LightGBM, XGBoost])
@pytest.mark.parametrize("loss_name", ["amount_weighted", "expected_value"])
def test_auto_split_financial_parameters_matches_explicit_manual_split(model_class, loss_name):
    from sklearn.model_selection import train_test_split

    X, y = make_classification(n_samples=100, n_features=4, random_state=15)
    amounts = np.arange(1, 101, dtype=float)
    train, valid = train_test_split(np.arange(100), test_size=0.2, random_state=15, stratify=y)
    parameters = (
        {"amounts": amounts} if loss_name == "amount_weighted" else {"ead": amounts, "lgd": np.linspace(0.2, 0.8, 100)}
    )
    loss = get_loss(loss_name, **parameters)
    options = dict(n_estimators=6, random_state=15, n_jobs=1, early_stopping_rounds=2)
    automatic = model_class(objective=loss, **options).fit(X, y)
    train_loss = get_loss(loss_name, **{key: value[train] for key, value in parameters.items()})
    valid_loss = get_loss(loss_name, **{key: value[valid] for key, value in parameters.items()})
    manual = model_class(objective=train_loss, eval_metric=valid_loss.metric(), **options).fit(
        X[train], y[train], eval_set=[(X[valid], y[valid])]
    )
    np.testing.assert_allclose(automatic.predict_proba(X), manual.predict_proba(X))
    assert automatic.training_summary_["金融参数切分"] == {"原始样本": 100, "训练样本": 80, "验证样本": 20}
    assert len(loss.amounts_ if loss_name == "amount_weighted" else loss.ead) == 100


def test_report_process_is_snapshot_when_tuner_continues():
    from hscredit import ModelTuner

    X, y = make_classification(n_samples=60, n_features=4, random_state=5)
    tuner = ModelTuner(RandomForest(n_estimators=3, n_jobs=1), search_space={}, n_jobs=1, random_state=5)
    tuner.fit(X, y, n_trials=1, show_progress_bar=False)
    report = ModelReport(tuner.get_best_model(), X_train=X, y_train=y, n_jobs=1)
    before = report.get_training_process()
    tuner.fit(X, y, n_trials=1, show_progress_bar=False)
    pd.testing.assert_frame_equal(report.get_training_process()["调参记录"], before["调参记录"])
    assert len(before["调参记录"]) == 1


def test_tuner_learning_curve_and_report_metric_convenience_entry():
    X, y = make_classification(n_samples=60, n_features=4, random_state=5)
    model = LightGBM(n_estimators=4, n_jobs=1).tune(
        X, y, search_space={}, n_trials=1, metric=get_metric("auc"), n_jobs=1
    )
    curves = model.tuner.get_learning_curve()
    assert set(["试验编号", "折编号", "迭代", "指标", "数值"]).issubset(curves.columns)
    assert len(curves) > 0
    report = model.report(X_train=X, y_train=y, metrics=[get_metric("brier")], verbose=False, n_jobs=1)
    assert "概率均方误差" in report.get_metrics()["统计项"].tolist()


def test_custom_metric_report_serializes_and_native_model_refit_is_detected():
    import pickle
    from sklearn.linear_model import LogisticRegression

    X, y = make_classification(n_samples=60, n_features=4, random_state=8)
    model = LogisticRegression().fit(X, y)
    metric = get_metric(lambda y, p: np.mean((y - p) ** 2), name="自定义误差", greater_is_better=False)
    report = ModelReport(model, X_train=X, y_train=y, metrics=[metric], n_jobs=1)
    restored = pickle.loads(pickle.dumps(report))
    pd.testing.assert_frame_equal(report.get_metrics(), restored.get_metrics())
    model.fit(X, 1 - y)
    with pytest.raises(ValueError, match="重新训练"):
        report.get_metrics()


def test_lift_plot_keeps_ties_together_and_uses_weights():
    from hscredit.core.viz import lift_plot
    import matplotlib.pyplot as plt

    y, p, weights = np.array([0, 0, 1, 1]), np.full(4, 0.5), np.array([1.0, 3.0, 2.0, 5.0])
    for order in (np.arange(4), np.arange(4)[::-1]):
        figure = lift_plot(y[order], p[order], n_bins=4, sample_weight=weights[order])
        np.testing.assert_allclose(figure.axes[0].lines[1].get_ydata(), [1])
        plt.close(figure)
