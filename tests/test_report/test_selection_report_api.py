"""筛选报告公共调用、数值与文件格式契约。"""

import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from sklearn.pipeline import Pipeline

from hscredit.core.selectors import NullSelector, SelectionReport, VarianceSelector, collect_selection_report
from hscredit.excel import ExcelWriter
from hscredit.exceptions import ValidationError
from hscredit.report import ReportResult, auto_selection_report, load_selection_report


@pytest.fixture
def selector():
    return NullSelector(threshold=0.4, n_jobs=1).fit(
        pd.DataFrame({"评分": [1, 2, 3, 4], "缺失字段": [1, np.nan, 3, np.nan], "全空": [np.nan] * 4})
    )


def test_report_uses_fitted_values_and_independent_counts(selector, monkeypatch):
    monkeypatch.setattr(NullSelector, "fit", lambda *a, **k: pytest.fail("不得重新拟合"))
    tables = auto_selection_report(source=selector)
    assert type(tables) is dict
    assert list(tables) == ["筛选汇总", "特征决策", "指标明细", "迭代记录"]
    summary = tables["筛选汇总"].iloc[0]
    assert summary["输入特征数"] == 3
    assert summary["选中特征数"] == 1
    assert summary["剔除特征数"] == 2
    assert summary["保留率"] == 1 / 3
    detail = tables["特征决策"]
    assert detail["特征"].tolist() == ["评分", "缺失字段", "全空"]
    assert detail["指标值"].tolist() == [0, 0.5, 1]
    detail.loc[:, "筛选结果"] = "修改副本"
    assert selector.selected_features_ == ["评分"]
    result = auto_selection_report(selector, return_result=True)
    assert isinstance(result, ReportResult)
    assert result.complete
    assert result.sections["特征决策"].input_rows is None
    assert result.sections["特征决策"].output_rows == 3


def test_pipeline_and_snapshot_return_identical_tables():
    data = pd.DataFrame({"评分": [1, 2, 3], "常量": [1, 1, 1]})
    pipeline = Pipeline([("缺失", NullSelector(n_jobs=1)), ("方差", VarianceSelector(n_jobs=1))]).fit(data)
    snapshot = collect_selection_report(pipeline)
    for name, frame in auto_selection_report(pipeline).items():
        pd.testing.assert_frame_equal(frame, auto_selection_report(snapshot)[name])


def test_direct_excel_path_preserves_precision_and_text(tmp_path):
    report = SelectionReport(details=[{"特征": "=1+1", "指标值": 1 / 3}, {"特征": "空值", "指标值": pd.NA}])
    destination = tmp_path / "nested" / "筛选.xlsx"
    result = report.to_excel(destination, return_result=True, writer_params={"theme_color": "123456"})
    assert destination.is_file()
    assert not list(tmp_path.rglob("*.json"))
    assert result.artifacts == [{"类型": "Excel", "路径": str(destination), "完整": True}]
    with_workbook = load_workbook(destination, data_only=False)
    try:
        sheet = with_workbook["特征决策"]
        assert sheet["F2"].value == "=1+1"
        assert sheet["F2"].data_type == "s"
        assert sheet["M2"].value == pytest.approx(1 / 3, abs=1e-15)
        assert sheet["M3"].value is None
        assert sheet["A1"].fill.fgColor.rgb.endswith("123456")
        assert sheet.freeze_panes == "A2"
        assert [cell.value for cell in with_workbook["参数说明"][1]] == ["阶段路径", "参数", "拟合标识"]
    finally:
        with_workbook.close()


def test_existing_writer_does_not_save_or_overwrite_sheet(selector, monkeypatch):
    writer = ExcelWriter()
    writer.get_sheet_by_name("特征决策")["A1"] = "原报告"
    monkeypatch.setattr(writer, "save", lambda *a, **k: pytest.fail("调用方负责保存"))
    result = auto_selection_report(selector, excel_writer=writer, max_rows_per_sheet=2, return_result=True)
    assert writer.workbook["特征决策"]["A1"].value == "原报告"
    assert writer.workbook["特征决策_1"].max_row == 3
    assert writer.workbook["特征决策_2"].max_row == 2
    assert result.artifacts == []
    auto_selection_report(selector, excel_writer=writer)
    assert writer.workbook["特征决策_3"]["A1"].value == "阶段编号"


def test_incomplete_snapshot_is_rejected_unless_best_effort(tmp_path):
    report = collect_selection_report([], strict=False)
    with pytest.raises(ValueError, match="未完成"):
        auto_selection_report(report, excel_writer=tmp_path / "拒绝.xlsx")
    assert not list(tmp_path.iterdir())
    result = auto_selection_report(report, mode="best_effort", return_result=True)
    assert not result.complete
    assert result.metadata["诊断"]


@pytest.mark.parametrize(
    "options",
    [{"max_excel_cells": True}, {"max_sheets": True}, {"max_rows_per_sheet": False}, {"by_selector": "False"}],
)
def test_invalid_export_limits_fail_before_creating_files(tmp_path, selector, options):
    report = collect_selection_report(selector)
    with pytest.raises(ValidationError):
        report.save(tmp_path / "归档", **options)
    with pytest.raises(ValidationError):
        report.to_excel(tmp_path / "直接.xlsx", **options)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("formats", [None, 1, [["json"]], ["json", "json"], {"json": True}])
def test_invalid_formats_raise_chinese_validation(tmp_path, selector, formats):
    with pytest.raises(ValidationError, match="formats"):
        collect_selection_report(selector).save(tmp_path / "报告", formats=formats)
    assert not list(tmp_path.iterdir())


def test_numpy_integer_limits_produce_serializable_mapping(tmp_path, selector):
    saved = collect_selection_report(selector).save(
        tmp_path / "整数预算", max_rows_per_sheet=np.int64(2), max_sheets=np.int64(20)
    )
    assert all(item["行数"] <= 2 for item in saved["工作表映射"])
    json.dumps(saved, ensure_ascii=False)


@pytest.mark.parametrize(
    "value", [{"类型": "浮点特殊值", "值": "损坏"}, {"类型": "日期", "值": "无效日期"}, {"类型": "映射", "值": [1]}]
)
def test_invalid_type_tags_raise_validation_error(tmp_path, value):
    saved = SelectionReport(metadata={"自定义": 1}).save(tmp_path / "报告", formats="json")
    path = Path(saved["路径"]["json"])
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["元数据"] = value
    path.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValidationError, match="JSON"):
        load_selection_report(path)


def test_real_workbook_null_rates_and_roundtrip(tmp_path):
    path = Path(__file__).resolve().parents[2] / "examples" / "hscredit_yyp.xlsx"
    if not path.exists():
        pytest.skip("本地演示工作簿 hscredit_yyp.xlsx 未提供")
    fields = ["衡枢鉴真分老客版", "近六个月非银多头机构数", "青云24", "商品类别"]
    data = pd.read_excel(path, usecols=fields + ["FPD"])
    selector = NullSelector(target="FPD", threshold=0.5, n_jobs=1).fit(data)
    result = auto_selection_report(selector, return_result=True)
    detail = result["特征决策"].set_index("特征")
    for field in fields:
        expected = sum(pd.isna(value) for value in data[field]) / len(data)
        assert detail.loc[field, "指标值"] == expected
    assert "FPD" not in detail.index
    snapshot = collect_selection_report(selector)
    snapshot.to_excel(tmp_path / "真实数据.xlsx")
    saved = snapshot.save(tmp_path / "快照", formats="json")
    restored = load_selection_report(saved["路径"]["json"])
    for name, frame in result.items():
        pd.testing.assert_frame_equal(frame, auto_selection_report(restored)[name])
