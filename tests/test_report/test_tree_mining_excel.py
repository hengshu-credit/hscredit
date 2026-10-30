"""树挖掘报告的统一 API、原始行统计与 Excel 落盘契约。"""

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from hscredit import MultiwayTreeExtractor
from hscredit.excel import ExcelWriter
from hscredit.report.mining import DecisionTreeAnalyzer, ManualTreeExtractor


@pytest.fixture
def data():
    return pd.DataFrame(
        {
            "评分": np.arange(48),
            "多头": np.tile(np.arange(12), 4),
            "FPD": ([0] * 8 + [1] * 4) * 4,
            "MOB1": [0, 2, 3, 4, 7, 8] * 8,
            "金额": np.arange(48) + 100,
        }
    )


@pytest.fixture(params=[MultiwayTreeExtractor, ManualTreeExtractor, DecisionTreeAnalyzer])
def tree(request, data):
    if request.param is MultiwayTreeExtractor:
        model = request.param(target="FPD", features=["评分", "多头"], max_depth=2, min_samples_leaf=2, n_jobs=1)
        model.fit(data).manual_split(feature="评分", thresholds=[15, 30])
    else:
        model = request.param(target="FPD", n_jobs=1)
        model.fit(data, features=["评分", "多头"], max_depth=2, min_samples_leaf=2)
        if request.param is ManualTreeExtractor:
            model.manual_split(data, feature="评分", threshold=22, node=0)
    return model


def test_shared_rule_apis_and_excel(tree, data, tmp_path):
    datasets = {"训练集": data, "验证集": data.iloc[::2], "测试集": data.iloc[1::2]}
    expected = tree.report(datasets, overdue=["MOB1"], dpds=[7, 3, 0], overdue_operator=">=", amount="金额")
    path = tmp_path / "reports" / "挖掘.xlsx"
    actual = tree.report(datasets, overdue=["MOB1"], dpds=[7, 3, 0], overdue_operator=">=", amount="金额", save=path)
    assert list(actual) == list(datasets)
    for key in datasets:
        pd.testing.assert_frame_equal(actual[key], expected[key])
    pd.testing.assert_frame_equal(tree.get_rule_table(datasets=data), tree.report(data))
    pd.testing.assert_frame_equal(tree.get_rule_table(data, "FPD"), tree.report(data, target="FPD"))
    all_rules = tree.get_rules(leaf_only=False)
    assert len(all_rules) == len(tree.get_rule_table())
    assert all_rules[0].expr == "True"
    assert len(tree.get_rules()) == len(tree.get_rule_table(leaf_only=True))
    book = load_workbook(path)
    assert book.sheetnames == ["组合策略汇总"]
    assert len(book["组合策略汇总"]._images) == 8
    detail = book["组合策略汇总"]
    assert not any(c.value == "完整节点信息" for row in detail for c in row)
    headers = [
        r
        for r in range(1, detail.max_row + 1)
        if [detail.cell(r, col).value for col in range(2, 5)] == ["节点编号", "是否叶子", "是否调整"]
    ]
    assert len(headers) == 8
    node_ids = tree.get_rule_table()["节点编号"].tolist()
    expressions = tree.get_rule_table()["指标名称"].tolist()
    for header, image in zip(headers, detail._images):
        image_start = image.anchor._from.row + 1
        assert image_start < header
        reserved_pixels = sum(
            (detail.row_dimensions[r].height or detail.sheet_format.defaultRowHeight or 15) / 0.75
            for r in range(image_start, header)
        )
        assert reserved_pixels >= image.anchor.ext.cy / 9525
        assert [detail.cell(header + 1 + i, 2).value for i in range(len(node_ids))] == node_ids
        assert [detail.cell(header + 1 + i, 5).value for i in range(len(node_ids))] == expressions
    sheet = book["组合策略汇总"]
    values = [cell.value for row in sheet for cell in row]
    assert sum(isinstance(v, str) and "最终树应用：" in v for v in values) == 6
    assert "MOB1>=7" in values
    assert "组合策略" in values and "命中数" in values and "命中率" in values
    assert "样本整体坏率" in values
    assert "节点编号" in values
    assert "是否调整" in values
    assert "规则数（含根节点）" in values and "树深度（根为0）" in values
    assert "AUC" in values and "KS" in values and "金额AUC" in values and "金额KS" in values
    assert "[0, 0.25)" in values and "[5, inf)" in values
    # 根节点金额按原始行直接核算，而非与另一个报告入口互相比较。
    first = actual["训练集"].iloc[0]
    assert first[("分箱详情", "样本总数")] == data["金额"].sum()
    assert first[("MOB1>=7", "坏样本数")] == data.loc[data.MOB1 >= 7, "金额"].sum()
    assert not any(c.data_type == "e" for ws in book for row in ws for c in row)
    assert all(not c.alignment.wrap_text for ws in book for row in ws for c in row if c.value is not None)
    assert all(dimension.width <= 50 for ws in book for dimension in ws.column_dimensions.values())
    for ws in book:
        dimensions = sorted(ws.column_dimensions.values(), key=lambda dim: dim.min)
        assert all(left.max < right.min for left, right in zip(dimensions, dimensions[1:]))
    assert sum(d.width < 30 for d in sheet.column_dimensions.values()) > 5
    book.close()


def test_report_containers_and_external_writer(tree, data, tmp_path):
    for frames in ([data, data.iloc[:24]], (data, data.iloc[:24])):
        assert len(tree.report(frames)) == 2
    path = tmp_path / "合并.xlsx"
    with ExcelWriter() as writer:
        writer.set_filename(str(path))
        writer.get_sheet_by_name("业务原表")["A1"] = "保留"
        tree.report(data, sheet="FPD", excel_writer=writer, save=path)
        tree.report(data, overdue=["MOB1"], dpds=[7, 3, 0], sheet="MOB", excel_writer=writer)
    book = load_workbook(path)
    assert book["业务原表"]["A1"].value == "保留"
    assert len(book["FPD"]._images) == len(book["MOB"]._images) == 2
    book.close()


def test_export_imported_binary_tree_without_training_data(data, tmp_path):
    from sklearn.tree import DecisionTreeClassifier

    clf = DecisionTreeClassifier(max_depth=2).fit(data[["评分", "多头"]], data.FPD)
    tree = ManualTreeExtractor.from_sklearn(clf, features=["评分", "多头"], target="FPD")
    path = tmp_path / "导入树.xlsx"
    tree.report(data, save=path)
    with path.open("rb") as stream:
        book = load_workbook(stream)
        assert len(book["组合策略汇总"]._images) == 2
        book.close()


def test_failed_export_preserves_existing_report(tree, data, tmp_path, monkeypatch):
    path = tmp_path / "原报告.xlsx"
    path.write_bytes(b"existing report")

    def fail(*args, **kwargs):
        raise ValueError("绘图失败")

    monkeypatch.setattr("hscredit.core.viz.plot_tree", fail)
    with pytest.raises(ValueError, match="绘图失败"):
        tree.report(data, save=path)
    assert path.read_bytes() == b"existing report"


def test_node_details_after_local_edits(data):
    tree = MultiwayTreeExtractor(target="FPD", features=["评分", "多头"], max_depth=3, n_jobs=1).fit(data)
    tree.manual_split(feature="评分", thresholds=[15, 30])
    child = tree.nodes_[0].children[1]
    tree.manual_split(node=child, feature="多头", thresholds=[4, 8])
    table = tree.get_node_table().set_index("节点编号")
    assert set(table.index) == set(tree.nodes_)
    assert str(table["父节点"].dtype) == "Int64"
    for node, rule in zip(tree.nodes_, tree.get_rules(leaf_only=False)):
        rows = data.loc[rule.predict(data)]
        assert table.loc[node, "样本总数"] == len(rows)
        assert table.loc[node, "坏样本数"] == rows.FPD.sum()
        assert table.loc[node, "坏样本率"] == pytest.approx(rows.FPD.mean())
        assert table.loc[node, "LIFT值"] == pytest.approx(rows.FPD.mean() / data.FPD.mean())
        assert table.loc[node, "指标含义"]
    tree.prune(node=child)
    after = tree.get_node_table()
    assert len(after) == len(tree.nodes_) < len(table)
    pd.testing.assert_frame_equal(tree.get_rule_table(data=data), tree.get_rule_table(datasets=data))
    pd.testing.assert_frame_equal(tree.get_rule_table(data, True), tree.get_rule_table(data, leaf_only=True))
    with pytest.raises(ValueError, match="不能同时"):
        tree.get_rule_table(datasets=data, data=data)


def test_no_export_without_save(tree, data, monkeypatch):
    def fail(*args, **kwargs):
        raise AssertionError("普通 report 不应触发绘图")

    monkeypatch.setattr("hscredit.core.viz.plot_tree", fail)
    assert not tree.report(data).empty


def test_dataset_overview_and_grey_labels_match_model_report(tree, data, tmp_path):
    from hscredit.report._sample_stats import build_sample_stats_table
    from hscredit.utils.targets import bind_targets

    data = data.assign(放款时间=pd.date_range("2026-01-01", periods=len(data), freq="D"))
    path = tmp_path / "多标签样本.xlsx"
    expected = tree.report(data, overdue=["MOB1"], dpds=[7, 3, 0], del_grey=True, leaf_only=True)
    result = tree.report(
        data,
        overdue=["MOB1"],
        dpds=[7, 3, 0],
        del_grey=True,
        leaf_only=True,
        date="放款时间",
        save=path,
        project_desc="按放款样本复核",
        data_source="样本表",
    )
    pd.testing.assert_frame_equal(result, expected)
    prepared, bound = bind_targets(data, overdue=["MOB1"], dpds=[7, 3, 0], del_grey=True)
    stats, _ = build_sample_stats_table(
        ["数据集1"], [bound.y_map(prepared)], bound.names(), target_specific_totals=True, n_jobs=1
    )
    book = load_workbook(path)
    sheet = book["组合策略汇总"]
    title = next(c.row for row in sheet for c in row if c.value == "3、数据样本统计")
    values = [sheet.cell(title + 4, col).value for col in range(3, 3 + len(stats.columns))]
    np.testing.assert_allclose(values, stats.iloc[0].to_numpy(dtype=float))
    assert any(c.value == "4、样本分布情况" for row in sheet for c in row)
    assert any(c.value == "样本区间" for row in sheet for c in row)
    assert any(c.value == "样本表" for row in sheet for c in row)
    book.close()


def test_export_target_callable_evaluated_once_per_dataset(tree, data, tmp_path):
    calls = []

    def target(frame):
        calls.append(len(frame))
        return frame.FPD

    result = tree.report({"训练集": data, "测试集": data.iloc[:20]}, target=target, save=tmp_path / "动态标签.xlsx")
    assert calls == [48, 20]
    assert len(result) == 2


@pytest.mark.parametrize("cls", [ManualTreeExtractor, DecisionTreeAnalyzer, MultiwayTreeExtractor])
def test_root_only_tree_rules_and_reports(cls, data):
    if cls is MultiwayTreeExtractor:
        tree = cls(target="FPD", features=["评分"], min_samples_split=100, n_jobs=1).fit(data)
    else:
        tree = cls(target="FPD", n_jobs=1).fit(data, features=["评分"], min_samples_split=100)
    assert tree.get_rules()[0].expr == "True"
    table = tree.get_rule_table(leaf_only=True)
    assert table["节点编号"].tolist() == [0]
    assert table["样本总数"].tolist() == [len(data)]
    if hasattr(tree, "get_leaf_node_ids"):
        assert tree.get_leaf_node_ids() == [0]


@pytest.mark.parametrize("cls", [ManualTreeExtractor, MultiwayTreeExtractor])
def test_rule_table_marks_adjusted_subtree_after_split_and_prune(cls, data):
    if cls is MultiwayTreeExtractor:
        tree = cls(target="FPD", features=["评分", "多头"], max_depth=3, min_samples_leaf=2, n_jobs=1).fit(data)
        assert set(tree.get_rule_table()["是否调整"]) == {"否"}
        tree.manual_split(node=0, feature="评分", thresholds=[15, 30])
        tree.prune(node=tree.nodes_[0].children[0])
    else:
        tree = cls(target="FPD", n_jobs=1).fit(data, features=["评分", "多头"], max_depth=3, min_samples_leaf=2)
        assert set(tree.get_rule_table()["是否调整"]) == {"否"}
        tree.manual_split(data, feature="评分", threshold=22, node=0)
        tree.delete_node(node=1)
    table = tree.get_rule_table()
    assert list(table.columns[:3]) == ["节点编号", "是否叶子", "是否调整"]
    assert set(table["是否调整"]) == {"是"}
    multi = tree.report(data, overdue=["MOB1"], dpds=[7, 3, 0])
    assert set(multi[("分箱详情", "是否调整")]) == {"是"}


def test_tree_views_reuse_final_routes_and_distinguish_count_and_amount(data):
    from hscredit.report.mining._tree_report import _tree_view
    from hscredit.core.viz.tree_plots import _extract_tree_data

    tree = MultiwayTreeExtractor(target="FPD", features=["评分"], n_jobs=1).fit(data)
    tree.manual_split(feature="评分", thresholds=[15, 30])
    counts_before = {node: current.counts.copy() for node, current in tree.nodes_.items()}
    evaluation = data.iloc[:20]
    original_edges = _extract_tree_data(tree)["edges"]
    for amount in [None, "金额"]:
        table = tree.report(evaluation, amount=amount, round_digits=False)
        view, _ = _tree_view(tree, table, amount is not None)
        for node, rule in zip(view["nodes"], tree.get_rules(leaf_only=False)):
            mask = rule.predict(evaluation)
            expected = evaluation.loc[mask, amount].sum() if amount else mask.sum()
            assert node["n_samples"] == expected
        assert view["nodes"][-1]["n_samples"] == 0
        assert len(view["edges"]) == 3
        assert view["edges"] == original_edges  # 连线只保留阈值/分组条件，数量在节点卡片显示。
    for node, before in counts_before.items():
        np.testing.assert_array_equal(tree.nodes_[node].counts, before)
    assert tree.edit_history_[-1]["操作"] == "指定阈值分裂"
    assert tree.edit_history_[-1]["切分点"] == [15, 30]


def test_model_metrics_match_raw_leaf_predictions_and_weights(tree, data):
    from sklearn.metrics import roc_auc_score
    from hscredit.report.mining._tree_report import _prediction_metrics

    sample = data.iloc[::2]
    scores = pd.Series(index=sample.index, dtype=float)
    for rule in tree.get_rules():
        training_rows = data.loc[rule.predict(data)]
        scores.loc[rule.predict(sample)] = training_rows.FPD.mean()
    result = _prediction_metrics(tree, {"验证集": sample}, [{"FPD": sample.FPD}], "金额")
    for basis, weights in [("订单", np.ones(len(sample))), ("金额", sample["金额"].to_numpy())]:
        expected_auc = roc_auc_score(sample.FPD, scores, sample_weight=weights)
        grouped = (
            pd.DataFrame(
                {
                    "分数": scores.to_numpy(),
                    "坏": weights * sample.FPD.to_numpy(),
                    "好": weights * (1 - sample.FPD.to_numpy()),
                }
            )
            .groupby("分数")[["坏", "好"]]
            .sum()
        )
        expected_ks = (
            (grouped["坏"].cumsum() / grouped["坏"].sum() - grouped["好"].cumsum() / grouped["好"].sum()).abs().max()
        )
        assert result[(basis, "FPD", "AUC")][0] == pytest.approx(expected_auc)
        assert result[(basis, "FPD", "KS")][0] == pytest.approx(expected_ks)


def test_zero_amount_keeps_order_report_and_blank_weighted_auc(data, tmp_path):
    from hscredit.report.mining._tree_report import _prediction_metrics

    data = data.assign(金额=0)
    tree = MultiwayTreeExtractor(target="FPD", features=["评分"], min_samples_leaf=2, n_jobs=1).fit(data)
    metrics = _prediction_metrics(tree, {"样本": data}, [{"FPD": data.FPD}], "金额")
    assert np.isnan(metrics[("金额", "FPD", "AUC")][0])
    assert np.isfinite(metrics[("订单", "FPD", "AUC")][0])
    path = tmp_path / "零金额.xlsx"
    result = tree.report({"样本": data}, amount="金额", save=path)
    assert result["样本"].iloc[0]["样本总数"] == 0
    book = load_workbook(path)
    assert len(book.active._images) == 4
    assert not any(c.data_type == "e" for row in book.active for c in row)
    book.close()
