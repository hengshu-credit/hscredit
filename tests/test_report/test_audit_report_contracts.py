"""审计反例转成报告正确性、明确口径及原子发布回归。"""

import hashlib

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from hscredit.core.eda import eda_summary, generate_report, vintage_analysis, vintage_summary
from hscredit.core.eda.report import export_report_to_excel
from hscredit.core.eda.strategy import score_strategy_simulation
from hscredit.excel import ExcelWriter
from hscredit.report import OverduePredictor, ModelReport, auto_model_report, auto_feature_analysis
from hscredit.report.result import ReportGenerationError, ReportResult


@pytest.mark.parametrize("direction", ["high", "low"])
def test_strategy_prefix_matches_independent_masks(direction):
    data = pd.DataFrame({"score": [100, 100, 200, 300], "y": [0, 1, 1, 0], "amount": ["100", "200", "300", "400"]})
    thresholds = [100, 300, 0, 100, 999]
    actual = score_strategy_simulation(data, "score", "y", thresholds, amount_col="amount", score_low_risk=direction)
    for i, threshold in enumerate(thresholds):
        mask = data.score.ge(threshold) if direction == "high" else data.score.le(threshold)
        assert actual.loc[i, "通过量(笔)"] == int(mask.sum())
        assert actual.loc[i, "通过金额"] == data.loc[mask, "amount"].astype(float).sum()
        assert actual.loc[i, "通过人群坏账金额"] == data.loc[mask & data.y.eq(1), "amount"].astype(float).sum()
    pd.testing.assert_series_equal(actual.iloc[0], actual.iloc[3], check_names=False)
    assert data.amount.dtype == object


def test_strategy_string_amount_regression():
    result = score_strategy_simulation(pd.DataFrame({"score": [100, 200], "y": [0, 1], "amount": ["100", "200"]}), "score", "y", [0, 0], amount_col="amount")
    assert result["通过金额"].tolist() == [300.0, 300.0]
    assert result["坏账率(金额,%)"].tolist() == [66.6667, 66.6667]


def panel():
    return pd.DataFrame({"id": [1, 2, 1, 2], "v": ["2026-01"] * 4, "mob": [1, 1, 2, 2], "y": [1, 0, 1, 0]})


def test_panel_vintage_counts_unique_entities_without_readding_ever():
    kwargs = dict(entity_col="id", input_mode="panel", label_mode="ever")
    data = panel()
    result = vintage_analysis(data, "v", "mob", "y", **kwargs)
    assert result["累积坏账率(%)"].tolist() == [50.0, 50.0]
    assert result["累积坏账户数"].tolist() == [1, 1]
    summary = vintage_summary(data, "v", "mob", "y", **kwargs)
    assert summary.loc[0, "总开户数"] == 2
    future = data[data.mob.eq(2)].assign(mob=3)
    extended = vintage_analysis(pd.concat([data, future]), "v", "mob", "y", **kwargs)
    pd.testing.assert_frame_equal(result, extended.iloc[:2].reset_index(drop=True))


def test_panel_vintage_requires_contract_and_valid_history():
    with pytest.raises(ValueError, match="entity_col"):
        vintage_analysis(panel(), "v", "mob", "y", input_mode="panel", label_mode="ever")
    with pytest.raises(ValueError, match="label_mode"):
        vintage_analysis(panel(), "v", "mob", "y", input_mode="panel", entity_col="id")
    with pytest.raises(ValueError, match="重复"):
        vintage_analysis(pd.concat([panel(), panel().iloc[:1]]), "v", "mob", "y", input_mode="panel", entity_col="id", label_mode="ever")
    restored = panel()
    restored.loc[2, "y"] = 0
    with pytest.raises(ValueError, match="ever"):
        vintage_analysis(restored, "v", "mob", "y", input_mode="panel", entity_col="id", label_mode="ever")
    current = vintage_analysis(restored, "v", "mob", "y", input_mode="panel", entity_col="id", label_mode="current")
    assert current["统计坏账率(%)"].tolist() == [50.0, 0.0]
    assert current["累积坏账率(%)"].isna().all()


def test_vintage_immature_and_missing_observations_are_not_good():
    result = vintage_analysis(panel(), "v", "mob", "y", max_mob=3, input_mode="panel", entity_col="id", label_mode="ever", observation_end="2026-03-31")
    assert result.loc[result.MOB.eq(3), "状态"].iloc[0] == "未成熟"
    assert result.loc[result.MOB.eq(3), "统计坏账率(%)"].isna().all()
    incomplete = vintage_analysis(panel().iloc[:-1], "v", "mob", "y", input_mode="panel", entity_col="id", label_mode="ever")
    assert incomplete.loc[incomplete.MOB.eq(2), "状态"].iloc[0] == "观测不足"
    assert incomplete.loc[incomplete.MOB.eq(2), "统计坏账率(%)"].isna().all()


def test_vintage_event_requires_independent_cohort_and_cutoff():
    data = pd.DataFrame({"id": [1], "v": ["2026-01"], "mob": [1], "y": [1]})
    with pytest.raises(ValueError, match="cohort_sizes"):
        vintage_analysis(data, "v", "mob", "y", entity_col="id", input_mode="event", label_mode="new_default")
    result = vintage_analysis(data, "v", "mob", "y", max_mob=3, entity_col="id", input_mode="event", label_mode="new_default", cohort_sizes={"2026-01": 2}, observation_end="2026-03-31")
    assert result["统计坏账率(%)"].iloc[:3].tolist() == [0.0, 50.0, 50.0]
    assert result.iloc[3]["状态"] == "未成熟"


def test_vintage_zero_event_cohort_still_has_mature_zero_rate():
    data = pd.DataFrame(columns=["id", "v", "mob", "y"])
    result = vintage_analysis(data, "v", "mob", "y", max_mob=2, entity_col="id", input_mode="event", label_mode="new_default", cohort_sizes={"2026-01": 100}, observation_end="2026-02-28")
    assert result["统计坏账率(%)"].iloc[:2].tolist() == [0.0, 0.0]
    assert result.iloc[2]["状态"] == "未成熟"


@pytest.mark.parametrize("mode", ["panel", "snapshot"])
def test_vintage_exposure_is_same_period_and_conserved(mode):
    data = panel().assign(amount=[100.0, 300.0, 90.0, 210.0])
    if mode == "snapshot":
        data = data.iloc[:2]
    result = vintage_analysis(data, "v", "mob", "y", entity_col="id", input_mode=mode, label_mode="ever", exposure_col="amount")
    np.testing.assert_allclose(result["有效金额分母"], result["好账户金额"] + result["坏账金额"])
    assert result.iloc[0]["金额坏账率(%)"] == 25.0
    assert result.iloc[0]["累积坏账率(%)"] == 50.0
    if mode == "panel":
        assert result.iloc[1]["有效金额分母"] == 300.0
        assert result.iloc[1]["金额坏账率(%)"] == 30.0
    with pytest.raises(ValueError, match="非缺失"):
        vintage_analysis(data.assign(amount=np.nan), "v", "mob", "y", entity_col="id", input_mode=mode, label_mode="ever", exposure_col="amount")


def reference_table():
    return pd.DataFrame({"分箱标签": ["[-inf, 0)", "[0, inf)", "合计"], "坏样本率": [0.1, 0.3, 0.12], "样本总数": [90, 10, 100]})


def test_overdue_auto_preserves_weighted_baseline_and_unknown_fallback():
    predictor = OverduePredictor("x", coefficients="auto", n_jobs=1).fit(reference_table())
    result = predictor.predict(pd.DataFrame({"x": [-1.0, 0.0, 1.0, np.nan]}))
    np.testing.assert_allclose(result, [0.1, 0.3, 0.3, 0.12])
    assert predictor.bin_weights_["_default"] == {"[-inf, 0)": 90.0, "[0, inf)": 10.0}
    categorical = reference_table().copy()
    categorical["分箱标签"] = ["A", "B", "合计"]
    other = OverduePredictor("x", n_jobs=1).fit(categorical).predict(pd.DataFrame({"x": ["A", "B", "new"]}))
    np.testing.assert_allclose(other, [0.1, 0.3, 0.12])


def test_overdue_auto_rejects_missing_weights_and_accounts_for_clipping():
    predictor = OverduePredictor("x", coefficients="auto", n_jobs=1).fit(reference_table().drop(columns="样本总数"))
    with pytest.raises(ValueError, match="样本总数"):
        predictor.predict(pd.DataFrame({"x": [-1]}))
    table = reference_table()
    table["样本总数"] = [50, 50, 100]
    table["坏样本率"] = [0.1, 0.9, 0.8]
    predictions = OverduePredictor("x", coefficients="auto", n_jobs=1).fit(table).predict(pd.DataFrame({"x": [-1, 1]}))
    assert predictions.mean() == pytest.approx(0.8)
    assert predictions.max() <= 1


def eda_data():
    return pd.DataFrame({"x": [1, 2, 3, 4], "y": [0, 1, 0, 1], "date": pd.to_datetime(["2026-01-01", "2026-01-02", "2026-02-01", "2026-02-02"])})


@pytest.mark.parametrize("function,key", [(eda_summary, "逾期率趋势"), (generate_report, "7.逾期率趋势")])
def test_eda_trend_sections_are_present_with_valid_dates(function, key):
    result = function(eda_data(), target="y", date_col="date", features=["x"], n_jobs=1, mode="strict", return_result=True)
    assert isinstance(result, dict)
    assert isinstance(result, ReportResult)
    assert len(result[key]) == 2
    assert result.sections[key].status == "成功"
    assert result.complete


def test_eda_best_effort_exposes_failure_strict_raises_and_preserves_status():
    data = eda_data().assign(date="invalid")
    result = eda_summary(data, target="y", date_col="date", features=["x"], n_jobs=1, return_result=True)
    assert not result.complete
    assert result.sections["逾期率趋势"].status == "失败"
    assert result.sections["逾期率趋势"].reason
    with pytest.raises(ReportGenerationError) as error:
        eda_summary(data, target="y", date_col="date", features=["x"], n_jobs=1, mode="strict")
    assert error.value.result.sections["逾期率趋势"].status == "失败"


def test_export_sections_get_unique_case_insensitive_clean_names(tmp_path):
    output = tmp_path / "report.xlsx"
    export_report_to_excel({"A/B": pd.DataFrame({"value": [1, 2]}), "A:B": pd.DataFrame({"value": [9]}), "a_b": pd.DataFrame({"value": [8]})}, str(output))
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["A_B", "A_B_2", "a_b_3"]
    assert workbook["A_B"]["B6"].value == 2
    assert workbook["A_B_2"]["B5"].value == 9
    assert workbook["A_B_2"]["B6"].value is None
    workbook.close()


def test_excel_exception_and_xml_failure_preserve_old_file(tmp_path, monkeypatch):
    output = tmp_path / "report.xlsx"
    writer = ExcelWriter()
    writer.insert_value2sheet(writer.get_sheet_by_name("old"), "A1", "old")
    writer.save(str(output))
    before = hashlib.sha256(output.read_bytes()).digest()
    with pytest.raises(RuntimeError):
        with ExcelWriter().set_filename(str(output)) as writer:
            writer.insert_value2sheet(writer.get_sheet_by_name("new"), "A1", "new")
            raise RuntimeError("受控计算失败")
    assert hashlib.sha256(output.read_bytes()).digest() == before
    writer = ExcelWriter()
    writer.insert_value2sheet(writer.get_sheet_by_name("new"), "A1", "new")
    def fail(_):
        raise RuntimeError("受控XML注入失败")
    monkeypatch.setattr(writer, "_inject_pivots", fail)
    with pytest.raises(RuntimeError):
        writer.save(str(output))
    assert hashlib.sha256(output.read_bytes()).digest() == before
    assert list(tmp_path.glob(".hscredit-*.xlsx")) == []


def test_report_export_discloses_partial_status(tmp_path):
    result = eda_summary(eda_data().assign(date="invalid"), target="y", date_col="date", features=["x"], n_jobs=1, return_result=True)
    output = tmp_path / "partial.xlsx"
    export_report_to_excel(result, str(output))
    workbook = load_workbook(output)
    assert workbook.sheetnames[0] == "报告执行状态"
    assert "不完整" in workbook.worksheets[0]["B2"].value
    assert result.artifacts[0]["完整"] is False
    assert result.sheet_mapping
    workbook.close()


@pytest.mark.parametrize("anchor", ["XFD1", "A1048576"])
def test_excel_shape_preflight_happens_before_cell_mutation(anchor):
    writer = ExcelWriter()
    worksheet = writer.get_sheet_by_name("预检")
    before = len(worksheet._cells)
    with pytest.raises(ValueError, match="工作表上限"):
        writer.insert_df2sheet(worksheet, pd.DataFrame({"a": [1], "b": [2]}), anchor)
    assert len(worksheet._cells) == before
    writer.workbook.close()


def test_excel_explicit_cell_budget_and_aggregated_feature_results():
    writer = ExcelWriter(max_table_cells=3)
    worksheet = writer.get_sheet_by_name("预检")
    with pytest.raises(ValueError, match="max_table_cells"):
        writer.insert_df2sheet(worksheet, pd.DataFrame({"a": [1, 2], "b": [3, 4]}), "B2")
    writer.workbook.close()
    from hscredit.report.feature_analyzer import _auto_feature_compute_call
    data = pd.DataFrame({"x": np.tile(np.arange(10), 100), "y": np.tile([0, 1], 500)})
    result = _auto_feature_compute_call((data, "x", "y", None, None, False, False, None, "x", False, {"method": "quantile", "n_jobs": 1}))
    assert "data" not in result
    assert len(result["sample_table"]) < 20


def small_model_report():
    from sklearn.linear_model import LogisticRegression
    data = eda_data()
    model = LogisticRegression().fit(data[["x"]], data.y)
    return ModelReport(model, X_train=data[["x"]], y_train=data.y, n_jobs=1)


def test_model_report_required_failure_is_structured_and_preserves_old_output(tmp_path, monkeypatch):
    report = small_model_report()
    output = tmp_path / "model.xlsx"
    output.write_bytes(b"previous-success")
    def fail(**kwargs):
        raise RuntimeError("受控核心计算失败")
    monkeypatch.setattr(report, "_precompute_excel_tables", fail)
    with pytest.raises(ReportGenerationError) as error:
        report.to_excel(str(output), mode="strict", with_plots=False)
    assert error.value.result.sections["核心数据计算"].status == "失败"
    assert output.read_bytes() == b"previous-success"
    result = report.to_excel(str(output), mode="best_effort", return_result=True, with_plots=False)
    assert not result.complete
    assert result.artifacts == []


def test_model_report_swallowed_plot_failure_blocks_strict_publication(tmp_path, monkeypatch):
    report = small_model_report()
    output = tmp_path / "model.xlsx"
    output.write_bytes(b"previous-success")
    def partial_plots(*args, **kwargs):
        report._report_warning("生成模型 KS 图失败: %s", "受控错误")
        return {}, {}
    monkeypatch.setattr(report, "_export_plots", partial_plots)
    with pytest.raises(ReportGenerationError) as error:
        report.to_excel(str(output), mode="strict", with_plots=True)
    assert any("KS" in s.reason for s in error.value.result.sections.values())
    assert output.read_bytes() == b"previous-success"
    result = report.to_excel(str(output), mode="best_effort", return_result=True, with_plots=True)
    assert not result.complete
    assert result.artifacts and result.artifacts[0]["完整"] is False
    workbook = load_workbook(output)
    assert workbook.sheetnames[0] == "报告执行状态"
    assert "不完整" in workbook.worksheets[0]["B2"].value
    workbook.close()


def test_auto_feature_and_model_report_protocol_exports(tmp_path):
    output = tmp_path / "features.xlsx"
    result = auto_feature_analysis(eda_data(), features=["x"], target="y", excel_writer=str(output), pictures=[], output_dir=str(tmp_path / "assets"), n_jobs=1, show_progress=False, mode="strict", return_result=True)
    assert isinstance(result, ReportResult) and result.complete
    assert result.artifacts and "特征分箱：x" in result.sections
    report = small_model_report()
    model_result = auto_model_report(report.model, X_train=eda_data()[["x"]], y_train=eda_data().y, verbose=False, n_jobs=1, return_result=True)
    assert isinstance(model_result, ReportResult) and model_result.complete


def test_auto_feature_fault_injection_reports_no_artifact(tmp_path, monkeypatch):
    import hscredit.report.feature_analyzer as analyzer
    output = tmp_path / "features.xlsx"
    output.write_bytes(b"previous-success")
    def fail(task):
        raise RuntimeError("受控特征分箱失败")
    monkeypatch.setattr(analyzer, "_auto_feature_compute_call", fail)
    kwargs = dict(features=["x"], target="y", excel_writer=str(output), pictures=[], output_dir=str(tmp_path / "assets"), n_jobs=1, show_progress=False)
    with pytest.raises(ReportGenerationError):
        auto_feature_analysis(eda_data(), mode="strict", **kwargs)
    result = auto_feature_analysis(eda_data(), return_result=True, **kwargs)
    assert result.sections["输入与分箱计算"].status == "失败"
    assert "受控特征分箱失败" in result.sections["输入与分箱计算"].reason
    assert not result.artifacts
    assert output.read_bytes() == b"previous-success"


def test_invalid_frequency_stays_failed_independently_of_fallback_table(tmp_path):
    output = tmp_path / "frequency.xlsx"
    output.write_bytes(b"previous-success")
    with pytest.raises(ReportGenerationError) as error:
        auto_feature_analysis(eda_data(), features=["x"], target="y", date="date", freq="not-a-frequency",
                              excel_writer=str(output), pictures=[], output_dir=str(tmp_path / "images"),
                              n_jobs=1, show_progress=False, mode="strict", return_result=True)
    result = error.value.result
    assert result.sections["时间频率回退"].status == "失败"
    assert result.sections["时间分布"].status == "成功"
    assert not result["时间分布"].empty
    assert not result.complete
    assert output.read_bytes() == b"previous-success"
