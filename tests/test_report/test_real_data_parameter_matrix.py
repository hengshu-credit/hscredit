"""使用真实放款数据验证各分析模块的统一参数及标签输入矩阵。"""

from functools import partial
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler

from hscredit.core.eda.target import bad_rate_overall, bad_rate_by_dimension, bad_rate_trend, bad_rate_by_bins
from hscredit.core.rules import Rule
from hscredit.core.viz import distribution_plot, bin_overdues_plot, bad_rate_trend_plot
from hscredit.report import (
    feature_bin_stats,
    feature_bin_stats_2d,
    feature_binning_summary,
    feature_group_binning_summary,
    auto_feature_analysis,
    auto_model_report,
    swap_out_report,
    ModelReport,
    OverduePredictor,
)
from hscredit.report.mining import DecisionTreeAnalyzer
from hscredit.utils.overdue import make_overdue_target

FEATURES = ["衡枢鉴真分老客版", "近六个月非银多头机构数", "青云24"]
TARGET, DATE, AMOUNT, GROUP = "核验标签", "放款时间", "放款金额", "商品类别"
KINDS = ["column", "series", "array", "list", "lambda", "partial", "dict_values", "dict_callable"]
ALL_KINDS = KINDS + ["overdue_gt", "overdue_ge", "overdue_lt", "overdue_le"]


@pytest.fixture(scope="module")
def loans():
    path = Path(__file__).resolve().parents[2] / "examples/hscredit_yyp.xlsx"
    if not path.exists():
        pytest.skip("缺少 examples/hscredit_yyp.xlsx")
    data = pd.read_excel(path)
    data = data.loc[:, FEATURES + ["FPD", "MOB1", DATE, AMOUNT, GROUP]].copy()
    assert len(data) == 970
    data[TARGET] = data.FPD.astype(float)
    # 仅在独立验证标签中注入灰客户，原始 FPD/逾期天数不变。
    data.loc[data.index[::11], TARGET] = -1
    data.loc[data.index[::13], TARGET] = np.nan
    return data


@pytest.fixture(scope="module")
def model(loans):
    return make_pipeline(SimpleImputer(), StandardScaler(), LogisticRegression(max_iter=500)).fit(
        loans[FEATURES], loans.FPD
    )


def _column(data, name):
    return data[name]


def _inputs(data, kind, grey, *, paired=False):
    if kind.startswith("overdue"):
        op = {"overdue_gt": ">", "overdue_ge": ">=", "overdue_lt": "<", "overdue_le": "<="}[kind]
        expected = [make_overdue_target(data.MOB1, d, grey, op) for d in (7, 3, 0)]
        return (
            dict(overdue=["MOB1"] * 3 if paired else ["MOB1"], dpds=[7, 3, 0], overdue_operator=op, del_grey=grey),
            expected,
        )
    raw = data[TARGET]
    second = data.MOB1.gt(0)
    values = {
        "column": TARGET,
        "series": raw.iloc[::-1],
        "array": raw.to_numpy(),
        "list": raw.tolist(),
        "lambda": lambda df: df[TARGET],
        "partial": partial(_column, name=TARGET),
        "dict_values": {"验证_7+": raw.iloc[::-1], "第二@7": second.to_numpy()},
        "dict_callable": {"验证_7+": lambda df: df[TARGET], "第二@7": lambda df: df.MOB1.gt(0)},
    }
    sources = [raw, second] if kind.startswith("dict") else [raw]
    expected = [s.eq(1).astype(float).where(s.isin([0, 1])) if grey else s.eq(1).astype(float) for s in sources]
    return dict(target=values[kind], del_grey=grey), expected


def _check_table(table, expected, amount=None):
    if isinstance(table.columns, pd.MultiIndex):
        groups = list(dict.fromkeys(c[0] for c in table.columns if c[-1] == "坏样本数"))
        actual = []
        for group in groups:
            total = (group, "样本总数") if (group, "样本总数") in table else ("分箱详情", "样本总数")
            actual.append((table.iloc[-1][total], table.iloc[-1][(group, "坏样本数")]))
    elif "逾期标签" in table:
        actual = [(t.iloc[-1]["样本总数"], t.iloc[-1]["坏样本数"]) for _, t in table.groupby("逾期标签", sort=False)]
    else:
        actual = [(table.iloc[-1]["样本总数"], table.iloc[-1]["坏样本数"])]
    assert len(actual) == len(expected)
    for (total, bad), y in zip(actual, expected):
        weight = pd.Series(1.0, index=y.index) if amount is None else amount
        assert total == pytest.approx(weight[y.notna()].sum())
        assert bad == pytest.approx(weight[y.eq(1)].sum())


@pytest.mark.parametrize("grey", [False, True])
@pytest.mark.parametrize("kind", ALL_KINDS)
@pytest.mark.parametrize(
    "entry",
    [
        "bin",
        "category_bin",
        "multi_bin",
        "2d",
        "summary",
        "rule",
        "predictor",
        "model",
        "eda",
        "eda_group",
        "eda_trend",
        "eda_bins",
    ],
)
def test_real_data_modules_count_and_amount_matrix(loans, model, entry, kind, grey):
    options, expected = _inputs(loans, kind, grey)
    shared = dict(method="quantile", max_n_bins=3, n_jobs=1, margins=True, amount=loans[AMOUNT].iloc[::-1])
    if entry == "multi_bin":
        table = feature_bin_stats(loans, features=np.array(FEATURES), long_format=True, **options, **shared)
        assert set(table["指标名称"]) == set(FEATURES)
        for _, part in table.groupby("指标名称", sort=False):
            _check_table(part, expected, loans[AMOUNT])
        return
    if entry == "category_bin":
        table = feature_bin_stats(loans, features=GROUP, long_format=True, **options, **shared)
    elif entry == "bin":
        table = feature_bin_stats(loans, features=np.array([FEATURES[0]]), long_format=True, **options, **shared)
    elif entry == "2d":
        table = feature_bin_stats_2d(loans, x1=FEATURES[0], x2=FEATURES[1], long_format=True, **options, **shared)
    elif entry == "summary":
        tables, _ = feature_binning_summary(
            loans,
            features=FEATURES[0],
            methods="quantile",
            **options,
            **{k: v for k, v in shared.items() if k != "method"},
        )
        table = tables[FEATURES[0]]["quantile"]
    elif entry == "rule":
        table = Rule(f"`{FEATURES[0]}` > 0.5").report(loans, amount=loans[AMOUNT], margins=True, **options)
    elif entry == "predictor":
        predictor = OverduePredictor(
            FEATURES[0], method="quantile", max_n_bins=3, bin_params={"amount": loans[AMOUNT]}, n_jobs=1, **options
        ).fit(loans)
        table = predictor.bin_table_
        assert len(predictor.transform(loans)) == len(loans)
    elif entry == "model":
        report = ModelReport(model, datasets={"样本": loans}, features=np.array(FEATURES), n_jobs=1, **options)
        labels = report._label_names or [None]
        assert len(labels) == len(expected)
        for label, y in zip(labels, expected):
            _check_table(
                report.get_bin_table("样本", label=label, max_n_bins=3, amount=loans[AMOUNT]), [y], loans[AMOUNT]
            )
            _check_table(
                report.get_feature_bin_table(FEATURES[0], "样本", label=label, max_n_bins=3, amount=loans[AMOUNT]),
                [y],
                loans[AMOUNT],
            )
        return
    else:
        if entry == "eda":
            result = bad_rate_overall(loans, **options)
            records = result.to_dict("records") if isinstance(result, pd.DataFrame) else [result]
            assert [r["样本总数"] for r in records] == [y.notna().sum() for y in expected]
            assert [r["坏样本数"] for r in records] == [y.sum() for y in expected]
            return
        if entry == "eda_group":
            result = bad_rate_by_dimension(loans, GROUP, **options)
            expected = [y.where(loans[GROUP].notna()) for y in expected]
        elif entry == "eda_trend":
            result = bad_rate_trend(loans, date=loans[DATE], freq="MS", **options)
        else:
            result = bad_rate_by_bins(loans, FEATURES[0], max_n_bins=3, **options)
        tables = list(result.values()) if isinstance(result, dict) else [result]
        assert len(tables) == len(expected)
        for table, y in zip(tables, expected):
            assert table["样本数"].sum() == y.notna().sum()
            assert table["坏样本数"].sum() == y.sum()
        return
    _check_table(table, expected, loans[AMOUNT])


@pytest.mark.parametrize("grey", [False, True])
@pytest.mark.parametrize("kind", ALL_KINDS)
@pytest.mark.parametrize("entry", ["feature", "model", "strategy"])
def test_real_data_report_export_matrix(loans, model, tmp_path, entry, kind, grey):
    options, _ = _inputs(loans, kind, grey)
    path = tmp_path / f"{entry}-{kind}-{grey}.xlsx"
    if entry == "feature":
        report = auto_feature_analysis(
            loans,
            features=FEATURES[0],
            date=loans[DATE],
            freq="MS",
            amount=loans[AMOUNT],
            pictures=[],
            show_progress=False,
            bin_params={"method": "quantile", "max_n_bins": 3},
            n_jobs=1,
            excel_writer=str(path),
            output_dir=str(tmp_path),
            mode="strict",
            return_result=True,
            **options,
        )
        assert report.complete
    elif entry == "model":
        report = auto_model_report(
            model,
            datasets={"样本": loans},
            features=FEATURES,
            date=loans[DATE],
            freq="MS",
            group=loans[GROUP],
            amount=loans[AMOUNT],
            max_n_bins=3,
            method="quantile",
            n_jobs=1,
            verbose=False,
            with_plots=False,
            excel_path=str(path),
            include_sample_records=False,
            mode="strict",
            return_result=True,
            **options,
        )
        assert report.complete
    else:
        writer = swap_out_report(
            loans,
            Rule(f"`{FEATURES[0]}` > 0.5"),
            features=FEATURES[0],
            methods="quantile",
            bin_params={"max_n_bins": 3},
            date=loans[DATE],
            freq="MS",
            amount=loans[AMOUNT],
            n_jobs=1,
            save=str(path),
            **options,
        )
        writer.workbook.close()
    book = load_workbook(path, read_only=True)
    cells = [c.value for sheet in book for row in sheet for c in row]
    assert book.sheetnames and any(isinstance(v, (int, float)) for v in cells)
    assert not any(isinstance(v, str) and v in {"#REF!", "#DIV/0!", "#VALUE!"} for v in cells)
    if kind.startswith("dict"):
        assert "验证_7+" in cells and "第二@7" in cells
    book.close()
    plt.close("all")


@pytest.mark.parametrize("grey", [False, True])
@pytest.mark.parametrize("kind", ["column", "array", "dict_values", "dict_callable", "overdue_gt"])
@pytest.mark.parametrize("entry", ["auto_feature_analysis", "auto_model_report", "swap_out_report"])
def test_real_data_skill_entry_matrix(loans, model, tmp_path, entry, kind, grey):
    from hscredit.skills_runtime import execute_skill

    options, _ = _inputs(loans, kind, grey)
    objects = {"data": loans, "model": model, "date": loans[DATE], "amount": loans[AMOUNT]}
    params = dict(features=FEATURES[0], n_jobs=1, freq="MS", **options)
    if "target" in params:
        objects["target"] = params.pop("target")
    if entry == "auto_feature_analysis":
        params.update(pictures=[], show_progress=False, bin_params={"method": "quantile", "max_n_bins": 3})
    elif entry == "auto_model_report":
        params.update(
            datasets={"样本": "data"},
            features=FEATURES,
            max_n_bins=3,
            with_plots=False,
            verbose=False,
            include_sample_records=False,
        )
    else:
        params.update(rules=[f"`{FEATURES[0]}` > 0.5"], methods="quantile", bin_params={"max_n_bins": 3})
    result = execute_skill(
        "hsreport",
        {
            "version": "1",
            "operation": entry,
            "inputs": {name: {"kind": "object_ref", "ref": name} for name in objects},
            "parameters": params,
            "output": {"directory": str(tmp_path), "name": entry, "overwrite": False},
            "environment": {"mode": "current", "install_missing": False},
        },
        objects=objects,
    )
    assert result["status"] == "success"
    assert any(a["type"] == "excel" and Path(a["path"]).is_file() for a in result["artifacts"])


@pytest.mark.parametrize("grey", [False, True])
@pytest.mark.parametrize("kind", KINDS + ["overdue_gt"])
@pytest.mark.parametrize(
    "entry", ["group", "ruleset", "rule_group", "tree", "efficiency", "benchmark", "swap", "plots"]
)
def test_real_data_composite_modules(loans, tmp_path, entry, kind, grey):
    from hscredit.report.feature_analyzer import feature_efficiency_analysis, benchmark_binning_methods
    from hscredit.report.rule_analysis import ruleset_analysis, rule_swap_analysis
    from hscredit.report.rule_strategy import rule_group_compare

    options, expected = _inputs(loans, kind, grey)
    rule = Rule(f"`{FEATURES[0]}` > 0.5")
    if entry == "group":
        tables, summary = feature_group_binning_summary(
            loans,
            features=FEATURES[0],
            methods="quantile",
            date=loans[DATE],
            freq="MS",
            amount=loans[AMOUNT],
            max_n_bins=3,
            margins=True,
            long_format=True,
            n_jobs=1,
            **options,
        )
        merged = pd.concat(list(tables[FEATURES[0]]["quantile"].values()), ignore_index=True)
        totals = (
            merged.loc[merged["分箱标签"].eq("合计")]
            .groupby("逾期标签", sort=False)[["样本总数", "坏样本数"]]
            .sum()
            .reset_index()
        )
        _check_table(totals, expected, loans[AMOUNT])
        assert not summary.empty
    elif entry == "ruleset":
        first = Rule(f"`{FEATURES[0]}` > {loans[FEATURES[0]].median()}")
        result = ruleset_analysis(loans, [first, rule], amount=loans[AMOUNT], n_jobs=1, **options)
        assert not result.empty
    elif entry == "rule_group":
        result = rule_group_compare(loans, rule, group=loans[GROUP], amount=loans[AMOUNT], n_jobs=1, **options)
        assert not result.empty
    elif entry == "tree":
        tree = DecisionTreeAnalyzer(target="FPD", features=FEATURES, tree_params={"max_depth": 2}, n_jobs=1).fit(loans)
        result = tree.report(loans, **options)
        assert not result.empty
    elif entry == "efficiency":
        try:
            result = feature_efficiency_analysis(
                loans,
                FEATURES[0],
                method="quantile",
                max_n_bins=3,
                amount=loans[AMOUNT],
                date=loans[DATE],
                freq="MS",
                margins=True,
                n_jobs=1,
                output_dir=str(tmp_path),
                **options,
            )
            _check_table(result["auto_table"], expected, loans[AMOUNT])
        finally:
            plt.close("all")
    elif entry == "benchmark":
        result = benchmark_binning_methods(loans, FEATURES[0], methods=["quantile"], n_jobs=1, **options)
        errors = [column for column in result if isinstance(column, tuple) and column[-1] == "错误信息"]
        assert errors and result[errors].isna().all().all()
    elif entry == "swap":
        working = loans
        if not grey and not kind.startswith("overdue"):
            # 置换业务保留未表现样本语义，非 OUT-IN 的缺失表现必须明确报错。
            with pytest.raises(ValueError, match="缺少实际表现"):
                rule_swap_analysis(
                    loans,
                    score=FEATURES[0],
                    reference_data=loans,
                    rules_out=[rule],
                    method="quantile",
                    max_n_bins=3,
                    n_jobs=1,
                    **options,
                )
            working = loans.assign(**{TARGET: loans[TARGET].fillna(0)})
            options, _ = _inputs(working, kind, grey)
        result = rule_swap_analysis(
            working,
            score=FEATURES[0],
            reference_data=working,
            rules_out=[rule],
            method="quantile",
            max_n_bins=3,
            amount=loans[AMOUNT],
            n_jobs=1,
            **options,
        )
        assert not result["swap_pipeline"].empty
    else:
        try:
            paired_options, _ = _inputs(loans, kind, grey, paired=True)
            frame = distribution_plot(loans, date=loans[DATE], freq="MS", result=True, **paired_options)
            assert not frame.empty and frame.iloc[:, 0].is_unique
            figures = [
                bin_overdues_plot(
                    loans, feature=FEATURES[0], method="quantile", max_n_bins=3, n_jobs=1, **paired_options
                ),
                bad_rate_trend_plot(loans, date=loans[DATE], freq="MS", **options),
            ]
            for figure in figures:
                assert figure.axes
                assert not any("Error:" in text.get_text() for ax in figure.axes for text in ax.texts)
        finally:
            plt.close("all")
