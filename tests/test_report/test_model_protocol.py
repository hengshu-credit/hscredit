"""报告标签口径、默认推理方法及概率转评分公式的独立核验。"""

import numpy as np
import pandas as pd
import pytest
from numpy.testing import assert_allclose
from openpyxl import load_workbook
from sklearn.pipeline import Pipeline
from sklearn.metrics import roc_auc_score

from hscredit import ScoreTransformer, LogisticRegression
from hscredit.report import ModelReport
from hscredit.report._model_protocol import label_definition_table


class ProbabilityModel:
    feature_names_in_ = np.array(["概率"])
    classes_ = np.array([0, 1])
    _estimator_type = "classifier"

    def predict_proba(self, X):
        p = np.asarray(X["概率"])
        return np.column_stack((1 - p, p))

    def predict(self, X):
        return np.zeros(len(X))

    def custom_score(self, X, multiplier=1):
        return np.asarray(X["概率"]) * multiplier


@pytest.fixture
def data():
    return pd.DataFrame({"概率": np.linspace(0.05, 0.95, 12), "MOB1": [0, 0, 1, 2, 3, 3, 4, 7, 8, 10, np.nan, 0]})


@pytest.mark.parametrize("operator", [">", ">=", "<", "<="])
@pytest.mark.parametrize("del_grey", [False, True])
def test_overdue_labels_and_basic_definitions_match_independent_rows(operator, del_grey, data, tmp_path):
    report = ModelReport(
        ProbabilityModel(),
        datasets={"样本": data},
        overdue=["MOB1"],
        dpds=[0, 3],
        overdue_operator=operator,
        del_grey=del_grey,
        n_jobs=1,
    )
    assert report._label_names == ["MOB1@0", "MOB1@3"]
    table = label_definition_table(report).set_index("标签")
    assert report.summary().index.tolist() == ["MOB1@0", "MOB1@3"]
    for threshold in [0, 3]:
        days = data.MOB1.to_numpy()
        compare = {">": np.greater, ">=": np.greater_equal, "<": np.less, "<=": np.less_equal}[operator]
        expected = compare(days, threshold).astype(float)
        grey = (
            ((days > 0) & (days <= threshold))
            if operator == ">"
            else ((days > 0) & (days < threshold)) if operator == ">=" else np.zeros(len(days), dtype=bool)
        )
        if del_grey:
            expected[grey] = np.nan
        assert_allclose(report._get_y("样本", f"MOB1@{threshold}"), expected, equal_nan=True)
        assert_allclose(report._get_y("样本", f"MOB1{operator}{threshold}"), expected, equal_nan=True)
        assert table.loc[f"MOB1@{threshold}", "操作符"] == operator
        assert table.loc[f"MOB1@{threshold}", "剔灰配置"] == ("是" if del_grey else "否")
        assert table.loc[f"MOB1@{threshold}", "有效样本数"] == np.isfinite(expected).sum()
    path = tmp_path / "模型报告.xlsx"
    report.to_excel(path, with_plots=False, include_sample_records=False, mode="strict")
    book = load_workbook(path)
    basic = [c.value for row in book["1-基本信息"] for c in row]
    assert {"配置项", "配置值", "标签来源", "操作符", "灰样本定义", "剔灰配置", "MOB1@3"} <= set(
        v for v in basic if isinstance(v, str)
    )
    assert "标签定义与剔灰规则" not in basic and "实际剔除规则" not in basic
    performance = [c.value for row in book["2-模型性能"] for c in row]
    assert "MOB1@3" in performance
    assert not any(
        c.value in {"MOB1>3", "MOB1>=3", "MOB1<3", "MOB1<=3", "MOB1大于3"}
        for sheet in book
        for row in sheet
        for c in row
    )
    book.close()


def test_explicit_target_names_and_external_y_are_described_honestly(data):
    report = ModelReport(
        ProbabilityModel(),
        datasets={"派生": data, "外部": (data, np.arange(len(data)) % 2)},
        target={"MOB1大于3": lambda d: d.MOB1.gt(3)},
        del_grey=True,
        n_jobs=1,
    )
    assert report._label_names == ["MOB1@3"]
    definitions = label_definition_table(report).set_index("数据集")
    assert definitions.loc["外部", "标签来源"] == "显式y"
    assert definitions.loc["外部", "实际剔除规则"] == "外部 y 优先，按原值使用"
    assert definitions.loc["派生", "操作符"] == "不适用（直接标签）"
    assert_allclose(report._get_y("派生", "MOB1大于3"), data.MOB1.gt(3))
    with pytest.raises(ValueError, match="标签展示名重复"):
        ModelReport(
            ProbabilityModel(),
            datasets={"样本": data},
            target={"MOB1>3": lambda d: d.MOB1.gt(3), "MOB1@3": lambda d: d.MOB1.gt(3)},
            n_jobs=1,
        )


@pytest.mark.parametrize("method", ["standard", "linear", "quantile", "boxcox"])
@pytest.mark.parametrize("direction", ["ascending", "descending"])
def test_score_transformer_defaults_to_transform_and_reports_actual_formula(method, direction, tmp_path):
    options = {"n_quantiles": 5} if method == "quantile" else {"lmbda": 0.5} if method == "boxcox" else {}
    p = np.array([0, 0.01, 0.1, 0.3, 0.6, 0.8, 0.95, 1.0])
    y = (p > 0.4).astype(int)
    transformer = ScoreTransformer(method=method, lower=450, upper=550, direction=direction, decimal=0, **options).fit(
        p
    )
    X = pd.DataFrame({"原始概率": p})
    report = ModelReport(transformer, datasets={"验证": (X, y)}, n_jobs=1)
    assert report.prediction_method == "transform"
    actual = transformer.transformer_
    if method == "standard":
        bounded = np.clip(p, np.finfo(float).eps, 1 - np.finfo(float).eps)
        expected = actual.A_ - actual.B_ * np.log(bounded / (1 - bounded))
    else:
        if method == "linear":
            u = p
        elif method == "quantile":
            ranks = (p[:, None] >= actual.quantile_values_[None, :]).sum(axis=1) - 1
            u = np.clip(ranks, 0, 4) / 4
        else:
            bounded = np.clip(p, 1e-10, 1 - 1e-10)
            z = ((bounded / (1 - bounded) + actual.shift) ** 0.5 - 1) / 0.5
            u = np.clip((z - actual.z_min_) / (actual.z_max_ - actual.z_min_), 0, 1)
        expected = 550 - 100 * u if direction == "descending" else 450 + 100 * u
    assert_allclose(report._datasets["验证"].prediction, expected)
    assert report.score_direction == ("higher_safe" if direction == "descending" else "higher_risk")
    assert_allclose(
        report.get_metrics().set_index("统计项").loc["AUC", "验证"],
        roc_auc_score(y, -expected if direction == "descending" else expected),
    )
    sections = report._get_score_conversion_sections()
    formula = sections["formula"].set_index("公式项")["公式内容"]
    assert "Score =" in formula["公式"] and "transform" in formula["推理输出"]
    if method == "quantile":
        assert_allclose(sections["knots"]["分位点概率"], actual.quantile_values_)
    path = tmp_path / "评分报告.xlsx"
    report.to_excel(path, with_plots=False, include_sample_records=False, mode="strict")
    book = load_workbook(path)
    values = [c.value for row in book["5-模型参数"] for c in row]
    assert any(isinstance(v, str) and v.endswith("、概率转评分公式") for v in values)
    assert formula["公式"] in values
    if method == "quantile":
        assert "分位点概率" in values
    book.close()


@pytest.mark.parametrize("n_jobs", [1, 2])
def test_classifier_default_probability_and_explicit_inference_kwargs(data, n_jobs):
    y = data.MOB1.gt(3).astype(int)
    default = ModelReport(ProbabilityModel(), datasets={"验证": (data, y)}, n_jobs=1)
    assert_allclose(default._datasets["验证"].prediction, data["概率"])
    assert default._get_score_conversion_sections() is None
    custom = ModelReport(
        ProbabilityModel(),
        datasets={"验证": (data, y), "验证2": (data, y)},
        prediction_method="custom_score",
        prediction_kwargs={"multiplier": 100},
        n_jobs=n_jobs,
        parallel_backend="threading",
    )
    assert_allclose(custom._datasets["验证"].prediction, data["概率"] * 100)
    assert_allclose(custom._datasets["验证2"].prediction, data["概率"] * 100)
    custom.add_dataset("追加", "追加", data, y)
    assert_allclose(custom._datasets["追加"].prediction, data["概率"] * 100)


def test_transformer_pipeline_direction_and_explicit_predict():
    p = np.linspace(0.001, 0.999, 20)
    X = pd.DataFrame({"概率": p})
    y = (p > 0.5).astype(int)
    transformer = ScoreTransformer(lower=450, upper=550, direction="descending").fit(p)
    pipeline = Pipeline([("评分转换", transformer)])
    report = ModelReport(pipeline, datasets={"样本": (X, y)}, n_jobs=1)
    assert report.prediction_method == "transform" and report.score_direction == "higher_safe"
    assert_allclose(report._datasets["样本"].prediction, transformer.transform(p))
    explicit = ModelReport(pipeline, datasets={"样本": (X, y)}, prediction_method="predict", n_jobs=1)
    assert_allclose(explicit._datasets["样本"].prediction, transformer.predict(p))
    model = LogisticRegression(max_iter=100).fit(X, y)
    probability_report = ModelReport(model, datasets={"样本": (X, y)}, n_jobs=1)
    assert probability_report.prediction_method == "predict_proba"


def test_single_overdue_basic_description_uses_effective_rows(data, tmp_path):
    report = ModelReport(ProbabilityModel(), datasets={"样本": data}, overdue="MOB1", dpds=3, del_grey=True, n_jobs=1)
    path = tmp_path / "单标签剔灰.xlsx"
    report.to_excel(path, with_plots=False, include_sample_records=False, mode="strict")
    book = load_workbook(path)
    basic = book["1-基本信息"]
    contents = {basic.cell(row, 2).value: basic.cell(row, 3).value for row in range(1, basic.max_row + 1)}
    assert "样本数: 8" in contents["整体样本"] and "50.0%" in contents["整体样本"]
    assert report._target_name == "MOB1@3"
    definition = label_definition_table(report).iloc[0]
    assert definition["有效样本数"] == 8 and definition["剔除样本数"] == 4
    book.close()


def test_score_conversion_keeps_configured_and_fitted_direction_separate():
    p = np.linspace(0.1, 0.9, 8)
    converter = ScoreTransformer(lower=450, upper=550, direction="auto").fit(p)
    report = ModelReport(converter, datasets={"样本": (pd.DataFrame({"概率": p}), (p > 0.5).astype(int))}, n_jobs=1)
    params = report._get_score_conversion_sections()["params"].set_index("配置项")["配置值"]
    assert params["direction"] == "auto" and params["direction_"] == "ascending"
    assert params["A_"] == converter.transformer_.A_ and params["B_"] == converter.transformer_.B_
    assert report.score_direction == "higher_risk"


@pytest.mark.parametrize("n_jobs", [1, 2])
def test_all_bad_combined_label_keeps_report_complete_without_undefined_curves(data, tmp_path, n_jobs):
    # >=0 使联合标签全部为1；不能伪造 ROC/LIFT，也不应把合法报告当作失败。
    report = ModelReport(
        ProbabilityModel(),
        datasets={"样本": data.fillna(0)},
        overdue="MOB1",
        dpds=[0, 3],
        overdue_operator=">=",
        del_grey=True,
        n_jobs=n_jobs,
        parallel_backend="threading",
    )
    assert np.unique(report._datasets["样本"].y).tolist() == [1.0]
    path = tmp_path / "单类联合标签.xlsx"
    result = report.to_excel(path, with_plots=True, include_sample_records=False, mode="strict", return_result=True)
    assert result.complete
    assert not report.report_warnings_
    assert any(item["章节"] == "模型 KS/LIFT 图" for item in report.report_skips_)
    assert not (tmp_path / "单类联合标签_assets/ks_样本.png").exists()
