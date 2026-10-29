"""连续目标报告用原始行独立核算，禁止回归目标进入坏率/KS/AUC 口径。"""

import numpy as np
import pandas as pd
import pytest
from numpy.testing import assert_allclose
from sklearn.base import BaseEstimator, RegressorMixin
from openpyxl import load_workbook

from hscredit.report import ModelReport, auto_model_report, compare_models


class EchoRegressor(RegressorMixin, BaseEstimator):
    def fit(self, X, y=None):
        self.feature_names_in_ = np.array(["特征"])
        return self

    def predict(self, X):
        return np.asarray(X["特征"], dtype=float)


@pytest.fixture
def data():
    train = pd.DataFrame(
        {"特征": [0.0, 2.0, 6.0, 4.0], "目标": [1.0, 2.0, 3.0, 4.0], "权重": [1.0, 2.0, 1.0, 2.0]}, index=[7, 8, 9, 10]
    )
    test = pd.DataFrame(
        {"特征": [-2.0, 2.0, 8.0], "目标": [-1.5, 1.25, 5.5], "权重": [2.0, 0.0, 1.0]}, index=[17, 18, 19]
    )
    for frame in (train, test):
        frame["日期"] = pd.date_range("2026-01-01", periods=len(frame), freq="20D")
        frame["类别"] = np.where(np.arange(len(frame)) % 2, "甲", "乙")
        frame["金额"] = 100.0
    return train, test


def test_regression_metrics_match_independent_weighted_rows(data):
    train, test = data
    model = EchoRegressor().fit(train)
    report = ModelReport(
        model, datasets={"训练集": train, "独立测试": test}, target="目标", sample_weight="权重", n_jobs=1
    )
    table = report.get_metrics().set_index("统计项")
    baseline = np.average(train["目标"], weights=train["权重"])
    for name, frame in [("训练集", train), ("独立测试", test)]:
        y, p, w = (frame[column].to_numpy() for column in ("目标", "特征", "权重"))
        error = p - y
        squared = np.sum(w * error**2) / w.sum()
        variance = np.sum(w * (y - np.average(y, weights=w)) ** 2)
        assert_allclose(
            table.loc[["均方误差", "均方根误差", "平均绝对误差", "决定系数"], name],
            [squared, np.sqrt(squared), np.sum(w * np.abs(error)) / w.sum(), 1 - np.sum(w * error**2) / variance],
        )
        assert_allclose(table.loc["基准均方根误差", name], np.sqrt(np.sum(w * (y - baseline) ** 2) / w.sum()))
    assert not {"AUC", "KS", "坏样本率"}.intersection(table.index)
    assert report.prediction_method == "predict"
    assert "TOP名单规则" not in report.model_context_
    assert report.summary().index.name == "目标变量"
    assert report.to_dict()["metrics"]
    comparison = compare_models({"甲": model, "乙": model}, train, train["目标"], test, test["目标"], n_jobs=1)
    assert "均方根误差" in comparison.columns.get_level_values(0)


def test_continuous_targets_preserve_decimals_alignment_and_explicit_y(data):
    train, test = data
    model = EchoRegressor().fit(train)
    report = auto_model_report(
        model, datasets={"测试": (test, test["目标"].iloc[::-1])}, target="错误字段", verbose=False, n_jobs=1
    )
    assert_allclose(report._datasets["测试"].y, test["目标"])
    report.add_dataset("追加", "追加", train, y=train["目标"] + 0.25)
    assert_allclose(report._datasets["追加"].y, train["目标"] + 0.25)
    with pytest.raises(ValueError, match="索引"):
        ModelReport(model, datasets={"测试": (test, pd.Series([1, 2, 3]))}, n_jobs=1)
    with pytest.raises(ValueError, match="无穷"):
        ModelReport(model, datasets={"测试": (test, [1, np.inf, 3])}, n_jobs=1)
    with pytest.raises(ValueError, match="二分类"):
        ModelReport(model, datasets={"测试": test}, target="目标", metrics="auc", n_jobs=1)
    with pytest.raises(ValueError, match="连续预测"):
        ModelReport(model, datasets={"测试": test}, target="目标", prediction_method="transform", n_jobs=1)
    missing = test.assign(目标=pd.array([1.25, pd.NA, 5.5], dtype="Float64"))
    report = ModelReport(model, datasets={"测试": missing}, target="目标", n_jobs=1)
    assert report.get_metrics().set_index("统计项").loc["样本数", "测试"] == 2


def test_regression_binning_totals_use_original_rows_and_training_cutpoints(data):
    train, test = data
    report = ModelReport(
        EchoRegressor().fit(train), datasets={"训练集": train, "测试集": test}, target="目标", n_jobs=1
    )
    table = report.get_bin_table("测试集", max_n_bins=2, amount="金额")
    total = table.iloc[-1]
    assert total["分箱标签"] == "合计"
    assert total["样本总数"] == len(test) == table.iloc[:-1]["样本总数"].sum()
    assert_allclose(total["目标值均值"], test["目标"].mean())
    assert_allclose(total["目标值标准差"], test["目标"].std(ddof=0))
    assert_allclose(total["均方根误差"], np.sqrt(np.mean((test["目标"] - test["特征"]) ** 2)))
    assert total["金额合计"] == 300.0
    assert not any("坏" in column or "KS" in column or "IV" in column for column in table)
    changed = test.assign(目标=test["目标"] * 100)
    report2 = ModelReport(report.model, datasets={"训练集": train, "测试集": changed}, target="目标", n_jobs=1)
    assert report2.get_bin_table("测试集", max_n_bins=2)["分箱标签"].tolist() == table["分箱标签"].tolist()


def test_regression_strict_excel_has_complete_continuous_sections(data, tmp_path):
    train, test = data
    model = EchoRegressor().fit(train)
    output = tmp_path / "回归.xlsx"
    result = auto_model_report(
        model,
        datasets={"训练集": train, "独立测试": test},
        target="目标",
        excel_path=output,
        date="日期",
        group="类别",
        amount="金额",
        with_plots=True,
        mode="strict",
        return_result=True,
        verbose=False,
        n_jobs=1,
    )
    assert result.complete and output.is_file()
    workbook = load_workbook(output)
    assert {
        "目录",
        "1-基本信息",
        "2-模型性能",
        "3-入模变量分析",
        "4-稳定性分析",
        "5-模型参数",
        "6-模型部署需求",
        "7-训练与调参过程",
    } <= set(workbook.sheetnames)
    values = [cell.value for row in workbook["2-模型性能"] for cell in row]
    assert "均方根误差" in values and "坏样本率" not in values
    assert len(workbook["2-模型性能"]._images) == 2
    assert len(list((tmp_path / "回归_assets").glob("*.png"))) == 2
    assert all(cell.data_type != "e" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


@pytest.mark.parametrize("name", ["ResNet", "FTTransformer", "TabM"])
def test_net_missing_tree_importance_is_capability_not_report_failure(name, tmp_path):
    import hscredit as hc

    torch = pytest.importorskip("torch")
    previous = torch.get_num_threads()
    torch.set_num_threads(2)
    try:
        rng = np.random.default_rng(4)
        X = pd.DataFrame(rng.normal(size=(32, 2)), columns=["分数", "次数"])
        y = (X["分数"] > 0).astype(int)
        kwargs = {"k": 2} if name == "TabM" else {}
        model = getattr(hc, name)(
            d_block=16, n_blocks=1, max_epochs=1, early_stopping_rounds=None, device="cpu", **kwargs
        ).fit(X, y)
        result = model.model_report(
            datasets={"测试": (X, y)},
            excel_path=tmp_path / f"{name}.xlsx",
            with_plots=False,
            mode="strict",
            return_result=True,
            verbose=False,
            n_jobs=1,
        )
        assert result.complete
        assert result.sections["原生特征重要性"].status == "不适用"
        assert not any("降级：" in title for title in result.sections)
    finally:
        torch.set_num_threads(previous)
