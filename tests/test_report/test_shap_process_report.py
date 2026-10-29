"""SHAP 重要性、阅读型训练过程与动态页号的真实输出契约。"""

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from sklearn.base import BaseEstimator, RegressorMixin
from hscredit import RandomForest, ModelTuner
from hscredit.report import ModelReport, auto_model_report
from hscredit.core.models.explainability import ModelExplainer
from hscredit.report._training_process_report import process_tables


def test_title_covers_table_merges_and_picture_using_actual_column_widths():
    from types import SimpleNamespace
    from openpyxl import Workbook
    from hscredit.report._model_layout import fit_report_titles

    sheet = Workbook().active
    sheet["B2"] = "7、训练与调参过程"
    sheet["B5"] = "7.1、学习曲线"
    sheet["B7"] = "迭代"
    sheet["C7"] = "指标"
    sheet.merge_cells("C7:E7")
    for column in "BCDEFGHI":
        sheet.column_dimensions[column].width = 10
    sheet._images.append(SimpleNamespace(anchor="B9", width=600, height=200))
    fit_report_titles(sheet)
    areas = {a.min_row: a for a in sheet.merged_cells.ranges if a.min_col == 2}
    assert areas[2].max_col >= 9 and areas[5].max_col >= 9
    assert sheet["B5"].value == "7.1、学习曲线"


class AdditiveRegressor(RegressorMixin, BaseEstimator):
    def fit(self, X, y=None):
        self.feature_names_in_ = X.columns.to_numpy()
        return self

    def predict(self, X):
        return 2 * X["数值"].to_numpy() + (X["类别"].to_numpy() == "甲") * 3


@pytest.mark.parametrize("category_dtype", [False, True])
def test_shap_categorical_regression_matches_exact_additive_formula(category_dtype):
    X = pd.DataFrame({"数值": [0.0, 1.0, 2.0, 3.0], "类别": ["甲", "乙", "甲", "乙"]})
    if category_dtype:
        X["类别"] = pd.Categorical(X["类别"], categories=["乙", "甲", "丙"], ordered=True)

    class TypedRegressor(AdditiveRegressor):
        def predict(self, frame):
            assert frame["类别"].dtype == X["类别"].dtype
            if category_dtype:
                assert frame["类别"].cat.categories.tolist() == ["乙", "甲", "丙"]
            return super().predict(frame)

    model = TypedRegressor().fit(X)
    explainer = ModelExplainer(model, background_data=X, model_output="prediction", random_state=3)
    result = explainer.explain(X, max_evals=5)
    expected = np.column_stack([2 * (X["数值"] - 1.5), (X["类别"].eq("甲").astype(int) - 0.5) * 3])
    np.testing.assert_allclose(result.values, expected, atol=1e-7)
    np.testing.assert_allclose(result.base_values + result.values.sum(axis=1), model.predict(X))
    assert result.data["类别"].tolist() == X["类别"].tolist()
    assert explainer.get_shap_importance().index.tolist() == ["数值", "类别"]


@pytest.fixture
def fitted():
    rng = np.random.default_rng(7)
    X = pd.DataFrame(rng.normal(size=(60, 3)), columns=["甲", "乙", "丙"])
    y = (X["甲"] > 0).astype(int)
    tuner = ModelTuner(
        RandomForest(n_estimators=4, n_jobs=1, random_state=4),
        search_space={"max_depth": [2, 3]},
        n_jobs=1,
        random_state=4,
    )
    tuner.fit(X.iloc[:40], y.iloc[:40], eval_set=(X.iloc[40:], y.iloc[40:]), n_trials=2, show_progress_bar=False)
    return tuner.get_best_model(), X, y


def test_compact_process_shap_and_dynamic_numbering(fitted, tmp_path):
    model, X, y = fitted
    report = ModelReport(
        model,
        datasets={"训练": (X.iloc[:40], y.iloc[:40]), "验证": (X.iloc[40:], y.iloc[40:])},
        n_jobs=1,
        explain_config={
            "enabled": True,
            "scope": "importance",
            "dataset": "验证",
            "max_samples": 8,
            "max_background": 8,
        },
    )
    table = report.get_shap_importance()
    assert {"原生特征重要性", "平均绝对SHAP值", "平均SHAP值"} <= set(table)
    result = report.get_model_explanation()["解释结果"]
    np.testing.assert_allclose(
        table.set_index("特征")["平均绝对SHAP值"].reindex(X.columns), np.abs(result.values).mean(axis=0)
    )
    display = process_tables(report)
    assert "模型口径" not in display and "预处理参数" not in display
    assert len(display["参数搜索对比"]) == 2
    assert "参数：max_depth" in display["参数搜索对比"]
    path = tmp_path / "报告.xlsx"
    report.to_excel(path, with_plots=False, mode="strict")
    book = load_workbook(path)
    assert "7-训练与调参过程" in book.sheetnames and "7-模型解释" not in book.sheetnames
    assert book["7-训练与调参过程"].max_row < 100
    assert book["7-训练与调参过程"]["B2"].value == "7、训练与调参过程"
    assert book["7-训练与调参过程"]["B5"].value == "7.1、训练概览"
    assert "1、SHAP 特征重要性" in [c.value for row in book["3-入模变量分析"] for c in row]
    book.close()
    report.to_excel(path, with_plots=False, training_sheet_index=10, mode="strict")
    book = load_workbook(path)
    assert "10-训练与调参过程" in book.sheetnames
    assert any(c.hyperlink and c.hyperlink.location == "#'10-训练与调参过程'!B2" for row in book["目录"] for c in row)
    book.close()
    with pytest.raises(ValueError, match="training_sheet_index"):
        report.to_excel(path, training_sheet_index=6)
    for invalid in (True, 7.5, "8", -1):
        with pytest.raises(ValueError, match="training_sheet_index"):
            report.to_excel(path, training_sheet_index=invalid)
    full = ModelReport(model, datasets={"训练": (X, y)}, n_jobs=1, explain_config={"enabled": True})
    with pytest.raises(ValueError, match="training_sheet_index"):
        full.to_excel(path, training_sheet_index=7)


def test_regression_shap_export(fitted, tmp_path):
    X = pd.DataFrame({"数值": [0.0, 1.0, 2.0, 3.0], "类别": ["甲", "乙", "甲", "乙"]})
    model = AdditiveRegressor().fit(X)
    path = tmp_path / "回归.xlsx"
    report = auto_model_report(
        model,
        datasets={"训练": (X, model.predict(X))},
        n_jobs=1,
        verbose=False,
        explain_config={"enabled": True, "scope": "importance", "max_background": 4, "max_evals": 5},
    )
    result = report.to_excel(path, with_plots=True, mode="strict", return_result=True, training_sheet_index=9)
    assert result.complete
    book = load_workbook(path)
    assert "9-训练与调参过程" in book.sheetnames
    assert len(book["3-入模变量分析"]._images) == 1
    book.close()


def test_training_titles_keep_sheet_number_after_learning_curve_plots(fitted, tmp_path):
    model, X, y = fitted
    report = ModelReport(model, datasets={"训练": (X, y)}, n_jobs=1)
    curve = pd.DataFrame({"数据集": ["验证"] * 3, "指标": ["损失"] * 3, "迭代": [1, 2, 3], "数值": [0.6, 0.5, 0.4]})
    report._training_process_snapshot["选中试验曲线"] = curve.assign(试验编号=1)
    report._training_process_snapshot["训练曲线"] = curve
    path = tmp_path / "学习曲线.xlsx"
    report.to_excel(path, with_plots=True, training_sheet_index=10, mode="strict")
    book = load_workbook(path)
    sheet = book["10-训练与调参过程"]
    titles = [
        c.value
        for row in sheet.iter_rows(min_col=2, max_col=2)
        for c in row
        if isinstance(c.value, str) and "、" in c.value
    ]
    assert titles[0] == "10、训练与调参过程"
    assert all(title.startswith(f"10.{index}、") for index, title in enumerate(titles[1:], 1))
    assert titles[-2].endswith("选中试验学习曲线") and titles[-1].endswith("最终模型学习曲线")
    assert len(sheet._images) == 2
    book.close()
