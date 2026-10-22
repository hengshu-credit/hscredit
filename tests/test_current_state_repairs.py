"""独立现状审查缺陷的行为回归，包含真实数据与资源前置校验。"""

import importlib
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from hscredit import (
    OneHotEncoder,
    OptimalBinning,
    Rule,
    RuleFlow,
    TreeRuleExtractor,
    feature_bin_stats,
    feature_summary,
)
from hscredit.database.adapters.base import BaseDatabaseAdapter
from hscredit.exceptions import ValidationError
from hscredit.report.population_drift import population_drift
from hscredit.report.result import ReportGenerationError


@pytest.fixture
def real_data():
    path = Path(__file__).resolve().parents[1] / "examples" / "hscredit_yyp.xlsx"
    if not path.exists():
        pytest.skip("缺少 examples/hscredit_yyp.xlsx")
    return pd.read_excel(path)


@pytest.mark.parametrize("multi", [False, True])
@pytest.mark.parametrize("return_rules,return_binner", [(False, False), (True, False), (False, True), (True, True)])
def test_real_categorical_report_roundtrip(real_data, multi, return_rules, return_binner):
    features = ["衡枢鉴真分老客版", "商品类别"] if multi else "商品类别"
    output = feature_bin_stats(
        real_data,
        features,
        target="FPD",
        method="quantile",
        amount="放款金额",
        return_rules=return_rules,
        return_binner=return_binner,
        n_jobs=1,
    )
    table = output[0] if isinstance(output, tuple) else output
    assert not table.empty
    if return_rules:
        rules = output[1]
        restored = OptimalBinning(method="quantile", n_jobs=1).fit(real_data[["商品类别"]], real_data.FPD)
        baseline = restored.transform(real_data[["商品类别"]])
        restored.import_rules({"商品类别": rules["商品类别"]})
        pd.testing.assert_frame_equal(restored.transform(real_data[["商品类别"]]), baseline)


@pytest.mark.parametrize("sparse", [False, True])
@pytest.mark.parametrize("return_df", [False, True])
def test_onehot_writable_output_and_missing_unknown(sparse, return_df):
    frame = pd.DataFrame({"类别": [1, 2, 1, np.nan]}, index=[1, 1, 4, 9])
    original = frame.copy(deep=True)
    encoder = OneHotEncoder(n_jobs=1, sparse_output=sparse, return_df=return_df)
    result = encoder.fit_transform(frame)
    array = (
        result.sparse.to_dense().to_numpy()
        if sparse and return_df
        else result.toarray() if sparse else np.asarray(result)
    )
    np.testing.assert_array_equal(array.sum(axis=1), np.ones(4))
    unknown = encoder.transform(pd.DataFrame({"类别": [999]}))
    assert float(np.asarray(unknown.toarray() if sparse and not return_df else unknown).sum()) == 0
    pd.testing.assert_frame_equal(frame, original)


@pytest.mark.parametrize("mode", ["serial", "parallel"])
@pytest.mark.parametrize("budget", [{"max_output_rows": 10}, {"memory_budget_bytes": 50000}])
def test_ruleflow_budget_rejects_before_rules(mode, budget, monkeypatch):
    flow = RuleFlow([Rule(f"值 > {i}") for i in range(100)], mode=mode, n_jobs=1, parallel_config=budget)
    monkeypatch.setattr(Rule, "predict", lambda *a, **k: pytest.fail("超预算时不得开始计算规则"))
    with pytest.raises(ValidationError, match="超过"):
        flow.predict(pd.DataFrame({"值": np.arange(1000)}))


def test_ruleflow_row_budget_is_rows_not_rule_cells():
    data = pd.DataFrame({"值": np.arange(20)})
    flow = RuleFlow(
        [Rule("值 > 3"), Rule("值 > 8")],
        n_jobs=1,
        mode="parallel",
        parallel_config={"max_output_rows": 20, "memory_budget_bytes": 1000000},
    )
    result = flow.predict(data)
    assert result.shape == (20, 6)
    assert result["是否命中"].sum() == 16


def test_auto_binning_projects_before_dispatch(monkeypatch):
    module = importlib.import_module("hscredit.core.binning.optimal_binning")
    frame = pd.DataFrame(np.ones((100, 100)))
    target = pd.Series(np.arange(100) % 2)

    def dispatch(function, tasks, **kwargs):
        assert tasks[0][1].columns.tolist() == [0]
        assert kwargs["workload"].data_bytes == tasks[0][1].memory_usage(deep=True).sum() + target.memory_usage(
            deep=True
        )
        return [("quantile", 1.0, None)]

    monkeypatch.setattr(module, "parallel_execute", dispatch)
    assert OptimalBinning.auto_select_method(frame, target, 0, methods=["quantile"]) == "quantile"


def test_auto_binning_invalid_criterion_fails_before_dispatch(monkeypatch):
    module = importlib.import_module("hscredit.core.binning.optimal_binning")
    monkeypatch.setattr(module, "parallel_execute", lambda *a, **k: pytest.fail("指标无效时不得启动任务"))
    with pytest.raises(ValueError, match="criterion"):
        OptimalBinning.auto_select_method(pd.DataFrame({"x": [0, 1]}), pd.Series([0, 1]), "x", criterion="wrong")


def test_summary_type_explicit_and_numeric_datetime():
    frame = pd.DataFrame(
        {
            "金额": np.arange(100) + 0.5,
            "日期": pd.date_range("2020-01-01", periods=100),
            "编码": np.arange(100),
            "强制数值": ["1"] * 100,
        }
    )
    result = feature_summary(frame, numeric_as_categorical=["编码"], force_numeric=["强制数值"], n_jobs=1).set_index(
        "特征名"
    )
    assert result["字段类型"].to_dict() == {
        "金额": "numerical",
        "日期": "datetime",
        "编码": "categorical",
        "强制数值": "numerical",
    }


@pytest.fixture
def registry(monkeypatch):
    module = importlib.import_module("hscredit.database.registry")
    for attribute in ("_ADAPTERS", "_ALIASES", "_REGISTERED_ALIASES"):
        monkeypatch.setattr(module, attribute, dict(getattr(module, attribute)))
    return module


def test_registry_conflicts_do_not_mutate(registry):
    class Adapter(BaseDatabaseAdapter):
        pass

    before = dict(registry._ADAPTERS), dict(registry._ALIASES)
    for name, aliases in [("custom", ["mysql"]), ("mongo", []), ("custom", ["custom"]), ("custom", ["dup", "dup"])]:
        with pytest.raises(ValidationError):
            registry.register_adapter(name, Adapter, aliases=aliases)
        assert (registry._ADAPTERS, registry._ALIASES) == before
    with pytest.raises(ValidationError, match="适配器类"):
        registry.register_adapter("custom", dict)
    registry.register_adapter("custom", Adapter, aliases=["private"])
    assert registry.get_adapter_class("private") is Adapter
    registry.register_adapter("private", Adapter, replace=True)
    assert registry.canonical_adapter_name("private") == "private"


class BrokenImportance:
    def get_feature_importances(self):
        raise RuntimeError("预期模型故障")


def test_importance_failure_has_status_or_strict_error():
    frame = pd.DataFrame({"x": [1, 2, 3]})
    with pytest.warns(UserWarning, match="预期模型故障"):
        table = feature_summary(frame, models={"故障": BrokenImportance()}, n_jobs=1)
    assert not table.attrs["完整"]
    assert table.attrs["分析状态"][0]["状态"] == "失败"
    with pytest.raises(ReportGenerationError, match="预期模型故障"):
        feature_summary(frame, models={"故障": BrokenImportance()}, mode="strict", n_jobs=1)


def test_importance_success_is_in_structured_result():
    class Model:
        def get_feature_importances(self):
            return pd.Series({"x": 0.8})

    result = feature_summary(pd.DataFrame({"x": [1, 2, 3]}), models={"模型": Model()}, return_result=True, n_jobs=1)
    assert result.complete
    assert result["特征摘要"]["模型重要性"].iloc[0] == 0.8


@pytest.fixture
def drift_data():
    return pd.DataFrame({"值": [1, 2, 3, 4, 5, 6], "日期": pd.to_datetime(["2024-01-01"] * 5 + [None])})


def test_drift_strict_missing_fields_never_publishes(drift_data, tmp_path):
    path = tmp_path / "报告.xlsx"
    with pytest.raises(ReportGenerationError, match="缺少"):
        population_drift(drift_data, drift_data, ["不存在"], output=path, mode="strict", n_jobs=1)
    assert not path.exists()


@pytest.mark.parametrize("policy, included, excluded", [("drop", 5, 1), ("separate", 6, 0)])
def test_drift_discloses_missing_fields_and_rows(drift_data, tmp_path, policy, included, excluded):
    path = tmp_path / "报告.xlsx"
    with pytest.warns(UserWarning):
        result = population_drift(
            drift_data,
            drift_data,
            ["值", "不存在"],
            date="日期",
            missing_date=policy,
            output=path,
            return_result=True,
            n_jobs=1,
        )
    row = result["样本口径"].set_index("数据集").loc["实际"]
    assert row["输入行数"] == row["纳入行数"] + row["排除行数"] == 6
    assert row["纳入行数"] == included and row["排除行数"] == excluded
    assert result.sections["特征:不存在"].status == "失败"
    assert not result.complete
    workbook = load_workbook(path, read_only=True)
    assert {"样本口径", "执行状态"}.issubset(workbook.sheetnames)
    workbook.close()


def test_drift_missing_dates_require_policy(drift_data, tmp_path):
    with pytest.raises(ValidationError, match="缺失日期"):
        population_drift(drift_data, drift_data, ["值"], date="日期", output=tmp_path / "报告.xlsx", n_jobs=1)


@pytest.mark.parametrize("mode", ["strict", "best_effort"])
def test_drift_score_failure_uses_report_status_and_publish_policy(tmp_path, mode):
    data = pd.DataFrame({"变量": [1, 2, 3, 4], "评分": [1.0, 2.0, np.inf, 4.0]})
    path = tmp_path / "监控.xlsx"
    options = dict(score_col="评分", mode=mode, return_result=True, n_jobs=1, output=path)
    if mode == "strict":
        with pytest.raises(ReportGenerationError, match="评分分布对比") as raised:
            population_drift(data, data, ["变量"], **options)
        assert raised.value.result.sections["评分分布对比"].status == "失败"
        assert not path.exists()
    else:
        with pytest.warns(UserWarning, match="评分分布对比"):
            result = population_drift(data, data, ["变量"], **options)
        assert not result.complete and "评分分布对比" not in result
        assert result.sections["评分分布对比"].status == "失败"
        assert not result["PSI总览"].empty
        assert path.exists()


@pytest.mark.parametrize("scores", [[None, None, None, None], [1.0, None, 2.0, None]])
def test_drift_score_distribution_includes_missing_samples(tmp_path, scores):
    data = pd.DataFrame({"变量": [1, 2, 3, 4], "评分": scores})
    result = population_drift(
        data,
        data,
        ["变量"],
        score_col="评分",
        mode="strict",
        return_result=True,
        output=tmp_path / "监控.xlsx",
        n_jobs=1,
    )
    assert result.complete
    distribution = result["评分分布对比"]
    missing = distribution["分箱"].eq("缺失值")
    assert distribution.loc[missing, "实际占比"].iloc[0] == pd.Series(scores).isna().mean()
    assert distribution["基准占比"].sum() == pytest.approx(1.0)
    assert result.sections["评分分布对比"].input_rows == len(data)


@pytest.mark.parametrize("dtype", [object, "string"])
def test_tree_missing_categories_and_literals_have_identical_leaf_rules(dtype):
    labels = ["nan", "None", None, "a'b", "a,b", "普通"] * 20
    frame = pd.DataFrame(
        {
            "类别": pd.Series(labels, dtype=dtype),
            "数值": [np.nan, 1.0, -1.0, 2.0, 3.0, 4.0] * 20,
            "target": [1, 1, 0, 0, 1, 0] * 20,
        }
    )
    miner = TreeRuleExtractor(
        algorithm="dt", max_depth=4, min_samples_leaf=1, min_samples_split=2, n_jobs=1, random_state=4
    ).fit(frame)
    assert len(miner._raw_category_values_["类别"]) == 6
    for item in miner.extract_rules():
        raw = frame.loc[miner.X_train_.index]
        expected = miner._apply_conditions(item["conditions"], miner.X_train_)
        actual = Rule(miner._rule_to_string(item)).predict(raw)
        np.testing.assert_array_equal(actual, expected)
