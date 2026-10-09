"""生成 README 的 hscredit 原生 Excel 功能演示。"""

from pathlib import Path
import json
import os
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("MPLBACKEND", "Agg")

import pandas as pd
from openpyxl import load_workbook

import hscredit
from hscredit.core.binning import OptimalBinning
from hscredit.excel import ExcelWriter, dataframe2excel


def main():
    output = ROOT / "docs/assets/readme/excel"
    output.mkdir(parents=True, exist_ok=True)
    df = pd.read_excel(ROOT / "examples/hscredit_yyp.xlsx")
    features = ["衡枢鉴真分老客版", "近六个月非银多头机构数", "手机号近一个月非银多头机构数"]
    binner = OptimalBinning(method="quantile", max_n_bins=4, n_jobs=1)
    binner.fit(df[features], df["FPD"])
    table = binner.get_bin_table("衡枢鉴真分老客版")
    percent_cols = ["样本占比", "好样本占比", "坏样本占比", "坏样本率", "坏账改善", "累积坏账改善", "分档KS值"]
    condition_cols = ["样本总数", "分档IV值", "LIFT值"]

    table.save(
        str(output / "dataframe-save.xlsx"),
        sheet_name="分箱统计",
        title="DataFrame.save 完整分箱表",
        index=True,
        auto_width=True,
        percent_cols=percent_cols,
        condition_cols=condition_cols,
    )

    writer = ExcelWriter()
    dataframe2excel(
        table,
        writer,
        sheet_name="分箱统计",
        title="dataframe2excel 完整分箱表",
        start_row=4,
        index=True,
        auto_width=True,
        percent_cols=percent_cols,
        condition_cols=condition_cols,
    )
    sheet = writer.get_sheet_by_name("分箱统计")
    writer.set_freeze_panes(sheet, "D7")
    writer.insert_value2sheet(sheet, "B2", value="查看数据透视表")
    writer.insert_hyperlink2sheet(sheet, "B2", sheet="透视分析", target_space="B2")

    writer.insert_pivot_table2sheet(
        worksheet="透视分析",
        data=table,
        pivot_anchor="B2",
        rows="分箱标签",
        values=[
            ("样本总数", "sum"),
            ("坏样本数", "sum"),
            {"field": "样本总数", "agg": "sum", "show_as": "全局占比", "name": "样本占比", "number_format": "0.00%"},
        ],
        source_sheet=sheet,
        source_anchor="C6",
        write_source=False,
        name="分箱样本透视表",
    )
    writer.save(str(output / "excel-features.xlsx"))

    # 透视表单元格在首次保存时生成，再由同一原生 Writer 完成自动列宽。
    formatted = ExcelWriter(style_excel=str(output / "excel-features.xlsx"))
    pivot = formatted.get_sheet_by_name("透视分析")
    formatted.adjust_columns_width(pivot)
    formatted.add_conditional_formatting(pivot, "C3", "C6")
    formatted.save(str(output / "excel-features.xlsx"))

    inspections = []
    for path in output.glob("*.xlsx"):
        book = load_workbook(path)
        for worksheet in book.worksheets:
            inspections.append({
                "file": path.name,
                "sheet": worksheet.title,
                "range": worksheet.calculate_dimension(),
                "freeze": worksheet.freeze_panes,
                "conditional_formats": len(worksheet.conditional_formatting),
                "pivots": len(worksheet._pivots),
            })
        book.close()
    workbook = load_workbook(output / "excel-features.xlsx")
    assert workbook["分箱统计"].freeze_panes == "D7"
    assert workbook["分箱统计"]["B2"].hyperlink.location == "#透视分析!B2"
    assert len(workbook["分箱统计"].conditional_formatting) == 3
    assert len(workbook["透视分析"]._pivots) == 1
    assert any("%" in cell.number_format for row in workbook["分箱统计"] for cell in row)
    workbook.close()
    with zipfile.ZipFile(output / "excel-features.xlsx") as archive:
        assert any(name.startswith("xl/pivotTables/pivotTable") for name in archive.namelist())
        assert any(name.startswith("xl/pivotCache/pivotCacheRecords") for name in archive.namelist())
    (output / "manifest.json").write_text(json.dumps(inspections, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(inspections, ensure_ascii=False))


if __name__ == "__main__":
    main()
