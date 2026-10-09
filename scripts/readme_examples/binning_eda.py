"""生成 README 分箱与 EDA 演示：保留完整返回表和方法直接保存的图片。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import sys
import textwrap

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
from sklearn.model_selection import train_test_split

import hscredit
from hscredit.excel import ExcelWriter
from hscredit.core.viz import bin_plot, bin_2d_plot, bin_trend_plot, bin_overdues_plot
from hscredit.report import (
    feature_bin_stats,
    feature_binning_summary,
    feature_group_binning_summary,
    feature_efficiency_analysis,
    auto_feature_analysis,
)

ASSET_DIR = ROOT / "docs/assets/readme/binning-eda"
SECTION_PATH = ROOT / ".audit_tmp/readme-complete/sections/binning_eda.md"
PERCENT_COLUMNS = [
    "样本占比", "好样本占比", "坏样本占比", "坏样本率", "坏账改善",
    "累积坏账改善", "累计好样本占比", "累计坏样本占比", "分档KS值",
    "缺失率", "唯一值占比", "最大值占比", "众数占比", "零值率", "负值率", "重复率", "KS",
]
CONDITION_COLUMNS = ["坏样本率", "LIFT值", "指标IV值", "IV", "KS", "PSI"]

# 每段代码直接调用公开 API；这里的名字也用于收集原始返回对象。
CASES = [
    {
        "id": "summary", "title": "DataFrame.summary：数据质量、区分度与跨月稳定性",
        "description": "一次返回三个数值特征和商品类别的完整综合统计；字段类型交由原生 API 判断。本数据中唯一值较多的评分被原生识别为 id，趋势返回 unknown，结果原样保留。",
        "code": '''
summary = df.summary(
    features=features + ["商品类别"], y="FPD",
    max_n_bins=4, psi_method="date_col", psi_date_col="放款时间", n_jobs=1,
)
display(summary)
''',
        "outputs": ["summary"],
    },
    {
        "id": "method-comparison", "title": "不同分箱方法：完整分箱明细与 KS、Lift、IV 评估",
        "description": "在相同数据、目标和最大箱数下比较等频、CART 和最优 IV；summary 的指标聚合由 feature_binning_summary 原生完成。",
        "code": '''
method_tables, method_summary = feature_binning_summary(
    df, feature=features, methods=["quantile", "cart", "best_iv"],
    target="FPD", max_n_bins=4, margins=True, random_state=42, n_jobs=1,
)
display(method_summary)
method_summary.save(
    str(ASSET_DIR / "method-summary.xlsx"), index=True, auto_width=True,
    percent_cols=[("分档KS值", "FPD"), ("坏样本率", "FPD")],
    condition_cols=[("LIFT值", "FPD"), ("指标IV值", "FPD")],
)
for feature_name, by_method in method_tables.items():
    for method_name, table in by_method.items():
        print(feature_name, method_name)
        display(table)
''',
        "outputs": ["method_summary", "method_tables"], "workbooks": ["method-summary.xlsx"],
    },
    {
        "id": "multi-dpd", "title": "feature_bin_stats：多标签、金额口径与类别特征",
        "description": "MOB1 与 7、3、0 三个阈值在同一次调用中展开；amount 加入放款金额口径，margins=True 保留方法生成的合计。",
        "code": '''
multi_dpd_table = feature_bin_stats(
    df, feature=features, overdue=["MOB1"], dpds=[7, 3, 0],
    amount="放款金额", method="quantile", max_n_bins=4, margins=True, n_jobs=1,
)
category_table = feature_bin_stats(
    df, feature="商品类别", target="FPD", method="quantile",
    max_n_bins=4, margins=True, n_jobs=1,
)
display(multi_dpd_table)
display(category_table)
''',
        "outputs": ["multi_dpd_table", "category_table"],
    },
    {
        "id": "group-comparison", "title": "feature_group_binning_summary：跨月份与商品类别比较",
        "description": "同一特征和方法在全量数据上拟合一次，各月份或商品类别复用切点，完整返回每组明细与汇总。",
        "code": '''
monthly_tables, monthly_summary = feature_group_binning_summary(
    df, feature=features, methods=["quantile", "cart"],
    date_col="放款时间", freq="M", target="FPD",
    max_n_bins=4, margins=True, random_state=42, n_jobs=1,
)
category_tables, category_summary = feature_group_binning_summary(
    df, feature=features, methods="quantile", group_col="商品类别", target="FPD",
    max_n_bins=4, margins=True, n_jobs=1,
)
display(monthly_summary)
display(category_summary)
for grouped_result in (monthly_tables, category_tables):
    for feature_name, by_method in grouped_result.items():
        for method_name, by_group in by_method.items():
            for group_name, table in by_group.items():
                print(feature_name, method_name, group_name)
                display(table)
''',
        "outputs": ["monthly_summary", "category_summary", "monthly_tables", "category_tables"],
    },
    {
        "id": "bin-plot", "title": "bin_plot：单变量的样本结构与坏率",
        "description": "直接传入完整原生分箱统计表；图片通过 save 参数输出。",
        "code": '''
single_bin_table = feature_bin_stats(
    df, feature="衡枢鉴真分老客版", target="FPD",
    method="quantile", max_n_bins=4, margins=True, n_jobs=1,
)
figure = bin_plot(
    single_bin_table, desc="衡枢鉴真分老客版",
    save=str(ASSET_DIR / "bin-plot.png"),
)
display(single_bin_table)
display(figure)
''',
        "outputs": ["single_bin_table"], "images": ["bin-plot.png"],
    },
    {
        "id": "bin-2d", "title": "bin_2d_plot：两变量交叉风险与二维分箱",
        "description": "原生九宫格同时展示单变量分箱、分箱后的 KS 曲线和五类交叉指标。",
        "code": '''
figure = bin_2d_plot(
    df, features=["衡枢鉴真分老客版", "近六个月非银多头机构数"],
    target="FPD", method="quantile", max_n_bins=4,
    binner_kwargs={"n_jobs": 1}, save=str(ASSET_DIR / "bin-2d.png"),
)
display(figure)
''',
        "outputs": [], "images": ["bin-2d.png"],
    },
    {
        "id": "bin-trend", "title": "bin_trend_plot：按放款月份检查分箱风险趋势",
        "description": "保留方法默认分箱行为和完整月份，不缩减分组或二次调整图片；各月份的实际切点以原图标签为准。",
        "code": '''
figure = bin_trend_plot(
    df, feature="衡枢鉴真分老客版", target="FPD", date_col="放款时间",
    method="quantile", max_n_bins=4, n_jobs=1,
    save=str(ASSET_DIR / "bin-trend.png"),
)
display(figure)
''',
        "outputs": [], "images": ["bin-trend.png"],
    },
    {
        "id": "bin-overdues", "title": "bin_overdues_plot：对比 MOB1 的三种逾期口径",
        "description": "feature_bin_stats 默认按 MOB1 > DPD 生成三种标签。完整多级表头统计表直接传给 bin_table；没有在方法外拆标签、筛字段或重算指标。当前表模式能绘制三种坏率，但其适配器未识别当前统计表的 IV/KS/Lift 列名，因此图顶摘要仅显示趋势；完整指标仍保留在下方原始表和 Excel。",
        "code": '''
overdue_table = feature_bin_stats(
    df, feature="衡枢鉴真分老客版", overdue=["MOB1"], dpds=[7, 3, 0],
    method="quantile", max_n_bins=4, margins=True, n_jobs=1,
)
figure = bin_overdues_plot(
    overdue_table, bin_table=overdue_table, n_jobs=1,
    save=str(ASSET_DIR / "bin-overdues.png"),
)
display(overdue_table)
display(figure)
''',
        "outputs": ["overdue_table"], "images": ["bin-overdues.png"],
    },
    {
        "id": "efficiency", "title": "feature_efficiency_analysis：手工与自动分箱效率",
        "description": "手工规则未传入时使用方法内置分位数切点；自动分箱最多四箱。返回两张完整统计表、原始规则、2×2 对比图及两张月份趋势图。",
        "code": '''
efficiency = feature_efficiency_analysis(
    df, feature="衡枢鉴真分老客版", target="FPD", auto_method="quantile",
    max_n_bins=4, date_col="放款时间", margins=True, n_jobs=1,
    trend_kwargs={"n_jobs": 1}, output_dir=str(ASSET_DIR / "efficiency"),
)
display(efficiency["manual_table"])
display(efficiency["auto_table"])
display(efficiency["manual_rules"])
display(efficiency["auto_rules"])
display(efficiency["comparison_figure"])
for figure in efficiency["trend_figures"].values():
    display(figure)
''',
        "outputs": ["efficiency"], "image_glob": "efficiency/*.png",
    },
    {
        "id": "auto-feature", "title": "auto_feature_analysis：多 DPD、金额与时间分析报告",
        "description": "三特征和三种 MOB1 阈值放入一个报告，原生生成样本概况、月份分布、综合统计、图表及订单/金额分箱表。strict 模式只有必需章节成功才发布工作簿。工作簿保留报告方法本身的格式，包括其将 Lift 显示为百分比的默认行为；下方完整返回表另存的 Excel 则把 Lift 保持为倍率。",
        "code": '''
feature_report = auto_feature_analysis(
    df, features=features, overdue=["MOB1"], dpds=[7, 3, 0],
    amount="放款金额", date="放款时间", margins=True,
    bin_params={"method": "quantile", "max_n_bins": 4, "n_jobs": 1},
    excel_writer=str(ASSET_DIR / "auto-feature-report.xlsx"),
    output_dir=str(ASSET_DIR / "auto-feature-plots"),
    n_jobs=1, mode="strict", return_result=True,
)
display(feature_report.status_table())
''',
        "outputs": ["feature_report"], "workbooks": ["auto-feature-report.xlsx"],
        "image_glob": "auto-feature-plots/*.png",
    },
]


def raw_frames(value, prefix):
    """逐个交付原始返回中的全部 DataFrame，不改动行列。"""
    if isinstance(value, pd.DataFrame):
        yield prefix, value
    elif isinstance(value, dict):
        for key, child in value.items():
            yield from raw_frames(child, prefix + "/" + str(key))
    elif hasattr(value, "status_table"):
        yield prefix + "/status_table", value.status_table()


def repo_link(path):
    return Path(path).relative_to(ROOT).as_posix()


def verify_workbook(path):
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        return {sheet.title: {"rows": sheet.max_row, "columns": sheet.max_column,
                              "images": len(sheet._images)} for sheet in workbook.worksheets}
    finally:
        workbook.close()


def save_raw_tables(case, namespace):
    frames = []
    for name in case["outputs"]:
        frames.extend(raw_frames(namespace[name], name))
    if not frames:
        return []
    workbook_path = ASSET_DIR / (case["id"] + "-tables.xlsx")
    table_metadata = []
    with ExcelWriter().set_filename(str(workbook_path)) as writer:
        for number, (label, table) in enumerate(frames, 1):
            html_path = ASSET_DIR / f'{case["id"]}-table-{number:02d}.html'
            html_path.write_text(table.to_html(), encoding="utf-8")
            sheet_name = f"原始表{number}"
            worksheet = writer.get_sheet_by_name(sheet_name)
            # 汇总表的指标在列索引第一层，分箱明细的指标在最后层。
            # 只选择格式化目标列标签，不切取或改写原始 DataFrame。
            percent_columns = [
                column for column in table.columns
                if any(level in PERCENT_COLUMNS for level in (column if isinstance(column, tuple) else (column,)))
            ]
            condition_columns = [
                column for column in table.columns
                if any(level in CONDITION_COLUMNS for level in (column if isinstance(column, tuple) else (column,)))
            ]
            table.save(
                writer, worksheet=worksheet, title=label, index=True, auto_width=True,
                percent_cols=percent_columns, condition_cols=condition_columns,
            )
            table_metadata.append({"label": label, "rows": len(table), "columns": len(table.columns),
                                   "html": repo_link(html_path), "workbook": repo_link(workbook_path),
                                   "sheet": sheet_name})
    verify_workbook(workbook_path)
    return table_metadata


def write_section(manifest):
    lines = ["### 数据概览与分箱分析", "", "所有 DataFrame 都保留 API 返回的完整行列和索引；Excel 使用原生 `.save(auto_width=True)` 导出，比例列使用百分比、Lift 保持倍率，图像只由方法的 `save` 或 `output_dir` 保存。", "",
             "以下各段接公共数据准备代码运行：", "", "```python", "from pathlib import Path",
             "from hscredit.core.viz import bin_plot, bin_2d_plot, bin_trend_plot, bin_overdues_plot",
             "from hscredit.report import (", "    feature_bin_stats, feature_binning_summary, feature_group_binning_summary,",
             "    feature_efficiency_analysis, auto_feature_analysis,", ")",
             'ASSET_DIR = Path("docs/assets/readme/binning-eda")', "ASSET_DIR.mkdir(parents=True, exist_ok=True)", "```", ""]
    for case in CASES:
        outcome = manifest.get(case["id"])
        if outcome is None:
            continue
        lines.extend(["#### " + case["title"], "", case["description"], "", "```python",
                      textwrap.dedent(case["code"]).strip(), "```", ""])
        if outcome.get("error"):
            lines.extend(["当前原生接口未完成该调用：`" + outcome["error"] + "`。没有重绘或构造替代结果。", ""])
            continue
        for image in outcome["images"]:
            lines.extend([f'![{case["id"]} 原生输出]({image["path"]})', ""])
        for workbook in outcome["workbooks"]:
            lines.extend([f'[下载完整原生 Excel]({workbook["path"]})', ""])
        if outcome["tables"]:
            lines.extend(["<details>", "<summary>查看所有完整返回表与 Excel</summary>", ""])
            for table in outcome["tables"]:
                lines.append(f'- `{table["label"]}`：{table["rows"]} 行 × {table["columns"]} 列；[完整 HTML]({table["html"]}) · [Excel：{table["sheet"]}]({table["workbook"]})')
            lines.extend(["", "</details>", ""])
            # 概览和方法汇总内嵌完整表；其余完整明细均由上方链接提供。
            if case["id"] in {"summary", "method-comparison"}:
                first_table = outcome["tables"][0]
                lines.extend(["<details>", f'<summary>展开 {first_table["label"]} 完整原始表</summary>', "",
                              (ROOT / first_table["html"]).read_text(encoding="utf-8"), "", "</details>", ""])
    lines.extend(["完整可重跑代码：[binning_eda.py](scripts/readme_examples/binning_eda.py)。", ""])
    SECTION_PATH.parent.mkdir(parents=True, exist_ok=True)
    SECTION_PATH.write_text("\n".join(lines), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only", nargs="*", choices=[case["id"] for case in CASES])
    args = parser.parse_args()
    os.chdir(ROOT)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    manifest_path = ASSET_DIR / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    df = pd.read_excel("examples/hscredit_yyp.xlsx")
    features = ["衡枢鉴真分老客版", "近六个月非银多头机构数", "手机号近一个月非银多头机构数"]
    X = df[features]
    y = df["FPD"]
    X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.25, stratify=y, random_state=42)
    namespace = dict(globals(), **locals())
    namespace["display"] = lambda value: None
    for case in CASES:
        if args.only and case["id"] not in args.only:
            continue
        print("开始：" + case["id"], flush=True)
        outcome = {"tables": [], "images": [], "workbooks": []}
        try:
            exec(compile(textwrap.dedent(case["code"]), case["id"], "exec"), namespace)
            outcome["tables"] = save_raw_tables(case, namespace)
            image_paths = [ASSET_DIR / filename for filename in case.get("images", [])]
            image_paths.extend(sorted(ASSET_DIR.glob(case["image_glob"])) if case.get("image_glob") else [])
            for path in image_paths:
                with Image.open(path) as picture:
                    outcome["images"].append({"path": repo_link(path), "size": picture.size,
                                              "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
            for filename in case.get("workbooks", []):
                path = ASSET_DIR / filename
                outcome["workbooks"].append({"path": repo_link(path), "sheets": verify_workbook(path)})
            if case["id"] == "efficiency":
                rules_path = ASSET_DIR / "efficiency-rules.json"
                rules_path.write_text(json.dumps({"manual_rules": namespace["efficiency"]["manual_rules"],
                                                  "auto_rules": namespace["efficiency"]["auto_rules"]}, ensure_ascii=False, indent=2), encoding="utf-8")
            print("完成：" + case["id"], flush=True)
        except Exception as exc:
            outcome["error"] = f"{type(exc).__name__}: {exc}"
            print("失败：" + outcome["error"], flush=True)
        finally:
            plt.close("all")
        manifest[case["id"]] = outcome
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        write_section(manifest)
    if any(value.get("error") for value in manifest.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
